import base64
import json
import math
import os
import re
import threading
import sys
import traceback
from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}
SOURCE_COLUMNS = [
    "Source File",
    "Source Relative Path",
    "Source Sheet",
    "Source Area",
    "Source Header Row",
]

ACTION_DELETE = "행 삭제"
ACTION_REPLACE = "값 변경"
ACTION_REVIEW = "검토 대상으로 복사"
ACTION_SORT_ASC = "오름차순 정렬"
ACTION_SORT_DESC = "내림차순 정렬"
ACTION_REMOVE_EMPTY = "완전 빈 행 삭제"

OPERATORS = [
    "정확히 일치",
    "포함",
    "포함하지 않음",
    "시작 문자열",
    "끝 문자열",
    "빈칸",
    "빈칸 아님",
    "목록 중 하나",
    "목록 중 하나 포함",
    "숫자 같음",
    "숫자 초과",
    "숫자 이상",
    "숫자 미만",
    "숫자 이하",
]


DISPLAY_FORMAT_HEADERS = [
    "NO", "Part", "Drawing Number",
    "Category1", "Category2", "Category3", "Category4", "Category5",
    "Category6", "Category7", "Category8", "Category9", "Category10",
    "CWP No.", "Fluid", "Serial", "Sheet", "Etc1", "Etc2", "Etc3",
    "Rev", "Insulation Symbol", "Insulation Temp (Operating Temp)",
    "Paint Symbol", "Class", "Item", "Unit Size", "Main", "Sub", "Qty",
    "Design Factor", "Total Qty", "3D BIMWelding_Point(Main)",
    "3D BIMWelding_Point(Sub)", "Field/Shop", "Assembly",
    "Error Description", "Remark", "Remark2",
    "FileName", "CD_SITE", "PrpsProj", "Times", "Purpose", "Mtrl Group",
]

DISPLAY_DEFAULT_MODES = {
    "NO": ("자동번호", ""),
    "Part": ("특수규칙", ""),
    "Drawing Number": ("3D BM 열 복사", "DRAWING NUMBER"),
    "Category1": ("PCWBS 파일 매핑", "Category1"),
    "Category2": ("PCWBS 파일 매핑", "Category2"),
    "Category3": ("PCWBS 파일 매핑", "Category3"),
    "Category4": ("PCWBS 파일 매핑", "Category4"),
    "Category5": ("PCWBS 파일 매핑", "Category5"),
    "Category6": ("빈칸", ""),
    "Category7": ("빈칸", ""),
    "Category8": ("빈칸", ""),
    "Category9": ("빈칸", ""),
    "Category10": ("빈칸", ""),
    "CWP No.": ("3D BM 열 복사", "ISO_DWG_ID"),
    "Fluid": ("3D BM 열 복사", "FLUID"),
    "Serial": ("3D BM 열 복사", "SERIAL"),
    "Sheet": ("3D BM 열 복사", "SHEET"),
    "Etc1": ("빈칸", ""),
    "Etc2": ("빈칸", ""),
    "Etc3": ("빈칸", ""),
    "Rev": ("빈칸", ""),
    "Insulation Symbol": ("3D BM 열 복사", "INS. SPEC"),
    "Insulation Temp (Operating Temp)": ("빈칸", ""),
    "Paint Symbol": ("빈칸", ""),
    "Class": ("3D BM 열 복사", "CLASS"),
    "Item": ("3D BM 열 복사", "SYMBOL"),
    "Unit Size": ("특수규칙", ""),
    "Main": ("3D BM 열 복사", "SIZE-1"),
    "Sub": ("3D BM 열 복사", "SIZE-2"),
    "Qty": ("3D BM 열 복사", "QTY"),
    "Design Factor": ("빈칸", ""),
    "Total Qty": ("빈칸", ""),
    "3D BIMWelding_Point(Main)": ("빈칸", ""),
    "3D BIMWelding_Point(Sub)": ("빈칸", ""),
    "Field/Shop": ("빈칸", ""),
    "Assembly": ("빈칸", ""),
    "Error Description": ("빈칸", ""),
    "Remark": ("3D BM 열 복사", "REMARK"),
    "Remark2": ("빈칸", ""),
    "FileName": ("고정값", ""),
    "CD_SITE": ("고정값", ""),
    "PrpsProj": ("고정값", "2"),
    "Times": ("고정값", ""),
    "Purpose": ("고정값", "PO"),
    "Mtrl Group": ("고정값", "RUSSIA"),
}

MAPPING_MODES = ["3D BM 열 복사", "PCWBS 파일 매핑", "고정값", "빈칸", "자동번호", "특수규칙"]
# v8.12 이전 매핑 JSON 호환용 (입력 방식 이름 변경)
LEGACY_MAPPING_MODES = {"BM 열 복사": "3D BM 열 복사"}
SPECIAL_RULE_HEADERS = ["Part", "Unit Size"]

MAPPING_MODE_HELP = {
    "3D BM 열 복사": "선택한 3D BM 파일의 열 값을 그대로 복사합니다. 아래 칸에서 3D BM 열을 고르세요.",
    "PCWBS 파일 매핑": (
        "3번 탭 비교키(또는 Mapping Table)로 PCWBS 행을 찾아, 아래 칸에 입력한 PCWBS 열 값을 가져옵니다. "
        "(Category1~10만 사용 가능)"
    ),
    "고정값": "아래 칸에 입력한 값을 모든 행에 동일하게 입력합니다.",
    "빈칸": "값을 비워 둡니다.",
    "자동번호": (
        "1부터 행 순서대로 1씩 증가하는 번호를 입력합니다 (별도 설정 없음). "
        "6번 File Split 시에는 분할 파일마다 1부터 다시 매깁니다."
    ),
    "특수규칙": (
        "Part / Unit Size 열 전용 내장 규칙입니다.\n"
        "· Part: 오른쪽 'Part 조건 규칙'을 위에서부터 검사해 처음 일치한 Part 값을 입력, 없으면 'Part 기본값'\n"
        "· Unit Size: 'Unit Size 예외 Item 열' 조건에 일치하면 'U x mm', 아니면 'Unit Size 기본값'"
    ),
}
PART_VALUES = ["PPA", "PPU", "FFA", "FFU", "CIA", "CIU"]

# 3번 PCWBS 비교 방식
COMPARE_METHOD_KEY = "비교키 생성 설정"
COMPARE_METHOD_TABLE = "Mapping Table 사용"

# 5번 Line No. 매칭 방식
LINE_MATCH_EXACT = "완전 일치"
LINE_MATCH_STRIP_SHEET = "Sh't No. 제거 후 일치 (Sh't No. = 마지막 구분자 뒤의 Text)"
LINE_MATCH_DISPLAY_CONTAINS = "Display 값에 Line No. 포함할 경우"
LINE_MATCH_LINE_CONTAINS = "Line No. 값에 Display 값 포함할 경우"
LINE_MATCH_MODES = [
    LINE_MATCH_EXACT,
    LINE_MATCH_STRIP_SHEET,
    LINE_MATCH_DISPLAY_CONTAINS,
    LINE_MATCH_LINE_CONTAINS,
]
# v8.11 이전 규칙 JSON 호환용
LEGACY_LINE_MATCH_MODES = {
    "Display 끝 Sheet No. 제거 후 일치": LINE_MATCH_STRIP_SHEET,
    "Display 값에 Line No. 포함": LINE_MATCH_DISPLAY_CONTAINS,
    "Line No. 값에 Display 값 포함": LINE_MATCH_LINE_CONTAINS,
}

# 5번 Insulation Temperature 입력 방식
INS_TEMP_SOURCE_LINE = "Line List 온도 열"
INS_TEMP_SOURCE_RULE = "Display 열 기준 규칙 파일"
INS_TEMP_SOURCES = [INS_TEMP_SOURCE_LINE, INS_TEMP_SOURCE_RULE]
# 이전 버전에서 저장한 규칙 JSON 호환용
LEGACY_INS_TEMP_SOURCES = {"규칙 파일 우선 → 없으면 Line List": INS_TEMP_SOURCE_RULE}

# 5번 Paint Symbol 입력 방식
PAINT_SOURCE_TABLE = "Painting Code Table 계산"
PAINT_SOURCE_RULE = "Display 열 기준 규칙 파일"
PAINT_SOURCES = [PAINT_SOURCE_TABLE, PAINT_SOURCE_RULE]


# -----------------------------------------------------------------------------
# 공통 유틸리티
# -----------------------------------------------------------------------------

def clean_text(value):
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    return re.sub(r"\s+", " ", text)


def canonical_header(value):
    text = clean_text(value)
    return re.sub(r"[\s_\-./\\]+", "", text).casefold()


def is_effectively_empty_row(values):
    return all(clean_text(value) == "" for value in values)


def make_unique_headers(values):
    headers = []
    counts = Counter()

    for index, value in enumerate(values, start=1):
        header = clean_text(value) or f"Unnamed_Column_{index}"
        key = canonical_header(header)
        counts[key] += 1
        if counts[key] > 1:
            header = f"{header}__{counts[key]}"
        headers.append(header)

    return headers


def detect_header_row(ws, scan_limit=50):
    max_row = min(ws.max_row, scan_limit)
    max_col = ws.max_column
    best = None

    for row_number in range(1, max_row + 1):
        values = [
            ws.cell(row=row_number, column=column_number).value
            for column_number in range(1, max_col + 1)
        ]
        non_empty = [clean_text(value) for value in values if clean_text(value)]
        if len(non_empty) < 2:
            continue

        text_count = sum(
            1 for value in values
            if isinstance(value, str) and clean_text(value)
        )
        unique_count = len({canonical_header(value) for value in non_empty})
        numeric_count = sum(
            1 for value in values
            if isinstance(value, (int, float)) and not isinstance(value, bool)
        )
        score = (
            len(non_empty) * 3
            + text_count * 2
            + unique_count
            - numeric_count
            - (row_number - 1) * 0.15
        )
        candidate = (score, -row_number, row_number, len(non_empty))
        if best is None or candidate > best:
            best = candidate

    if best is None:
        return None, 0, 0
    return best[2], best[0], best[3]


def find_last_nonempty_column(ws, header_row):
    for column_number in range(ws.max_column, 0, -1):
        if clean_text(ws.cell(row=header_row, column=column_number).value):
            return column_number
    return 0


def derive_area_from_filename(file_path):
    return file_path.stem


def safe_relative_path(file_path, base_folder):
    try:
        return str(file_path.relative_to(base_folder))
    except ValueError:
        return str(file_path)


def auto_fit_columns(ws, min_width=10, max_width=45, scan_rows=500):
    for column_number in range(1, ws.max_column + 1):
        longest = 0
        for row_number in range(1, min(ws.max_row, scan_rows) + 1):
            value = ws.cell(row=row_number, column=column_number).value
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(column_number)].width = min(
            max(longest + 2, min_width), max_width
        )


def style_header(ws):
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions


def normalize_numeric(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    text = clean_text(value).replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_list_value(text):
    return [item.strip() for item in re.split(r"[,;\n]", text or "") if item.strip()]


def compare_value(cell_value, operator, rule_value, case_sensitive=False):
    cell_text = clean_text(cell_value)
    target_text = clean_text(rule_value)

    if not case_sensitive:
        left = cell_text.casefold()
        right = target_text.casefold()
    else:
        left = cell_text
        right = target_text

    if operator == "정확히 일치":
        return left == right
    if operator == "포함":
        return right in left
    if operator == "포함하지 않음":
        return right not in left
    if operator == "시작 문자열":
        return left.startswith(right)
    if operator == "끝 문자열":
        return left.endswith(right)
    if operator == "빈칸":
        return cell_text == ""
    if operator == "빈칸 아님":
        return cell_text != ""
    if operator == "목록 중 하나":
        candidates = parse_list_value(rule_value)
        if not case_sensitive:
            candidates = [item.casefold() for item in candidates]
        return left in candidates
    if operator == "목록 중 하나 포함":
        candidates = parse_list_value(rule_value)
        if not case_sensitive:
            candidates = [item.casefold() for item in candidates]
        return any(candidate in left for candidate in candidates)

    left_num = normalize_numeric(cell_value)
    right_num = normalize_numeric(rule_value)
    if left_num is None or right_num is None:
        return False

    if operator == "숫자 같음":
        return left_num == right_num
    if operator == "숫자 초과":
        return left_num > right_num
    if operator == "숫자 이상":
        return left_num >= right_num
    if operator == "숫자 미만":
        return left_num < right_num
    if operator == "숫자 이하":
        return left_num <= right_num

    return False


def write_rows_to_sheet(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    style_header(ws)
    auto_fit_columns(ws)


# -----------------------------------------------------------------------------
# 메인 애플리케이션
# -----------------------------------------------------------------------------

class BmDmcsTool:
    def __init__(self, root):
        self.root = root
        self.root.title("Piping Engineering - 3D BIM BM to Display Format Tool v8.13 Beta")
        self.root.geometry("1520x980")
        self.root.minsize(1200, 820)

        # In a PyInstaller one-file executable, __file__ points to a temporary
        # extraction directory. Use the executable folder so all input/output
        # folders remain next to the distributed EXE.
        if getattr(sys, "frozen", False):
            self.project_root = Path(sys.executable).resolve().parent
        else:
            self.project_root = Path(__file__).resolve().parent

        for folder_name in ["02_test_input", "03_reference", "04_output", "05_logs"]:
            (self.project_root / folder_name).mkdir(parents=True, exist_ok=True)

        self.is_running = False
        self.rules = []
        self.cleanup_headers = []
        self.cleanup_sheet_names = []
        self.pcwbs_bm_headers = []
        self.pcwbs_ref_headers = []
        self.display_bm_headers = []
        self.display_mappings = {}
        self.part_rules = []
        self.material_rules = []
        self.insulation_rules = []

        self.status_text = tk.StringVar(value="대기 중")
        self.current_item_text = tk.StringVar(value="현재 작업: 없음")
        self.progress_text = tk.StringVar(value="진행률: 0%")

        self._setup_modern_theme()
        self._set_window_icon()
        self._build_ui()
        self._apply_widget_theme(self.root)
        self._restore_sidebar_status_theme()
        self._show_page("merge")

    # ------------------------------------------------------------------ UI
    def _setup_modern_theme(self):
        """Configure a clean engineering-tool theme without changing business logic."""
        self.colors = {
            "navy": "#0B2A4A",
            "navy_hover": "#123E68",
            "blue": "#1F5F99",
            "blue_light": "#E8F1F8",
            "accent": "#2D74B8",
            "background": "#EEF1F4",
            "surface": "#FFFFFF",
            "surface_alt": "#F7F9FB",
            "border": "#D4DAE1",
            "text": "#17212B",
            "muted": "#65727E",
            "success": "#1F8A5B",
            "danger": "#D9363E",
            "danger_hover": "#B9272E",
            "gray_button": "#667381",
            "gray_hover": "#4F5B67",
        }

        self.default_font = ("Malgun Gothic", 9)
        self.small_font = ("Malgun Gothic", 8)
        self.bold_font = ("Malgun Gothic", 9, "bold")
        self.title_font = ("Malgun Gothic", 18, "bold")
        self.section_font = ("Malgun Gothic", 10, "bold")

        self.root.configure(bg=self.colors["background"])
        self.root.option_add("*Font", self.default_font)
        self.root.option_add("*tearOff", False)

        style = ttk.Style(self.root)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        style.configure(
            "TNotebook",
            background=self.colors["background"],
            borderwidth=0,
            tabmargins=(0, 0, 0, 0),
        )
        style.configure(
            "TNotebook.Tab",
            background="#DCE3EA",
            foreground=self.colors["text"],
            padding=(18, 9),
            font=self.bold_font,
            borderwidth=0,
        )
        style.map(
            "TNotebook.Tab",
            background=[
                ("selected", self.colors["navy"]),
                ("active", self.colors["blue_light"]),
            ],
            foreground=[
                ("selected", "#FFFFFF"),
                ("active", self.colors["navy"]),
            ],
        )

        style.configure(
            "TCombobox",
            fieldbackground=self.colors["surface"],
            background=self.colors["surface"],
            foreground=self.colors["text"],
            arrowcolor=self.colors["navy"],
            bordercolor=self.colors["border"],
            lightcolor=self.colors["border"],
            darkcolor=self.colors["border"],
            padding=4,
        )
        style.map(
            "TCombobox",
            fieldbackground=[("readonly", self.colors["surface"])],
            selectbackground=[("readonly", self.colors["blue_light"])],
            selectforeground=[("readonly", self.colors["text"])],
        )

        style.configure(
            "Treeview",
            background=self.colors["surface"],
            fieldbackground=self.colors["surface"],
            foreground=self.colors["text"],
            bordercolor=self.colors["border"],
            rowheight=25,
            font=self.default_font,
        )
        style.configure(
            "Treeview.Heading",
            background=self.colors["navy"],
            foreground="#FFFFFF",
            relief="flat",
            padding=(6, 7),
            font=self.bold_font,
        )
        style.map(
            "Treeview",
            background=[("selected", self.colors["blue"])],
            foreground=[("selected", "#FFFFFF")],
        )

        style.configure(
            "TProgressbar",
            troughcolor="#D9E0E7",
            background=self.colors["accent"],
            bordercolor="#D9E0E7",
            lightcolor=self.colors["accent"],
            darkcolor=self.colors["accent"],
            thickness=13,
        )
        style.configure(
            "Vertical.TScrollbar",
            background="#B9C4CE",
            troughcolor=self.colors["surface_alt"],
            arrowcolor=self.colors["navy"],
        )
        style.configure(
            "Horizontal.TScrollbar",
            background="#B9C4CE",
            troughcolor=self.colors["surface_alt"],
            arrowcolor=self.colors["navy"],
        )

    def _build_header_bar(self):
        header = tk.Frame(
            self.root,
            bg=self.colors["navy"],
            height=82,
            bd=0,
            highlightthickness=0,
        )
        header.grid(row=0, column=0, sticky="ew")
        header.grid_propagate(False)
        header.grid_columnconfigure(1, weight=1)

        brand = tk.Frame(header, bg=self.colors["navy"])
        brand.grid(row=0, column=0, rowspan=2, sticky="nsw", padx=(22, 25))
        tk.Label(
            brand,
            text="PIPING",
            bg=self.colors["navy"],
            fg="#FFFFFF",
            font=("Arial", 15, "bold"),
            anchor="w",
        ).pack(anchor="w", pady=(13, 0))
        tk.Label(
            brand,
            text="ENGINEERING",
            bg=self.colors["navy"],
            fg="#FFFFFF",
            font=("Arial", 15, "bold"),
            anchor="w",
        ).pack(anchor="w")

        tk.Label(
            header,
            text="3D BIM BM → Display Format Automation",
            bg=self.colors["navy"],
            fg="#FFFFFF",
            font=self.title_font,
            anchor="w",
        ).grid(row=0, column=1, sticky="sw", pady=(11, 0))

        tk.Label(
            header,
            text="Excel 병합 · 데이터 정제 · PCWBS 검증 · 온도/Paint 자동 매핑 · 분할 Export",
            bg=self.colors["navy"],
            fg="#C7D7E6",
            font=("Malgun Gothic", 9),
            anchor="w",
        ).grid(row=1, column=1, sticky="nw", pady=(2, 9))

        version_box = tk.Frame(
            header,
            bg="#163E63",
            padx=12,
            pady=7,
        )
        version_box.grid(row=0, column=2, rowspan=2, sticky="e", padx=22)
        tk.Label(
            version_box,
            text="REV. 8.0",
            bg="#163E63",
            fg="#FFFFFF",
            font=("Arial", 10, "bold"),
        ).pack()
        tk.Label(
            version_box,
            text="UI Preview",
            bg="#163E63",
            fg="#BFD4E6",
            font=("Arial", 8),
        ).pack()

    def _set_window_icon(self):
        """Set a compact window icon using the supplied DL logo."""
        try:
            self._window_icon = tk.PhotoImage(data="""iVBORw0KGgoAAAANSUhEUgAAAIAAAABnCAYAAADIf0rxAAAI30lEQVR4nO2dbYgkRxnHf1XdMz09va+3R0IueDEJaCBH5O4S7i4JGo54okKCKIpwMSIcvly4SDBfPHyJORMjRIwxnOIXlbxIPihGjxBEL4LGN4JBOUXy6dQ79+6yubvdnZnd2a4qP8zLzs7ObPfuzGa2t+oHyw4zVV1Pd/276nmqq6qFMcbgsBY5aAMcg8UJwHKcACzHCcBynAAsxwnAcpwALMdPSlCtVpFSIqXTStZIU2ddBWCMQQjB1NQUQRAwNDSEUgohxLI0SaRJt1KatOWkoZ/H2ujl5PP5RBEkSkRK2SygvaC0J5gm3Upp+nkh34pKyVI5iQJwI8WbG9exW06iAN6qpswxGFwXYDmuBbAc1wJYjnMCLccJwHKcACzHCcBynAAsxwnAcpwALMcJwHKcACzHCcBynAAsxwnAcpwALMcJwHKcACwncV3ARsMYs6o5CkKI5qSWjnmFQNCfiS+92DYoMieAfC6H5/tAtwstWn4TVKtVtNYA5HwfP5dbktdoTawUcaxQSiOlQAjJWuol2baldi5Uq6i6bYMiMwIwxuB5kjOTFzh/YQrpeRgMixXetmYB0Fpz3TVXExVDAC68eYkzZ8/V8hqDlJKoWGBibJSJLaMgaj3i3NwcsVJ40kslhMaxzp57g8nzb+B5XqvEaBdlw7a3v+0qRoYjlDYMqh3IjACU0oRhyA+e+TnHHjlOcesYcay6phdCMD83z8mfHueOfTtBSE78+hUOHf4K4fhoc5VTkM8zPjrEtdu3sXfXjRx4zx5uveUmhqKI+fl54ljheSu7Sg3bfvT8Cb700HcobhkjVt1tk1IyVyrzy2e/zQfvvJ1SqZxYxnqRKIBB91HtCCGgvlZRyu5NrRAC0bYsqjVvo6+em6/y3/9d4PTps7x88o889tTT7NrxDg4dvJt7PvIBoqhIqVRCel7iXbrEthV8ASkltKy4GiTZnBQqOn7slmRZj1xz+mqfPU+Sz+coDkcMTYwRhgGv/v1ffObzx7jtrkO8dPIPRFEEq3TwmmV1qOTBV/si2QwDTceP3ZJ0/K1Rlw3PXWtNHCu0NhSLIUMTY7x26nXe//H7OfrocYIgQEqJXqUIOolmI91Sbl1ABxpiiIohYRTyyDeOc/DwlxFC4LV0H5uBbHYBq2StElZaY4xh+MqtPPfsL7j3yNfw/ZrbtAkuC5DVLmCV9FJXxsDCQszwlRP85LkXOPrY9wjDEK27e/lZwgoB9IN4ISbauoVvPvFDXvzN74miCLVCqJcVrBBAP7yY5jCOlDz48HeZmS3h+37muwIrBNAvlNYUhyJOvfZPfvz8CYIgQGW8K3ACWCVaa2Qh4PvPvEClUiHnZ2YwtSNOAKtEa02hWODUP17nd3/5W70VGOwDnV5wAlgDUkj0fJVf/fbPQLZDZSeANWCMgZzPn/56Cq0UXob3UMyu5QNEG4PM5Tj9n0kuXp7G971Bm7RmrBDAejTQvie5NDPLxcsz+J4TwIam308zGhNAqtUFKpX5TG+jm13LV8F6tABmcYJPprFCAP2uIyEExmgKQZ4oLLgw0EZipRgbGWZ8bCTTzwSyPYw1IKQU6GrMdddsY2xkmJlSmbDHYxpj0k82WT4HdgmtU1CTcAJYAwIBCzH7du9ASNmcdr5WDIZcLofv5/pkYXqcANaA0hq/WODAHXvAGGSPs6bCQoGXX3mVx596mnAoRGmz/C5vczoFNUdUiMX/S2xUmicefoBrt29b8Z0CTgCrxJOSSqnMzbt3sGfnjczN9xoGGjzP499nzvHiz16C8RGIe2hRGu2/0jz04KFaCU4AfUQI9ELM4Xs/TD6fp1Qq9eWw+VwOb2yUaHQYpTRLb/n2xSWJRqKUwksxQGWFABJ8ptR4nqQ8M8uevTv52N13Mjc3h/Q81AoLVJJZXLeolKr/9R5WKqVTPaRyYWBKhBBopcnlcnzrq0fI5/MorfsgrPV7kihSWGfFtPB+3f2Vi5f5+tHPcest76JcLm/4p4AmhbismBbeyxkIIfA9j9nzU9x3+B6+8NmDlMrlVP1rFtjYEh4wvu+htWZ26iL3H/kkTx57gEql0nPYt5HI3OLQtbDaM2is1J29NENxKOTJx7/IfZ/6KJVKpXa8TXBNGiQKYCN2AULUHJykdw02dwahviC05fvWvM3vgAWlKF2eBQEH9u/j0aOH2XXTDZTKZWTSjh5iaRlpbFuSd4ltvcUuaTWauTDQGINRGqVrIVOnCLlRCbr+kKZxLbQxmGaopTAGtNG1MK66AMYQjY+w/7238elPfIi73vduhBDMlkqpJn0YXbdNLdrWSsNOr25H7eYSLXkXw8BueRvn0j5K0JoOIE4ZSmZOAPlcjjAKiYph4gYRnuchpWheuJzvEUZFomKIUgopBWEh4IqJcd55/Xb27t7B/ttvZscN1wOCSqWCMSb1jJ9cq20JG0SI5nFr1vm+18zbr3GANCOUwnRp4xvDh5OTkwRBQBRFPT/06BUhBNMzJaZnSrVNHpL8e2PYOjFGkM8DUCpXuHhppikKT0rCMGB0eIggCOp5NOXKXH1LmvSevhCC6dkS09Ol5vFXNs2wdcsohXq55cocb16aRorkvKkwcPVVVxAE+ZXtzpIAoOag+Z7XeUlW+6ipqC3sbJyilLI2gbOZxqC1IW42u7UHO2sd2/c8WRNN+9ZF7bOH6t/HcYzWHWxLov2YncoxtWVsSQ5r5rqAOK7t6NXlmi67fq0hm1JqSdPc+KXWXfQeEcexZqFDt9SpCkwX25LOp90H6FaOMVAoBJtPAM03mbd/3/a/W971DOAa0UmqtMvyLs2ZdD7J+xWlMsMNBNmOE4DlOAFYjhOA5TgBWI4TgOU4AViOE4DlOAFYjhOA5TgBWI4TgOU4AViOE4DlOAFYjhUrgxzdsWJlkKM7rguwHNcFWI5rASwnUQBaL2400OoPpPUNVutDtL+Aeb3KWWv+zVZO13UBDarVav0tna6xyBpp3k6eKADH5sbd1pbjBGA5TgCW4wRgOU4AluMEYDlOAJbzfxHkh9sUC7F1AAAAAElFTkSuQmCC""")
            self.root.iconphoto(True, self._window_icon)
        except Exception:
            self._window_icon = None

    def _button_style_for_text(self, text):
        label = clean_text(text).lower()
        if any(word in label for word in ["삭제", "지우기", "clear"]):
            return self.colors["danger"], self.colors["danger_hover"], "#FFFFFF"
        if any(word in label for word in ["실행", "생성", "최종", "저장"]):
            return self.colors["navy"], self.colors["navy_hover"], "#FFFFFF"
        if any(word in label for word in ["선택", "불러오기", "browse", "미리보기", "적용", "추가"]):
            return self.colors["gray_button"], self.colors["gray_hover"], "#FFFFFF"
        return "#E3E8ED", "#D2DAE2", self.colors["text"]

    def _apply_widget_theme(self, widget):
        """Recursively style existing tkinter widgets while preserving their layout."""
        try:
            children = widget.winfo_children()
        except tk.TclError:
            return

        for child in children:
            class_name = child.winfo_class()

            try:
                if isinstance(child, tk.LabelFrame):
                    child.configure(
                        bg=self.colors["surface"],
                        fg=self.colors["navy"],
                        font=self.section_font,
                        bd=0,
                        relief="flat",
                        highlightbackground="#E2E8F0",
                        highlightcolor="#E2E8F0",
                        highlightthickness=1,
                    )
                elif isinstance(child, tk.Frame):
                    # Header / sidebar frames keep their explicit navy tones.
                    if child.cget("bg") in {
                        self.colors["navy"],
                        "#163E63",
                    }:
                        pass
                    else:
                        parent = child.master
                        if isinstance(parent, tk.LabelFrame):
                            child.configure(bg=self.colors["surface"])
                        else:
                            child.configure(bg=self.colors["background"])
                elif isinstance(child, tk.Button):
                    bg, active_bg, fg = self._button_style_for_text(
                        child.cget("text")
                    )
                    child.configure(
                        bg=bg,
                        activebackground=active_bg,
                        fg=fg,
                        activeforeground=fg,
                        relief="flat",
                        bd=0,
                        cursor="hand2",
                        padx=max(child.winfo_pixels(child.cget("padx") or 0), 8),
                        pady=max(child.winfo_pixels(child.cget("pady") or 0), 4),
                        highlightthickness=0,
                        font=self.bold_font,
                    )
                elif isinstance(child, tk.Entry):
                    child.configure(
                        bg=self.colors["surface"],
                        fg=self.colors["text"],
                        insertbackground=self.colors["navy"],
                        relief="solid",
                        bd=1,
                        highlightthickness=1,
                        highlightbackground=self.colors["border"],
                        highlightcolor=self.colors["accent"],
                    )
                elif isinstance(child, tk.Listbox):
                    child.configure(
                        bg=self.colors["surface"],
                        fg=self.colors["text"],
                        selectbackground=self.colors["blue"],
                        selectforeground="#FFFFFF",
                        relief="solid",
                        bd=1,
                        highlightthickness=1,
                        highlightbackground=self.colors["border"],
                    )
                elif isinstance(child, scrolledtext.ScrolledText):
                    child.configure(
                        bg="#F8FAFC",
                        fg=self.colors["text"],
                        insertbackground=self.colors["navy"],
                        relief="solid",
                        bd=1,
                        highlightthickness=1,
                        highlightbackground=self.colors["border"],
                    )
                elif isinstance(child, tk.Text):
                    child.configure(
                        bg="#F8FAFC",
                        fg=self.colors["text"],
                        insertbackground=self.colors["navy"],
                    )
                elif isinstance(child, tk.Checkbutton):
                    child.configure(
                        bg=self.colors["surface"],
                        fg=self.colors["text"],
                        activebackground=self.colors["surface"],
                        activeforeground=self.colors["navy"],
                        selectcolor=self.colors["surface"],
                    )
                elif isinstance(child, tk.Radiobutton):
                    child.configure(
                        bg=self.colors["surface"],
                        fg=self.colors["text"],
                        activebackground=self.colors["surface"],
                        activeforeground=self.colors["navy"],
                        selectcolor=self.colors["surface"],
                    )
                elif isinstance(child, tk.Label):
                    current_bg = child.cget("bg")
                    if current_bg not in {
                        self.colors["navy"],
                        "#163E63",
                    }:
                        parent = child.master
                        parent_bg = (
                            parent.cget("bg")
                            if "bg" in parent.keys()
                            else self.colors["background"]
                        )
                        child.configure(
                            bg=parent_bg,
                            fg=(
                                child.cget("fg")
                                if child.cget("fg") not in {"#000000", "black"}
                                else self.colors["text"]
                            ),
                        )
            except (tk.TclError, ValueError):
                pass

            self._apply_widget_theme(child)

    def start_thread(self, target):
        """Run a long operation in a daemon thread so the GUI remains responsive."""
        threading.Thread(target=target, daemon=True).start()

    def _build_sidebar(self):
        self.page_title_var = tk.StringVar(value="1. Excel 병합")
        self.page_subtitle_var = tk.StringVar(
            value="여러 Area BM 파일을 자동으로 통합합니다."
        )

        self.sidebar_frame = tk.Frame(
            self.root,
            bg=self.colors["navy"],
            width=220,
            padx=10,
            pady=12,
        )
        self.sidebar_frame.grid(row=0, column=0, sticky="ns")
        self.sidebar_frame.grid_propagate(False)
        self.sidebar_frame.grid_rowconfigure(8, weight=1)

        tk.Label(
            self.sidebar_frame,
            text="PIPING\nENGINEERING",
            bg=self.colors["navy"],
            fg="#FFFFFF",
            font=("Arial", 19, "bold"),
            justify="left",
            anchor="w",
        ).grid(row=0, column=0, sticky="ew", padx=10, pady=(6, 18))

        self.sidebar_buttons = {}
        self.page_meta = [
            ("merge", "1. Excel 병합", "여러 Area BM 파일을 자동으로 통합합니다."),
            ("cleanup", "2. 병합 파일 수정", "병합 파일의 불필요 행 삭제 및 데이터를 수정합니다."),
            ("pcwbs", "3. PCWBS 비교", "PCWBS 기준과 3D BM 조합키를 비교합니다."),
            ("display", "4. Display Format 매핑", "Display Format 기본 파일을 생성합니다."),
            ("paint", "5. Temperature & Painting", "온도 및 Paint Symbol을 자동 적용합니다."),
            ("split", "6. File Split & Export", "최종 파일 분할 및 Export를 수행합니다."),
        ]

        for idx, (key, title, subtitle) in enumerate(self.page_meta, start=1):
            btn = tk.Button(
                self.sidebar_frame,
                text=title,
                command=lambda page_key=key: self._show_page(page_key),
                anchor="w",
                justify="left",
                relief="flat",
                bd=0,
                padx=14,
                pady=10,
                font=("Malgun Gothic", 10, "bold"),
                cursor="hand2",
                width=22,
            )
            btn.grid(row=idx, column=0, sticky="ew", padx=8, pady=4)
            self.sidebar_buttons[key] = btn

        footer = tk.Frame(
            self.sidebar_frame,
            bg="#163E63",
            padx=10,
            pady=9,
        )
        footer.grid(row=10, column=0, sticky="ew", padx=8, pady=(0, 8))
        tk.Label(
            footer,
            text="REV. 8.13 BETA",
            bg="#163E63",
            fg="#FFFFFF",
            font=("Arial", 10, "bold"),
        ).pack(anchor="w")
        tk.Label(
            footer,
            text="Beta Test Build",
            bg="#163E63",
            fg="#C3D4E5",
            font=("Arial", 8),
        ).pack(anchor="w", pady=(2, 0))

    def _show_page(self, page_key):
        page_lookup = {
            "merge": self.merge_tab,
            "cleanup": self.cleanup_tab,
            "pcwbs": self.pcwbs_tab,
            "display": self.display_tab,
            "paint": self.paint_tab,
            "split": self.split_tab,
        }
        titles = {key: (title, subtitle) for key, title, subtitle in self.page_meta}

        for key, frame in page_lookup.items():
            if key == page_key:
                frame.grid()
            else:
                frame.grid_remove()

        for key, button in self.sidebar_buttons.items():
            if key == page_key:
                button.configure(
                    bg=self.colors["accent"],
                    activebackground=self.colors["blue"],
                    fg="#FFFFFF",
                    activeforeground="#FFFFFF",
                )
            else:
                button.configure(
                    bg=self.colors["navy"],
                    activebackground=self.colors["navy_hover"],
                    fg="#FFFFFF",
                    activeforeground="#FFFFFF",
                )

        title, subtitle = titles.get(page_key, ("", ""))
        self.page_title_var.set(title)
        self.page_subtitle_var.set(subtitle)

    def _build_ui(self):
        self.root.grid_rowconfigure(0, weight=1)
        self.root.grid_columnconfigure(0, weight=0)
        self.root.grid_columnconfigure(1, weight=1)

        self._build_sidebar()

        content_shell = tk.Frame(self.root, bg=self.colors["background"])
        content_shell.grid(row=0, column=1, sticky="nsew", padx=(10, 12), pady=(12, 6))
        content_shell.grid_rowconfigure(1, weight=1)
        content_shell.grid_columnconfigure(0, weight=1)

        header_card = tk.Frame(
            content_shell,
            bg=self.colors["surface"],
            bd=0,
            relief="flat",
            highlightbackground=self.colors["border"],
            highlightthickness=1,
        )
        header_card.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        header_card.grid_columnconfigure(1, weight=1)

        tk.Label(
            header_card,
            textvariable=self.page_title_var,
            bg=self.colors["surface"],
            fg=self.colors["navy"],
            font=self.title_font,
            anchor="w",
        ).grid(row=0, column=0, sticky="w", padx=(18, 10), pady=12)

        tk.Label(
            header_card,
            textvariable=self.page_subtitle_var,
            bg=self.colors["surface"],
            fg=self.colors["muted"],
            font=("Malgun Gothic", 9),
            anchor="w",
        ).grid(row=0, column=1, sticky="w", padx=(0, 18), pady=14)

        self.page_container = tk.Frame(content_shell, bg=self.colors["background"])
        self.page_container.grid(row=1, column=0, sticky="nsew")
        self.page_container.grid_rowconfigure(0, weight=1)
        self.page_container.grid_columnconfigure(0, weight=1)

        self.merge_tab = tk.Frame(self.page_container, bg=self.colors["background"])
        self.cleanup_tab = tk.Frame(self.page_container, bg=self.colors["background"])
        self.pcwbs_tab = tk.Frame(self.page_container, bg=self.colors["background"])
        self.display_tab = tk.Frame(self.page_container, bg=self.colors["background"])
        self.paint_tab = tk.Frame(self.page_container, bg=self.colors["background"])
        self.split_tab = tk.Frame(self.page_container, bg=self.colors["background"])

        for frame in [
            self.merge_tab,
            self.cleanup_tab,
            self.pcwbs_tab,
            self.display_tab,
            self.paint_tab,
            self.split_tab,
        ]:
            frame.grid(row=0, column=0, sticky="nsew")

        self._build_merge_tab()
        self._build_cleanup_tab()
        self._build_pcwbs_tab()
        self._build_display_tab()
        self._build_paint_tab()
        self._build_split_tab()
        self._build_status_area()
        self._show_page("merge")

    def _build_pcwbs_tab(self):
        default_output = self.project_root / "04_output"
        self.pcwbs_bm_file = tk.StringVar(value=str(default_output / "BM_Modified_Result.xlsx"))
        self.pcwbs_bm_sheet = tk.StringVar(value="Final_Data")
        self.pcwbs_bm_header_row = tk.StringVar(value="1")
        self.pcwbs_ref_file = tk.StringVar(value="")
        self.pcwbs_ref_sheet = tk.StringVar(value="")
        self.pcwbs_ref_header_row = tk.StringVar(value="1")
        self.pcwbs_output_file = tk.StringVar(value=str(default_output / "PCWBS_Check_Result.xlsx"))

        self.pcwbs_bm_key_mode = tk.StringVar(value="열 조합")
        self.pcwbs_bm_subtitle = tk.StringVar(value="SUBTITLE")
        self.pcwbs_bm_cia = tk.StringVar(value="CIA")
        self.pcwbs_bm_iso = tk.StringVar(value="ISO_DWG_ID")
        self.pcwbs_mid_start = tk.StringVar(value="19")
        self.pcwbs_mid_length = tk.StringVar(value="7")
        self.pcwbs_key_separator = tk.StringVar(value="-")
        self.pcwbs_result_column = tk.StringVar(value="3D BM Match Result")
        self.pcwbs_status_text = tk.StringVar(value="파일과 열을 불러온 뒤 미리보기를 실행하세요.")

        # PCWBS 조합키: 최대 5개 구성요소를 사용자가 순서대로 설정
        self.pcwbs_component_enabled = [tk.BooleanVar(value=(i < 2)) for i in range(5)]
        self.pcwbs_component_column = [tk.StringVar(value="") for _ in range(5)]
        self.pcwbs_component_extract = [tk.StringVar(value="원문") for _ in range(5)]
        self.pcwbs_component_pad = [tk.StringVar(value="0") for _ in range(5)]
        self.pcwbs_component_column[0].set("Category5")
        self.pcwbs_component_column[1].set("Category4")
        self.pcwbs_component_extract[1].set("마지막 괄호 안")
        self.pcwbs_component_pad[0].set("2")

        # 사용자 정의 BM 조합키: 최대 5개 구성요소
        self.bm_component_enabled = [tk.BooleanVar(value=(i < 2)) for i in range(5)]
        self.bm_component_column = [tk.StringVar(value="") for _ in range(5)]
        self.bm_component_extract = [tk.StringVar(value="원문") for _ in range(5)]
        self.bm_component_pad = [tk.StringVar(value="0") for _ in range(5)]
        self.bm_component_column[0].set("SUBTITLE")
        self.bm_component_column[1].set("CIA")
        self.bm_component_pad[0].set("2")

        files_frame = tk.LabelFrame(self.pcwbs_tab, text="비교 파일", padx=10, pady=8)
        files_frame.pack(fill="x", padx=10, pady=(10, 5))
        files_frame.columnconfigure(1, weight=1)
        self._file_row(files_frame, 0, "수정 완료 3D BM", self.pcwbs_bm_file, self.choose_pcwbs_bm_file)
        tk.Label(files_frame, text="3D BM 시트").grid(row=1, column=0, sticky="w", pady=4)
        self.pcwbs_bm_sheet_combo = ttk.Combobox(files_frame, textvariable=self.pcwbs_bm_sheet, state="readonly", width=28)
        self.pcwbs_bm_sheet_combo.grid(row=1, column=1, sticky="w", padx=6)
        self.pcwbs_bm_sheet_combo.bind("<<ComboboxSelected>>", lambda _e: self.load_pcwbs_bm_headers())
        tk.Label(files_frame, text="헤더 행").grid(row=1, column=2, sticky="e")
        tk.Entry(files_frame, textvariable=self.pcwbs_bm_header_row, width=6).grid(row=1, column=3, sticky="w", padx=4)
        tk.Button(files_frame, text="3D BM 열 불러오기", command=self.load_pcwbs_bm_headers, width=16).grid(row=1, column=4, padx=4)

        self._file_row(files_frame, 2, "PCWBS 기준 파일", self.pcwbs_ref_file, self.choose_pcwbs_ref_file)
        tk.Label(files_frame, text="PCWBS 시트").grid(row=3, column=0, sticky="w", pady=4)
        self.pcwbs_ref_sheet_combo = ttk.Combobox(files_frame, textvariable=self.pcwbs_ref_sheet, state="readonly", width=28)
        self.pcwbs_ref_sheet_combo.grid(row=3, column=1, sticky="w", padx=6)
        self.pcwbs_ref_sheet_combo.bind("<<ComboboxSelected>>", lambda _e: self.load_pcwbs_ref_headers())
        tk.Label(files_frame, text="헤더 행").grid(row=3, column=2, sticky="e")
        tk.Entry(files_frame, textvariable=self.pcwbs_ref_header_row, width=6).grid(row=3, column=3, sticky="w", padx=4)
        tk.Button(files_frame, text="PCWBS 열 불러오기", command=self.load_pcwbs_ref_headers, width=16).grid(row=3, column=4, padx=4)
        self._file_row(files_frame, 4, "결과 파일", self.pcwbs_output_file, self.choose_pcwbs_output_file, save=True)

        # 비교 방식 선택: 비교키 생성 설정 / Mapping Table 중 하나만 활성화
        self.pcwbs_compare_method = tk.StringVar(value=COMPARE_METHOD_KEY)
        self.pcwbs_map_file = tk.StringVar(value="")
        self.pcwbs_map_sheet = tk.StringVar(value="")
        self.pcwbs_map_header_row = tk.StringVar(value="1")
        self.pcwbs_map_separator = tk.StringVar(value="-")
        self.pcwbs_map_info = tk.StringVar(
            value="Mapping Table을 불러오면 1행 헤더(3D BM 열 / PCWBS 열)와 등록 건수가 표시됩니다."
        )

        method_frame = tk.LabelFrame(self.pcwbs_tab, text="비교 방식 선택", padx=10, pady=4)
        method_frame.pack(fill="x", padx=10, pady=(5, 0))
        for method in [COMPARE_METHOD_KEY, COMPARE_METHOD_TABLE]:
            tk.Radiobutton(
                method_frame,
                text=method,
                value=method,
                variable=self.pcwbs_compare_method,
                command=self.update_pcwbs_compare_method_ui,
                font=self.bold_font,
            ).pack(side="left", padx=(0, 24))
        tk.Label(
            method_frame,
            text="선택한 방식만 활성화되고, 나머지 설정은 비활성화됩니다. (4번 Category 매핑에도 동일 적용)",
            fg="#555555",
        ).pack(side="left")

        self.pcwbs_mapping_frame = tk.LabelFrame(
            self.pcwbs_tab, text="Mapping Table (3D BM 값 → PCWBS 값)", padx=10, pady=5
        )
        self.pcwbs_mapping_frame.pack(fill="x", padx=10, pady=(5, 0))
        self.pcwbs_mapping_frame.columnconfigure(1, weight=1)
        self._file_row(
            self.pcwbs_mapping_frame, 0, "Mapping Table", self.pcwbs_map_file, self.choose_pcwbs_map_file
        )
        tk.Label(self.pcwbs_mapping_frame, text="Mapping 시트").grid(row=1, column=0, sticky="w", pady=3)
        self.pcwbs_map_sheet_combo = ttk.Combobox(
            self.pcwbs_mapping_frame, textvariable=self.pcwbs_map_sheet, state="readonly", width=28
        )
        self.pcwbs_map_sheet_combo.grid(row=1, column=1, sticky="w", padx=6)
        tk.Label(self.pcwbs_mapping_frame, text="헤더 행").grid(row=1, column=2, sticky="e")
        tk.Entry(self.pcwbs_mapping_frame, textvariable=self.pcwbs_map_header_row, width=6).grid(
            row=1, column=3, sticky="w", padx=4
        )
        tk.Button(
            self.pcwbs_mapping_frame, text="Mapping Table 예시", command=self.show_pcwbs_mapping_example, width=16
        ).grid(row=1, column=4, padx=4)
        tk.Button(
            self.pcwbs_mapping_frame, text="Mapping Table 불러오기", command=self.preview_pcwbs_mapping_table, width=18
        ).grid(row=1, column=5, padx=4)
        separator_frame = tk.Frame(self.pcwbs_mapping_frame)
        separator_frame.grid(row=2, column=0, columnspan=6, sticky="w", pady=(3, 0))
        tk.Label(separator_frame, text="조합 구분자").pack(side="left")
        tk.Entry(separator_frame, textvariable=self.pcwbs_map_separator, width=4).pack(side="left", padx=4)
        tk.Label(
            separator_frame,
            text="헤더 예: A1 = SUBTITLE-CIA, B1 = Category5-Category4 (열 2개 이상 조합 가능)",
            fg="#555555",
        ).pack(side="left", padx=6)
        tk.Label(
            self.pcwbs_mapping_frame, textvariable=self.pcwbs_map_info, fg="#1f4e79", anchor="w"
        ).grid(row=3, column=0, columnspan=6, sticky="w", pady=(3, 0))

        key_frame = tk.LabelFrame(self.pcwbs_tab, text="비교키 생성 설정", padx=8, pady=6)
        key_frame.pack(fill="x", padx=10, pady=5)
        self.pcwbs_key_frame = key_frame
        key_frame.columnconfigure(0, weight=1)
        key_frame.columnconfigure(1, weight=1)

        pcwbs_box = tk.LabelFrame(key_frame, text="PCWBS 조합키 (위에서 아래 순서)", padx=6, pady=5)
        pcwbs_box.grid(row=0, column=0, sticky="nsew", padx=(0, 5))
        headers = ["사용", "열", "추출 방식", "숫자 자릿수"]
        for c, label in enumerate(headers):
            tk.Label(pcwbs_box, text=label, font=("Arial", 9, "bold")).grid(row=0, column=c, padx=3, sticky="w")
        extract_values = ["원문", "마지막 괄호 안", "첫 괄호 안", "숫자만", "마지막 숫자"]
        self.pcwbs_component_combos = []
        for i in range(5):
            tk.Checkbutton(pcwbs_box, variable=self.pcwbs_component_enabled[i]).grid(row=i + 1, column=0)
            combo = ttk.Combobox(pcwbs_box, textvariable=self.pcwbs_component_column[i], state="readonly", width=18)
            combo.grid(row=i + 1, column=1, padx=3, pady=1)
            self.pcwbs_component_combos.append(combo)
            ttk.Combobox(pcwbs_box, textvariable=self.pcwbs_component_extract[i], values=extract_values, state="readonly", width=15).grid(row=i + 1, column=2, padx=3)
            tk.Entry(pcwbs_box, textvariable=self.pcwbs_component_pad[i], width=6).grid(row=i + 1, column=3, padx=3)
        tk.Label(pcwbs_box, text="구분자").grid(row=6, column=0, sticky="e")
        tk.Entry(pcwbs_box, textvariable=self.pcwbs_key_separator, width=6).grid(row=6, column=1, sticky="w", padx=3)
        tk.Label(pcwbs_box, text="예: Category5(2자리) + Category4(괄호 안) → 04-4300", fg="#555").grid(row=7, column=0, columnspan=4, sticky="w", pady=(4, 0))

        bm_box = tk.LabelFrame(key_frame, text="3D BM 비교키 생성", padx=6, pady=5)
        bm_box.grid(row=0, column=1, sticky="nsew", padx=(5, 0))

        tk.Label(bm_box, text="생성 방식", font=("Arial", 9, "bold")).grid(
            row=0, column=0, sticky="w", padx=3
        )
        self.pcwbs_bm_key_mode_combo = ttk.Combobox(
            bm_box,
            textvariable=self.pcwbs_bm_key_mode,
            values=["열 조합", "MID 텍스트 추출", "두 방식 교차검증"],
            state="readonly",
            width=22,
        )
        self.pcwbs_bm_key_mode_combo.grid(row=0, column=1, columnspan=3, sticky="w", padx=3)
        self.pcwbs_bm_key_mode_combo.bind(
            "<<ComboboxSelected>>",
            lambda _event: self.update_bm_key_mode_ui(),
        )

        tk.Label(
            bm_box,
            text="A. 열 조합 (위에서 아래 순서)",
            font=("Arial", 9, "bold"),
        ).grid(row=1, column=0, columnspan=4, sticky="w", pady=(6, 2))
        for c, label in enumerate(["사용", "열", "추출 방식", "숫자 자릿수"]):
            tk.Label(bm_box, text=label, font=("Arial", 9, "bold")).grid(
                row=2, column=c, padx=3, sticky="w"
            )

        self.bm_component_combos = []
        self.bm_component_extract_combos = []
        self.bm_component_pad_entries = []
        for i in range(5):
            tk.Checkbutton(bm_box, variable=self.bm_component_enabled[i]).grid(
                row=i + 3, column=0
            )
            combo = ttk.Combobox(
                bm_box,
                textvariable=self.bm_component_column[i],
                state="readonly",
                width=18,
            )
            combo.grid(row=i + 3, column=1, padx=3, pady=1)
            self.bm_component_combos.append(combo)

            extract_combo = ttk.Combobox(
                bm_box,
                textvariable=self.bm_component_extract[i],
                values=extract_values,
                state="readonly",
                width=15,
            )
            extract_combo.grid(row=i + 3, column=2, padx=3)
            self.bm_component_extract_combos.append(extract_combo)

            pad_entry = tk.Entry(
                bm_box,
                textvariable=self.bm_component_pad[i],
                width=6,
            )
            pad_entry.grid(row=i + 3, column=3, padx=3)
            self.bm_component_pad_entries.append(pad_entry)

        tk.Label(
            bm_box,
            text="PCWBS 조합키와 동일한 구분자를 사용",
            fg="#555",
        ).grid(row=8, column=0, columnspan=4, sticky="w", pady=(3, 5))

        tk.Label(
            bm_box,
            text="B. MID 텍스트 추출",
            font=("Arial", 9, "bold"),
        ).grid(row=9, column=0, columnspan=4, sticky="w", pady=(4, 2))
        tk.Label(bm_box, text="대상 열").grid(row=10, column=0, sticky="w")
        self.pcwbs_bm_iso_combo = ttk.Combobox(
            bm_box,
            textvariable=self.pcwbs_bm_iso,
            state="readonly",
            width=18,
        )
        self.pcwbs_bm_iso_combo.grid(row=10, column=1, sticky="w", padx=3)
        tk.Label(bm_box, text="MID 시작").grid(row=10, column=2, sticky="e")
        self.pcwbs_mid_start_entry = tk.Entry(
            bm_box,
            textvariable=self.pcwbs_mid_start,
            width=6,
        )
        self.pcwbs_mid_start_entry.grid(row=10, column=3, sticky="w", padx=3)
        tk.Label(bm_box, text="길이").grid(row=11, column=2, sticky="e")
        self.pcwbs_mid_length_entry = tk.Entry(
            bm_box,
            textvariable=self.pcwbs_mid_length,
            width=6,
        )
        self.pcwbs_mid_length_entry.grid(row=11, column=3, sticky="w", padx=3)
        tk.Label(
            bm_box,
            text="Excel MID와 동일하게 시작 위치는 1부터 계산",
            fg="#555",
        ).grid(row=11, column=0, columnspan=2, sticky="w", pady=(3, 0))

        self.update_bm_key_mode_ui()
        self.update_pcwbs_compare_method_ui()

        result_frame = tk.LabelFrame(self.pcwbs_tab, text="결과 표시 설정", padx=10, pady=8)
        result_frame.pack(fill="x", padx=10, pady=5)
        tk.Label(result_frame, text="결과 열명").grid(row=0, column=0, sticky="w")
        tk.Entry(result_frame, textvariable=self.pcwbs_result_column, width=28).grid(row=0, column=1, sticky="w", padx=4)
        tk.Label(
            result_frame,
            text="각 PCWBS 행별로 현재 3D BM 키 존재 여부만 '일치/불일치'로 표시합니다.",
            fg="#555555",
        ).grid(row=0, column=2, sticky="w", padx=12)

        button_frame = tk.Frame(self.pcwbs_tab)
        button_frame.pack(fill="x", padx=10, pady=8)
        self.pcwbs_preview_button = tk.Button(button_frame, text="PCWBS 비교 미리보기", width=20, height=2, command=self.start_pcwbs_preview_thread)
        self.pcwbs_preview_button.pack(side="left")
        self.pcwbs_run_button = tk.Button(button_frame, text="PCWBS 비교 실행 및 저장", width=22, height=2, command=self.start_pcwbs_run_thread)
        self.pcwbs_run_button.pack(side="left", padx=8)
        tk.Label(button_frame, textvariable=self.pcwbs_status_text, anchor="w").pack(side="left", padx=12)

        note = (
            "v6는 이전 날짜 결과와 비교하지 않습니다. 각 PCWBS 행의 조합키가 현재 3D BM 키에 "
            "존재하는지만 '일치/불일치'로 표시하며, 결과 시트도 최소화했습니다."
        )
        tk.Label(self.pcwbs_tab, text=note, justify="left", anchor="w", fg="#444444", wraplength=1150).pack(fill="x", padx=14, pady=(0, 8))

    def _build_status_area(self):
        """Dark sidebar status/log panel with developer contact information."""
        panel_bg = "#163E63"
        panel_text = "#FFFFFF"
        panel_muted = "#D2D7DC"

        self.sidebar_status_box = tk.Frame(
            self.sidebar_frame,
            bg=panel_bg,
            padx=10,
            pady=9,
        )
        self.sidebar_status_box.grid(
            row=8,
            column=0,
            sticky="sew",
            padx=8,
            pady=(10, 8),
        )
        self.sidebar_status_box.grid_columnconfigure(0, weight=1)

        self.sidebar_status_title = tk.Label(
            self.sidebar_status_box,
            text="PROCESS STATUS",
            bg=panel_bg,
            fg=panel_muted,
            font=("Arial", 8, "bold"),
            anchor="w",
        )
        self.sidebar_status_title.grid(row=0, column=0, sticky="ew")

        self.sidebar_status_label = tk.Label(
            self.sidebar_status_box,
            textvariable=self.status_text,
            bg=panel_bg,
            fg=panel_text,
            font=("Malgun Gothic", 9, "bold"),
            anchor="w",
        )
        self.sidebar_status_label.grid(row=1, column=0, sticky="ew", pady=(4, 0))

        self.sidebar_current_label = tk.Label(
            self.sidebar_status_box,
            textvariable=self.current_item_text,
            bg=panel_bg,
            fg=panel_text,
            font=("Malgun Gothic", 8),
            anchor="w",
            justify="left",
            wraplength=175,
        )
        self.sidebar_current_label.grid(row=2, column=0, sticky="ew", pady=(2, 0))

        self.sidebar_progress_label = tk.Label(
            self.sidebar_status_box,
            textvariable=self.progress_text,
            bg=panel_bg,
            fg=panel_text,
            font=("Arial", 9, "bold"),
            anchor="w",
        )
        self.sidebar_progress_label.grid(row=3, column=0, sticky="ew", pady=(3, 3))

        self.sidebar_progress_bar = ttk.Progressbar(
            self.sidebar_status_box,
            maximum=100,
            mode="determinate",
        )
        self.sidebar_progress_bar.grid(
            row=4,
            column=0,
            sticky="ew",
            pady=(0, 8),
        )

        self.sidebar_log_title = tk.Label(
            self.sidebar_status_box,
            text="PROCESS LOG",
            bg=panel_bg,
            fg=panel_muted,
            font=("Arial", 8, "bold"),
            anchor="w",
        )
        self.sidebar_log_title.grid(row=5, column=0, sticky="ew")

        self.log_text = scrolledtext.ScrolledText(
            self.sidebar_status_box,
            height=6,
            width=22,
            wrap=tk.WORD,
            font=("Consolas", 7),
            bg="#0F2E4A",
            fg="#FFFFFF",
            insertbackground="#FFFFFF",
            relief="solid",
            bd=1,
            highlightthickness=0,
        )
        self.log_text.grid(row=6, column=0, sticky="ew", pady=(4, 0))

        self.developer_box = tk.Frame(
            self.sidebar_frame,
            bg="#163E63",
            padx=10,
            pady=10,
        )
        self.developer_box.grid(
            row=9,
            column=0,
            sticky="ew",
            padx=8,
            pady=(0, 8),
        )
        self.developer_box.grid_columnconfigure(0, weight=1)

        tk.Label(
            self.developer_box,
            text="Developed by 김영철",
            bg="#163E63",
            fg="#FFFFFF",
            font=("Malgun Gothic", 9, "bold"),
            anchor="w",
        ).grid(row=0, column=0, sticky="ew")

        tk.Label(
            self.developer_box,
            text="02-369-5379",
            bg="#163E63",
            fg="#D5E2EE",
            font=("Arial", 8),
            anchor="w",
        ).grid(row=1, column=0, sticky="ew", pady=(4, 0))

        tk.Label(
            self.developer_box,
            text="kyc4888@dlenc.co.kr",
            bg="#163E63",
            fg="#D5E2EE",
            font=("Arial", 8),
            anchor="w",
        ).grid(row=2, column=0, sticky="ew", pady=(2, 0))

    def _restore_sidebar_status_theme(self):
        """Reapply sidebar status/contact colors after global widget styling."""
        if not hasattr(self, "sidebar_status_box"):
            return

        panel_bg = "#163E63"
        panel_text = "#FFFFFF"
        panel_muted = "#D2D7DC"

        self.sidebar_status_box.configure(bg=panel_bg)
        self.sidebar_status_title.configure(bg=panel_bg, fg=panel_muted)
        self.sidebar_status_label.configure(bg=panel_bg, fg=panel_text)
        self.sidebar_current_label.configure(bg=panel_bg, fg=panel_text)
        self.sidebar_progress_label.configure(bg=panel_bg, fg=panel_text)
        self.sidebar_log_title.configure(bg=panel_bg, fg=panel_muted)

        self.log_text.configure(
            bg="#0F2E4A",
            fg="#FFFFFF",
            insertbackground="#FFFFFF",
            relief="solid",
            bd=1,
            highlightthickness=0,
        )

        if hasattr(self, "developer_box"):
            self.developer_box.configure(bg="#163E63")
            for child in self.developer_box.winfo_children():
                if isinstance(child, tk.Label):
                    child.configure(bg="#163E63")

    def _build_merge_tab(self):
        default_input = self.project_root / "02_test_input"
        default_output = self.project_root / "04_output"

        self.merge_input_folder = tk.StringVar(value=str(default_input))
        self.merge_output_folder = tk.StringVar(value=str(default_output))
        self.merge_output_filename = tk.StringVar(value="BM_Merged_Step1.xlsx")
        self.merge_exclude_folders = tk.StringVar(
            value="archive, backup, old, 01_original, 03_reference, 04_output, 05_logs, .venv"
        )
        self.merge_include_subfolders = tk.BooleanVar(value=True)
        self.merge_header_row = tk.StringVar(value="0")
        self.merge_sheet_keyword = tk.StringVar(value="")
        self.merge_preview_count = tk.StringVar(value="미리보기 파일 수: 0")
        self.merge_preview_files = []

        frame = tk.LabelFrame(self.merge_tab, text="병합 설정", padx=10, pady=10)
        frame.pack(fill="x", padx=10, pady=10)
        frame.columnconfigure(1, weight=1)

        self._path_row(frame, 0, "입력 폴더", self.merge_input_folder, self.choose_merge_input_folder)
        self._path_row(frame, 1, "결과 저장 폴더", self.merge_output_folder, self.choose_merge_output_folder)

        tk.Label(frame, text="결과 파일명").grid(row=2, column=0, sticky="w", pady=4)
        tk.Entry(frame, textvariable=self.merge_output_filename, width=40).grid(row=2, column=1, sticky="w", padx=6)

        tk.Label(frame, text="제외 폴더명").grid(row=3, column=0, sticky="w", pady=4)
        tk.Entry(frame, textvariable=self.merge_exclude_folders).grid(row=3, column=1, sticky="ew", padx=6)
        tk.Label(frame, text="쉼표 구분").grid(row=3, column=2, sticky="w")

        options = tk.Frame(frame)
        options.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(8, 0))
        tk.Checkbutton(options, text="하위 폴더 포함", variable=self.merge_include_subfolders).pack(side="left")
        tk.Label(options, text="헤더 행 (0=자동)").pack(side="left", padx=(20, 4))
        tk.Entry(options, textvariable=self.merge_header_row, width=6).pack(side="left")
        tk.Label(options, text="시트명 키워드 (빈칸=전체)").pack(side="left", padx=(20, 4))
        tk.Entry(options, textvariable=self.merge_sheet_keyword, width=24).pack(side="left")

        buttons = tk.Frame(self.merge_tab)
        buttons.pack(fill="x", padx=10, pady=(0, 5))
        self.merge_preview_button = tk.Button(buttons, text="파일 미리보기", width=16, height=2, command=self.preview_merge_files)
        self.merge_preview_button.pack(side="left")
        self.merge_run_button = tk.Button(buttons, text="병합 실행", width=16, height=2, command=self.start_merge_thread)
        self.merge_run_button.pack(side="left", padx=8)
        tk.Label(buttons, textvariable=self.merge_preview_count).pack(side="left", padx=15)

        preview_frame = tk.LabelFrame(self.merge_tab, text="병합 대상 파일", padx=8, pady=8)
        preview_frame.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.merge_preview_listbox = tk.Listbox(preview_frame, font=("Consolas", 9))
        self.merge_preview_listbox.pack(side="left", fill="both", expand=True)
        scroll = tk.Scrollbar(preview_frame, command=self.merge_preview_listbox.yview)
        scroll.pack(side="right", fill="y")
        self.merge_preview_listbox.config(yscrollcommand=scroll.set)

    def _build_cleanup_tab(self):
        default_output = self.project_root / "04_output"
        self.cleanup_input_file = tk.StringVar(value=str(default_output / "BM_Merged_Step1.xlsx"))
        self.cleanup_output_file = tk.StringVar(value=str(default_output / "BM_Modified_Result.xlsx"))
        self.cleanup_sheet = tk.StringVar(value="Merged_Data")
        self.cleanup_header_row = tk.StringVar(value="1")

        source_frame = tk.LabelFrame(self.cleanup_tab, text="수정 대상", padx=10, pady=10)
        source_frame.pack(fill="x", padx=10, pady=(10, 5))
        source_frame.columnconfigure(1, weight=1)

        self._file_row(source_frame, 0, "병합 파일", self.cleanup_input_file, self.choose_cleanup_input_file)
        self._file_row(source_frame, 1, "결과 파일", self.cleanup_output_file, self.choose_cleanup_output_file, save=True)

        tk.Label(source_frame, text="데이터 시트").grid(row=2, column=0, sticky="w", pady=4)
        self.cleanup_sheet_combo = ttk.Combobox(source_frame, textvariable=self.cleanup_sheet, state="readonly", width=35)
        self.cleanup_sheet_combo.grid(row=2, column=1, sticky="w", padx=6)
        self.cleanup_sheet_combo.bind("<<ComboboxSelected>>", lambda _event: self.load_cleanup_headers())

        tk.Label(source_frame, text="헤더 행").grid(row=2, column=2, sticky="e", padx=(10, 4))
        tk.Entry(source_frame, textvariable=self.cleanup_header_row, width=6).grid(row=2, column=3, sticky="w")
        tk.Button(source_frame, text="열 불러오기", command=self.load_cleanup_headers, width=12).grid(row=2, column=4, padx=(8, 0))

        rule_frame = tk.LabelFrame(self.cleanup_tab, text="수정 규칙 작성", padx=10, pady=8)
        rule_frame.pack(fill="x", padx=10, pady=5)

        self.rule_action = tk.StringVar(value=ACTION_DELETE)
        self.rule_condition_column = tk.StringVar()
        self.rule_operator = tk.StringVar(value="정확히 일치")
        self.rule_condition_value = tk.StringVar()
        self.rule_second_enabled = tk.BooleanVar(value=False)
        self.rule_second_logic = tk.StringVar(value="AND")
        self.rule_second_column = tk.StringVar()
        self.rule_second_operator = tk.StringVar(value="정확히 일치")
        self.rule_second_value = tk.StringVar()
        self.rule_target_column = tk.StringVar()
        self.rule_new_value = tk.StringVar()
        self.rule_case_sensitive = tk.BooleanVar(value=False)
        self.rule_exception_enabled = tk.BooleanVar(value=False)
        self.rule_exception_column = tk.StringVar()
        self.rule_exception_operator = tk.StringVar(value="정확히 일치")
        self.rule_exception_value = tk.StringVar()

        tk.Label(rule_frame, text="동작").grid(row=0, column=0, sticky="w")
        action_combo = ttk.Combobox(
            rule_frame,
            textvariable=self.rule_action,
            values=[ACTION_DELETE, ACTION_REPLACE, ACTION_REVIEW, ACTION_SORT_ASC, ACTION_SORT_DESC, ACTION_REMOVE_EMPTY],
            state="readonly",
            width=18,
        )
        action_combo.grid(row=0, column=1, sticky="w", padx=4)
        action_combo.bind("<<ComboboxSelected>>", lambda _event: self.update_rule_input_state())

        tk.Label(rule_frame, text="조건 열").grid(row=0, column=2, sticky="e")
        self.condition_column_combo = ttk.Combobox(rule_frame, textvariable=self.rule_condition_column, state="readonly", width=24)
        self.condition_column_combo.grid(row=0, column=3, sticky="w", padx=4)

        tk.Label(rule_frame, text="조건").grid(row=0, column=4, sticky="e")
        self.operator_combo = ttk.Combobox(rule_frame, textvariable=self.rule_operator, values=OPERATORS, state="readonly", width=17)
        self.operator_combo.grid(row=0, column=5, sticky="w", padx=4)

        tk.Label(rule_frame, text="조건 값").grid(row=0, column=6, sticky="e")
        self.condition_value_entry = tk.Entry(rule_frame, textvariable=self.rule_condition_value, width=28)
        self.condition_value_entry.grid(row=0, column=7, sticky="ew", padx=4)

        tk.Label(rule_frame, text="수정 열").grid(row=1, column=0, sticky="w", pady=(8, 0))
        self.target_column_combo = ttk.Combobox(rule_frame, textvariable=self.rule_target_column, state="readonly", width=24)
        self.target_column_combo.grid(row=1, column=1, columnspan=2, sticky="w", padx=4, pady=(8, 0))

        tk.Label(rule_frame, text="새 값").grid(row=1, column=3, sticky="e", pady=(8, 0))
        self.new_value_entry = tk.Entry(rule_frame, textvariable=self.rule_new_value, width=28)
        self.new_value_entry.grid(row=1, column=4, columnspan=2, sticky="w", padx=4, pady=(8, 0))

        tk.Checkbutton(rule_frame, text="대소문자 구분", variable=self.rule_case_sensitive).grid(row=1, column=6, sticky="w", pady=(8, 0))

        second_frame = tk.Frame(rule_frame)
        second_frame.grid(row=2, column=0, columnspan=8, sticky="ew", pady=(8, 0))
        tk.Checkbutton(second_frame, text="두 번째 조건 사용", variable=self.rule_second_enabled).pack(side="left")
        ttk.Combobox(second_frame, textvariable=self.rule_second_logic, values=["AND", "OR"], state="readonly", width=6).pack(side="left", padx=(8, 4))
        tk.Label(second_frame, text="조건 열").pack(side="left", padx=(6, 4))
        self.second_column_combo = ttk.Combobox(second_frame, textvariable=self.rule_second_column, state="readonly", width=22)
        self.second_column_combo.pack(side="left")
        tk.Label(second_frame, text="조건").pack(side="left", padx=(8, 4))
        self.second_operator_combo = ttk.Combobox(second_frame, textvariable=self.rule_second_operator, values=OPERATORS, state="readonly", width=16)
        self.second_operator_combo.pack(side="left")
        tk.Label(second_frame, text="값").pack(side="left", padx=(8, 4))
        self.second_value_entry = tk.Entry(second_frame, textvariable=self.rule_second_value, width=22)
        self.second_value_entry.pack(side="left")

        exception_frame = tk.Frame(rule_frame)
        exception_frame.grid(row=3, column=0, columnspan=8, sticky="ew", pady=(8, 0))
        tk.Checkbutton(
            exception_frame,
            text="예외 조건 사용 (조건에 해당해도 남김/수정 안 함)",
            variable=self.rule_exception_enabled,
        ).pack(side="left")
        tk.Label(exception_frame, text="예외 열").pack(side="left", padx=(12, 4))
        self.exception_column_combo = ttk.Combobox(exception_frame, textvariable=self.rule_exception_column, state="readonly", width=22)
        self.exception_column_combo.pack(side="left")
        tk.Label(exception_frame, text="예외 조건").pack(side="left", padx=(10, 4))
        ttk.Combobox(exception_frame, textvariable=self.rule_exception_operator, values=OPERATORS, state="readonly", width=16).pack(side="left")
        tk.Label(exception_frame, text="예외 값").pack(side="left", padx=(10, 4))
        tk.Entry(exception_frame, textvariable=self.rule_exception_value, width=24).pack(side="left")

        rule_frame.columnconfigure(7, weight=1)

        rule_buttons = tk.Frame(self.cleanup_tab)
        rule_buttons.pack(fill="x", padx=10, pady=5)
        tk.Button(rule_buttons, text="규칙 추가", width=12, command=self.add_rule).pack(side="left")
        tk.Button(rule_buttons, text="선택 수정", width=12, command=self.update_selected_rule).pack(side="left", padx=4)
        tk.Button(rule_buttons, text="선택 삭제", width=12, command=self.delete_selected_rule).pack(side="left", padx=4)
        tk.Button(rule_buttons, text="위로", width=8, command=lambda: self.move_rule(-1)).pack(side="left", padx=(12, 2))
        tk.Button(rule_buttons, text="아래로", width=8, command=lambda: self.move_rule(1)).pack(side="left", padx=2)
        tk.Button(rule_buttons, text="규칙 저장", width=12, command=self.save_rules).pack(side="right", padx=4)
        tk.Button(rule_buttons, text="규칙 불러오기", width=12, command=self.load_rules).pack(side="right", padx=4)

        tree_frame = tk.LabelFrame(self.cleanup_tab, text="적용 규칙 목록 (위에서 아래 순서로 실행)", padx=8, pady=8)
        tree_frame.pack(fill="both", expand=True, padx=10, pady=5)
        columns = ("no", "action", "condition", "target", "exception")
        self.rule_tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=10)
        self.rule_tree.heading("no", text="#")
        self.rule_tree.heading("action", text="동작")
        self.rule_tree.heading("condition", text="조건")
        self.rule_tree.heading("target", text="수정 내용")
        self.rule_tree.heading("exception", text="예외")
        self.rule_tree.column("no", width=45, anchor="center")
        self.rule_tree.column("action", width=120)
        self.rule_tree.column("condition", width=430)
        self.rule_tree.column("target", width=260)
        self.rule_tree.column("exception", width=330)
        self.rule_tree.pack(side="left", fill="both", expand=True)
        tree_scroll = tk.Scrollbar(tree_frame, command=self.rule_tree.yview)
        tree_scroll.pack(side="right", fill="y")
        self.rule_tree.config(yscrollcommand=tree_scroll.set)
        self.rule_tree.bind("<<TreeviewSelect>>", self.populate_rule_form_from_selection)

        execution_frame = tk.Frame(self.cleanup_tab)
        execution_frame.pack(fill="x", padx=10, pady=(5, 10))
        self.preview_rules_button = tk.Button(execution_frame, text="적용 건수 미리보기", width=18, height=2, command=self.start_preview_rules_thread)
        self.preview_rules_button.pack(side="left")
        self.execute_rules_button = tk.Button(execution_frame, text="수정 실행 및 저장", width=18, height=2, command=self.start_cleanup_thread)
        self.execute_rules_button.pack(side="left", padx=8)
        tk.Button(execution_frame, text="전체 규칙 삭제", width=14, command=self.clear_all_rules).pack(side="left", padx=8)

        self.update_rule_input_state()

    def _path_row(self, parent, row, label, variable, command):
        tk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=4)
        tk.Entry(parent, textvariable=variable).grid(row=row, column=1, sticky="ew", padx=6)
        tk.Button(parent, text="폴더 선택", command=command, width=12).grid(row=row, column=2)

    def _file_row(self, parent, row, label, variable, command, save=False):
        tk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=4)
        tk.Entry(parent, textvariable=variable).grid(row=row, column=1, columnspan=3, sticky="ew", padx=6)
        tk.Button(parent, text="저장 위치" if save else "파일 선택", command=command, width=12).grid(row=row, column=4)

    # ------------------------------------------------------------- 공통 상태
    def log(self, message):
        self.root.after(0, self._append_log, message)

    def _append_log(self, message):
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert(tk.END, f"[{timestamp}] {message}\n")
        self.log_text.see(tk.END)

    def set_status(self, status, current_item="현재 작업: 없음", percent=None):
        self.root.after(0, self.status_text.set, status)
        self.root.after(0, self.current_item_text.set, current_item)
        if percent is not None:
            self.root.after(0, self._set_progress, percent)

    def _set_progress(self, percent):
        percent = max(0, min(100, percent))
        self.progress_text.set(f"진행률: {percent:.1f}%")
        if hasattr(self, "sidebar_progress_bar"):
            self.sidebar_progress_bar["value"] = percent

    def set_running(self, running):
        self.is_running = running
        state = "disabled" if running else "normal"
        for button in [
            self.merge_preview_button,
            self.merge_run_button,
            self.preview_rules_button,
            self.execute_rules_button,
            self.pcwbs_preview_button,
            self.pcwbs_run_button,
            self.display_preview_button,
            self.display_generate_button,
            self.paint_run_button,
            self.paint_save_rules_button,
            self.paint_load_rules_button,
            self.split_preview_button,
            self.split_run_button,
            self.display_apply_mapping_button,
            self.display_auto_map_button,
        ]:
            self.root.after(0, button.config, {"state": state})

    # ----------------------------------------------------------------- 병합
    def choose_merge_input_folder(self):
        folder = filedialog.askdirectory(initialdir=self.merge_input_folder.get())
        if folder:
            self.merge_input_folder.set(folder)

    def choose_merge_output_folder(self):
        folder = filedialog.askdirectory(initialdir=self.merge_output_folder.get())
        if folder:
            self.merge_output_folder.set(folder)

    def parse_excluded_folder_names(self):
        return {
            clean_text(name).casefold()
            for name in self.merge_exclude_folders.get().split(",")
            if clean_text(name)
        }

    def collect_merge_files(self):
        input_path = Path(self.merge_input_folder.get().strip())
        output_path = Path(self.merge_output_folder.get().strip())
        output_filename = self.merge_output_filename.get().strip()
        excluded_names = self.parse_excluded_folder_names()

        candidates = list(input_path.rglob("*")) if self.merge_include_subfolders.get() else list(input_path.glob("*"))
        result = []
        for file_path in candidates:
            if not file_path.is_file() or file_path.suffix.casefold() not in SUPPORTED_EXTENSIONS:
                continue
            if file_path.name.startswith("~$"):
                continue
            try:
                folder_parts = file_path.relative_to(input_path).parts[:-1]
            except ValueError:
                folder_parts = file_path.parts[:-1]
            if any(part.casefold() in excluded_names for part in folder_parts):
                continue
            try:
                if file_path.resolve() == (output_path / output_filename).resolve():
                    continue
            except OSError:
                pass
            result.append(file_path)
        return sorted(result, key=lambda path: str(path).casefold())

    def preview_merge_files(self):
        try:
            files = self.collect_merge_files()
            self.merge_preview_files = files
            self.merge_preview_listbox.delete(0, tk.END)
            base = Path(self.merge_input_folder.get().strip())
            for file_path in files:
                self.merge_preview_listbox.insert(tk.END, safe_relative_path(file_path, base))
            self.merge_preview_count.set(f"미리보기 파일 수: {len(files)}")
            self.log(f"병합 대상 파일 미리보기: {len(files)}개")
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def start_merge_thread(self):
        if self.is_running:
            return
        threading.Thread(target=self.run_merge, daemon=True).start()

    def run_merge(self):
        self.set_running(True)
        try:
            input_path = Path(self.merge_input_folder.get().strip())
            output_folder = Path(self.merge_output_folder.get().strip())
            output_name = self.merge_output_filename.get().strip()
            if not input_path.is_dir():
                raise ValueError("유효한 입력 폴더를 선택하세요.")
            output_folder.mkdir(parents=True, exist_ok=True)
            if not output_name.lower().endswith(".xlsx"):
                raise ValueError("결과 파일명은 .xlsx로 끝나야 합니다.")
            try:
                manual_header = int(self.merge_header_row.get().strip())
            except ValueError as exc:
                raise ValueError("헤더 행은 0 이상의 정수여야 합니다.") from exc
            sheet_keyword = clean_text(self.merge_sheet_keyword.get())

            files = self.collect_merge_files()
            if not files:
                raise ValueError("병합할 Excel 파일이 없습니다.")

            self.log("=== 병합 시작 ===")
            self.set_status("병합 중", "현재 작업: 파일 구조 분석", 0)

            all_headers = list(SOURCE_COLUMNS)
            seen_headers = {canonical_header(header) for header in all_headers}
            data_blocks = []
            summary_rows = []
            validation_rows = []

            total_files = len(files)
            for file_index, file_path in enumerate(files, start=1):
                self.set_status("병합 중", f"현재 작업: {file_path.name}", (file_index - 1) / total_files * 85)
                self.log(f"파일 처리: {file_path.name}")
                wb = load_workbook(file_path, data_only=False, read_only=False)
                relative_path = safe_relative_path(file_path, input_path)

                for ws in wb.worksheets:
                    if sheet_keyword and sheet_keyword.casefold() not in ws.title.casefold():
                        validation_rows.append({
                            "Level": "INFO", "File": file_path.name, "Sheet": ws.title,
                            "Issue": "시트명 키워드 불일치", "Detail": f"'{sheet_keyword}' 없음",
                        })
                        continue

                    header_row = manual_header if manual_header > 0 else detect_header_row(ws)[0]
                    if not header_row:
                        validation_rows.append({
                            "Level": "WARNING", "File": file_path.name, "Sheet": ws.title,
                            "Issue": "헤더 탐지 실패", "Detail": "처리하지 않음",
                        })
                        continue

                    last_col = find_last_nonempty_column(ws, header_row)
                    if last_col == 0:
                        continue
                    raw_headers = [ws.cell(header_row, col).value for col in range(1, last_col + 1)]
                    headers = make_unique_headers(raw_headers)

                    for header in headers:
                        key = canonical_header(header)
                        if key not in seen_headers:
                            all_headers.append(header)
                            seen_headers.add(key)

                    rows = []
                    for row_number in range(header_row + 1, ws.max_row + 1):
                        values = [ws.cell(row_number, col).value for col in range(1, last_col + 1)]
                        if is_effectively_empty_row(values):
                            continue
                        row_dict = dict(zip(headers, values))
                        row_dict.update({
                            "Source File": file_path.name,
                            "Source Relative Path": relative_path,
                            "Source Sheet": ws.title,
                            "Source Area": derive_area_from_filename(file_path),
                            "Source Header Row": header_row,
                        })
                        rows.append(row_dict)

                    if rows:
                        data_blocks.extend(rows)
                    summary_rows.append({
                        "File": file_path.name,
                        "Relative Path": relative_path,
                        "Sheet": ws.title,
                        "Header Row": header_row,
                        "Columns": len(headers),
                        "Data Rows": len(rows),
                    })
                wb.close()

            self.set_status("병합 중", "현재 작업: 결과 파일 작성", 90)
            out_wb = Workbook()
            data_ws = out_wb.active
            data_ws.title = "Merged_Data"
            write_rows_to_sheet(data_ws, all_headers, data_blocks)

            summary_ws = out_wb.create_sheet("Summary")
            summary_headers = ["File", "Relative Path", "Sheet", "Header Row", "Columns", "Data Rows"]
            write_rows_to_sheet(summary_ws, summary_headers, summary_rows)

            validation_ws = out_wb.create_sheet("Validation")
            validation_headers = ["Level", "File", "Sheet", "Issue", "Detail"]
            write_rows_to_sheet(validation_ws, validation_headers, validation_rows)

            info_ws = out_wb.create_sheet("Run_Info")
            info_ws.append(["Item", "Value"])
            info_ws.append(["Run Time", datetime.now().isoformat(timespec="seconds")])
            info_ws.append(["Input Folder", str(input_path)])
            info_ws.append(["Files", len(files)])
            info_ws.append(["Merged Rows", len(data_blocks)])
            style_header(info_ws)
            auto_fit_columns(info_ws)

            output_path = output_folder / output_name
            out_wb.save(output_path)
            self.cleanup_input_file.set(str(output_path))
            self.load_cleanup_file_metadata(show_message=False)
            self.set_status("완료", "현재 작업: 없음", 100)
            self.log(f"병합 완료: {output_path}")
            self.root.after(0, lambda: messagebox.showinfo("완료", f"병합이 완료되었습니다.\n\n{output_path}"))

        except Exception as exc:
            self.log(traceback.format_exc())
            self.set_status("오류", "현재 작업: 중단", 0)
            self.root.after(0, lambda exc=exc: messagebox.showerror("오류", str(exc)))
        finally:
            self.set_running(False)

    # --------------------------------------------------------------- 정제 UI
    def choose_cleanup_input_file(self):
        file_path = filedialog.askopenfilename(
            initialdir=str(Path(self.cleanup_input_file.get()).parent),
            filetypes=[("Excel files", "*.xlsx *.xlsm")],
        )
        if file_path:
            self.cleanup_input_file.set(file_path)
            self.load_cleanup_file_metadata()

    def choose_cleanup_output_file(self):
        file_path = filedialog.asksaveasfilename(
            initialdir=str(Path(self.cleanup_output_file.get()).parent),
            initialfile=Path(self.cleanup_output_file.get()).name,
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if file_path:
            self.cleanup_output_file.set(file_path)

    def load_cleanup_file_metadata(self, show_message=True):
        path = Path(self.cleanup_input_file.get().strip())
        if not path.exists():
            if show_message:
                messagebox.showwarning("안내", "정제 대상 파일이 없습니다.")
            return
        try:
            wb = load_workbook(path, read_only=True, data_only=False)
            self.cleanup_sheet_names = list(wb.sheetnames)
            wb.close()
            self.cleanup_sheet_combo["values"] = self.cleanup_sheet_names
            if "Merged_Data" in self.cleanup_sheet_names:
                self.cleanup_sheet.set("Merged_Data")
            elif self.cleanup_sheet_names:
                self.cleanup_sheet.set(self.cleanup_sheet_names[0])
            self.load_cleanup_headers(show_message=False)
            if show_message:
                self.log(f"정제 대상 파일 불러옴: {path.name}")
        except Exception as exc:
            messagebox.showerror("오류", f"파일 정보를 읽지 못했습니다.\n{exc}")

    def load_cleanup_headers(self, show_message=True):
        path = Path(self.cleanup_input_file.get().strip())
        if not path.exists():
            if show_message:
                messagebox.showwarning("안내", "정제 대상 파일을 선택하세요.")
            return
        try:
            header_row = int(self.cleanup_header_row.get().strip())
            if header_row < 1:
                raise ValueError
        except ValueError:
            messagebox.showerror("오류", "헤더 행은 1 이상의 정수여야 합니다.")
            return

        sheet_name = self.cleanup_sheet.get().strip()
        try:
            wb = load_workbook(path, read_only=True, data_only=False)
            if sheet_name not in wb.sheetnames:
                wb.close()
                raise ValueError("선택한 시트가 없습니다.")
            ws = wb[sheet_name]
            last_col = find_last_nonempty_column(ws, header_row)
            raw_headers = [ws.cell(header_row, col).value for col in range(1, last_col + 1)]
            self.cleanup_headers = make_unique_headers(raw_headers)
            wb.close()

            for combo in [
                self.condition_column_combo,
                self.target_column_combo,
                self.exception_column_combo,
                self.second_column_combo,
            ]:
                combo["values"] = self.cleanup_headers
            if self.cleanup_headers:
                if self.rule_condition_column.get() not in self.cleanup_headers:
                    self.rule_condition_column.set(self.cleanup_headers[0])
                if self.rule_target_column.get() not in self.cleanup_headers:
                    self.rule_target_column.set(self.cleanup_headers[0])
                if self.rule_exception_column.get() not in self.cleanup_headers:
                    self.rule_exception_column.set(self.cleanup_headers[0])
                if self.rule_second_column.get() not in self.cleanup_headers:
                    self.rule_second_column.set(self.cleanup_headers[0])
            if show_message:
                self.log(f"열 불러오기 완료: {len(self.cleanup_headers)}개")
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def update_rule_input_state(self):
        action = self.rule_action.get()
        condition_required = action in [ACTION_DELETE, ACTION_REPLACE, ACTION_REVIEW]
        replace_required = action == ACTION_REPLACE
        sort_required = action in [ACTION_SORT_ASC, ACTION_SORT_DESC]

        self.condition_column_combo.config(state="readonly" if (condition_required or sort_required) else "disabled")
        self.operator_combo.config(state="readonly" if condition_required else "disabled")
        self.condition_value_entry.config(state="normal" if condition_required else "disabled")
        self.target_column_combo.config(state="readonly" if replace_required else "disabled")
        self.new_value_entry.config(state="normal" if replace_required else "disabled")

    def build_rule_from_form(self):
        action = self.rule_action.get()
        rule = {
            "action": action,
            "condition_column": self.rule_condition_column.get().strip(),
            "operator": self.rule_operator.get().strip(),
            "condition_value": self.rule_condition_value.get(),
            "second_enabled": bool(self.rule_second_enabled.get()),
            "second_logic": self.rule_second_logic.get().strip() or "AND",
            "second_column": self.rule_second_column.get().strip(),
            "second_operator": self.rule_second_operator.get().strip(),
            "second_value": self.rule_second_value.get(),
            "target_column": self.rule_target_column.get().strip(),
            "new_value": self.rule_new_value.get(),
            "case_sensitive": bool(self.rule_case_sensitive.get()),
            "exception_enabled": bool(self.rule_exception_enabled.get()),
            "exception_column": self.rule_exception_column.get().strip(),
            "exception_operator": self.rule_exception_operator.get().strip(),
            "exception_value": self.rule_exception_value.get(),
        }

        if action in [ACTION_DELETE, ACTION_REPLACE, ACTION_REVIEW] and not rule["condition_column"]:
            raise ValueError("조건 열을 선택하세요.")
        if action in [ACTION_SORT_ASC, ACTION_SORT_DESC] and not rule["condition_column"]:
            raise ValueError("정렬 기준 열을 선택하세요.")
        if action == ACTION_REPLACE and not rule["target_column"]:
            raise ValueError("수정할 열을 선택하세요.")
        if rule["second_enabled"] and not rule["second_column"]:
            raise ValueError("두 번째 조건 열을 선택하세요.")
        if rule["exception_enabled"] and not rule["exception_column"]:
            raise ValueError("예외 조건 열을 선택하세요.")
        return rule

    def add_rule(self):
        try:
            rule = self.build_rule_from_form()
            self.rules.append(rule)
            self.refresh_rule_tree()
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def update_selected_rule(self):
        selected = self.rule_tree.selection()
        if not selected:
            messagebox.showwarning("안내", "수정할 규칙을 선택하세요.")
            return
        try:
            index = int(selected[0])
            self.rules[index] = self.build_rule_from_form()
            self.refresh_rule_tree()
            self.rule_tree.selection_set(str(index))
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def delete_selected_rule(self):
        selected = self.rule_tree.selection()
        if not selected:
            return
        index = int(selected[0])
        del self.rules[index]
        self.refresh_rule_tree()

    def move_rule(self, direction):
        selected = self.rule_tree.selection()
        if not selected:
            return
        index = int(selected[0])
        new_index = index + direction
        if new_index < 0 or new_index >= len(self.rules):
            return
        self.rules[index], self.rules[new_index] = self.rules[new_index], self.rules[index]
        self.refresh_rule_tree()
        self.rule_tree.selection_set(str(new_index))

    def clear_all_rules(self):
        if self.rules and not messagebox.askyesno("확인", "모든 규칙을 삭제할까요?"):
            return
        self.rules.clear()
        self.refresh_rule_tree()

    def refresh_rule_tree(self):
        for item in self.rule_tree.get_children():
            self.rule_tree.delete(item)
        for index, rule in enumerate(self.rules):
            condition = ""
            target = ""
            exception = ""
            action = rule["action"]
            if action in [ACTION_DELETE, ACTION_REPLACE, ACTION_REVIEW]:
                condition = f"{rule['condition_column']} | {rule['operator']} | {rule['condition_value']}"
                if rule.get("second_enabled"):
                    condition += (
                        f" {rule.get('second_logic', 'AND')} "
                        f"{rule.get('second_column', '')} | {rule.get('second_operator', '')} | {rule.get('second_value', '')}"
                    )
            elif action in [ACTION_SORT_ASC, ACTION_SORT_DESC]:
                condition = f"기준 열: {rule['condition_column']}"
            else:
                condition = "행 전체가 빈 경우"

            if action == ACTION_REPLACE:
                target = f"{rule['target_column']} → {rule['new_value']}"
            elif action == ACTION_REVIEW:
                target = "Review_Items 시트에 복사 (Final_Data에는 유지)"
            if rule.get("exception_enabled"):
                exception = f"{rule['exception_column']} | {rule['exception_operator']} | {rule['exception_value']}"

            self.rule_tree.insert("", "end", iid=str(index), values=(index + 1, action, condition, target, exception))

    def populate_rule_form_from_selection(self, _event=None):
        selected = self.rule_tree.selection()
        if not selected:
            return
        rule = self.rules[int(selected[0])]
        self.rule_action.set(rule.get("action", ACTION_DELETE))
        self.rule_condition_column.set(rule.get("condition_column", ""))
        self.rule_operator.set(rule.get("operator", "정확히 일치"))
        self.rule_condition_value.set(rule.get("condition_value", ""))
        self.rule_second_enabled.set(rule.get("second_enabled", False))
        self.rule_second_logic.set(rule.get("second_logic", "AND"))
        self.rule_second_column.set(rule.get("second_column", ""))
        self.rule_second_operator.set(rule.get("second_operator", "정확히 일치"))
        self.rule_second_value.set(rule.get("second_value", ""))
        self.rule_target_column.set(rule.get("target_column", ""))
        self.rule_new_value.set(rule.get("new_value", ""))
        self.rule_case_sensitive.set(rule.get("case_sensitive", False))
        self.rule_exception_enabled.set(rule.get("exception_enabled", False))
        self.rule_exception_column.set(rule.get("exception_column", ""))
        self.rule_exception_operator.set(rule.get("exception_operator", "정확히 일치"))
        self.rule_exception_value.set(rule.get("exception_value", ""))
        self.update_rule_input_state()

    def save_rules(self):
        file_path = filedialog.asksaveasfilename(
            initialdir=str(self.project_root / "03_reference"),
            initialfile="cleanup_rules.json",
            defaultextension=".json",
            filetypes=[("JSON files", "*.json")],
        )
        if not file_path:
            return
        payload = {
            "version": 2,
            "saved_at": datetime.now().isoformat(timespec="seconds"),
            "rules": self.rules,
        }
        Path(file_path).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        self.log(f"규칙 저장: {file_path}")

    def load_rules(self):
        file_path = filedialog.askopenfilename(
            initialdir=str(self.project_root / "03_reference"),
            filetypes=[("JSON files", "*.json")],
        )
        if not file_path:
            return
        payload = json.loads(Path(file_path).read_text(encoding="utf-8"))
        rules = payload.get("rules", payload if isinstance(payload, list) else [])
        if not isinstance(rules, list):
            raise ValueError("올바른 규칙 파일이 아닙니다.")
        self.rules = rules
        self.refresh_rule_tree()
        self.log(f"규칙 불러오기: {file_path}")

    # -------------------------------------------------------------- 정제 엔진
    def read_cleanup_rows(self):
        path = Path(self.cleanup_input_file.get().strip())
        if not path.exists():
            raise ValueError("정제 대상 파일을 선택하세요.")
        try:
            header_row = int(self.cleanup_header_row.get().strip())
        except ValueError as exc:
            raise ValueError("헤더 행은 1 이상의 정수여야 합니다.") from exc
        sheet_name = self.cleanup_sheet.get().strip()

        wb = load_workbook(path, data_only=False, read_only=False)
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise ValueError("선택한 데이터 시트가 없습니다.")
        ws = wb[sheet_name]
        last_col = find_last_nonempty_column(ws, header_row)
        headers = make_unique_headers([ws.cell(header_row, col).value for col in range(1, last_col + 1)])
        rows = []
        for row_number in range(header_row + 1, ws.max_row + 1):
            values = [ws.cell(row_number, col).value for col in range(1, last_col + 1)]
            row_dict = dict(zip(headers, values))
            row_dict["__source_row__"] = row_number
            rows.append(row_dict)
        wb.close()
        return headers, rows

    def row_matches_rule(self, row, rule):
        if rule["action"] == ACTION_REMOVE_EMPTY:
            return is_effectively_empty_row([row.get(header) for header in self.cleanup_headers])

        first = compare_value(
            row.get(rule.get("condition_column", "")),
            rule.get("operator", "정확히 일치"),
            rule.get("condition_value", ""),
            rule.get("case_sensitive", False),
        )
        if not rule.get("second_enabled"):
            return first

        second = compare_value(
            row.get(rule.get("second_column", "")),
            rule.get("second_operator", "정확히 일치"),
            rule.get("second_value", ""),
            rule.get("case_sensitive", False),
        )
        if rule.get("second_logic", "AND") == "OR":
            return first or second
        return first and second

    def row_is_exception(self, row, rule):
        if not rule.get("exception_enabled"):
            return False
        return compare_value(
            row.get(rule.get("exception_column", "")),
            rule.get("exception_operator", "정확히 일치"),
            rule.get("exception_value", ""),
            rule.get("case_sensitive", False),
        )

    def apply_rules(self, headers, original_rows, preview_only=False):
        rows = [dict(row) for row in original_rows]
        deleted_rows = []
        modified_rows = []
        review_rows = []
        rule_summary = []
        review_seen = set()

        for rule_index, rule in enumerate(self.rules, start=1):
            action = rule["action"]
            affected = 0

            if action in [ACTION_SORT_ASC, ACTION_SORT_DESC]:
                column = rule["condition_column"]
                reverse = action == ACTION_SORT_DESC

                def sort_key(row):
                    value = row.get(column)
                    numeric = normalize_numeric(value)
                    if numeric is not None:
                        return (0, numeric)
                    text = clean_text(value)
                    if text:
                        return (1, text.casefold())
                    return (2, "")

                rows.sort(key=sort_key, reverse=reverse)
                affected = len(rows)

            elif action in [ACTION_DELETE, ACTION_REMOVE_EMPTY]:
                remaining = []
                for row in rows:
                    matches = self.row_matches_rule(row, rule)
                    if matches and not self.row_is_exception(row, rule):
                        affected += 1
                        if not preview_only:
                            deleted = {header: row.get(header) for header in headers}
                            deleted["Rule No"] = rule_index
                            deleted["Action"] = action
                            deleted["Reason"] = self.rule_description(rule)
                            deleted["Original Excel Row"] = row.get("__source_row__")
                            deleted_rows.append(deleted)
                    else:
                        remaining.append(row)
                rows = remaining

            elif action == ACTION_REVIEW:
                for row in rows:
                    if self.row_matches_rule(row, rule) and not self.row_is_exception(row, rule):
                        affected += 1
                        if not preview_only:
                            unique_key = (rule_index, row.get("__source_row__"))
                            if unique_key not in review_seen:
                                review_seen.add(unique_key)
                                reviewed = {header: row.get(header) for header in headers}
                                reviewed["Rule No"] = rule_index
                                reviewed["Reason"] = self.rule_description(rule)
                                reviewed["Original Excel Row"] = row.get("__source_row__")
                                review_rows.append(reviewed)

            elif action == ACTION_REPLACE:
                for row in rows:
                    if self.row_matches_rule(row, rule) and not self.row_is_exception(row, rule):
                        target_column = rule["target_column"]
                        old_value = row.get(target_column)
                        new_value = rule.get("new_value", "")
                        if old_value != new_value:
                            affected += 1
                            if not preview_only:
                                row[target_column] = new_value
                                modified_rows.append({
                                    "Rule No": rule_index,
                                    "Original Excel Row": row.get("__source_row__"),
                                    "Target Column": target_column,
                                    "Old Value": old_value,
                                    "New Value": new_value,
                                    "Reason": self.rule_description(rule),
                                })

            rule_summary.append({
                "Rule No": rule_index,
                "Action": action,
                "Description": self.rule_description(rule),
                "Affected Rows": affected,
                "Rows Remaining": len(rows),
            })

        return rows, deleted_rows, modified_rows, review_rows, rule_summary

    def rule_description(self, rule):
        action = rule["action"]
        if action == ACTION_REMOVE_EMPTY:
            return "모든 데이터 열이 빈 행 삭제"
        if action in [ACTION_SORT_ASC, ACTION_SORT_DESC]:
            return f"{rule['condition_column']} 기준 {action}"
        text = f"{rule['condition_column']} {rule['operator']} '{rule['condition_value']}'"
        if rule.get("second_enabled"):
            text += (
                f" {rule.get('second_logic', 'AND')} "
                f"{rule.get('second_column', '')} {rule.get('second_operator', '')} '{rule.get('second_value', '')}'"
            )
        if action == ACTION_REPLACE:
            text += f" → {rule['target_column']}='{rule['new_value']}'"
        if rule.get("exception_enabled"):
            text += (
                f" / 예외: {rule['exception_column']} "
                f"{rule['exception_operator']} '{rule['exception_value']}'"
            )
        return text

    def start_preview_rules_thread(self):
        if self.is_running:
            return
        threading.Thread(target=self.preview_rules, daemon=True).start()

    def preview_rules(self):
        self.set_running(True)
        try:
            if not self.rules:
                raise ValueError("등록된 규칙이 없습니다.")
            self.set_status("미리보기 중", "현재 작업: 데이터 읽기", 10)
            headers, rows = self.read_cleanup_rows()
            self.cleanup_headers = headers
            self.set_status("미리보기 중", "현재 작업: 규칙 시뮬레이션", 50)
            _, _, _, _, summary = self.apply_rules(headers, rows, preview_only=True)
            lines = [f"원본 데이터 행: {len(rows):,}개", ""]
            for item in summary:
                lines.append(
                    f"규칙 {item['Rule No']}: {item['Affected Rows']:,}행 영향 / "
                    f"적용 후 {item['Rows Remaining']:,}행"
                )
            self.set_status("미리보기 완료", "현재 작업: 없음", 100)
            self.root.after(0, lambda: messagebox.showinfo("적용 건수 미리보기", "\n".join(lines)))
        except Exception as exc:
            self.log(traceback.format_exc())
            self.root.after(0, lambda exc=exc: messagebox.showerror("오류", str(exc)))
        finally:
            self.set_running(False)

    def start_cleanup_thread(self):
        if self.is_running:
            return
        threading.Thread(target=self.run_cleanup, daemon=True).start()

    def run_cleanup(self):
        self.set_running(True)
        try:
            if not self.rules:
                raise ValueError("등록된 규칙이 없습니다.")
            output_path = Path(self.cleanup_output_file.get().strip())
            if not output_path.name.lower().endswith(".xlsx"):
                raise ValueError("결과 파일은 .xlsx로 저장하세요.")
            output_path.parent.mkdir(parents=True, exist_ok=True)

            self.log("=== 정제 실행 시작 ===")
            self.set_status("수정 중", "현재 작업: 데이터 읽기", 10)
            headers, rows = self.read_cleanup_rows()
            self.cleanup_headers = headers

            self.set_status("수정 중", "현재 작업: 규칙 적용", 35)
            final_rows, deleted_rows, modified_rows, review_rows, rule_summary = self.apply_rules(headers, rows, preview_only=False)

            self.set_status("수정 중", "현재 작업: 결과 파일 작성", 75)
            out_wb = Workbook()
            final_ws = out_wb.active
            final_ws.title = "Final_Data"
            write_rows_to_sheet(final_ws, headers, final_rows)

            deleted_ws = out_wb.create_sheet("Deleted_Rows")
            deleted_headers = ["Rule No", "Action", "Reason", "Original Excel Row"] + headers
            write_rows_to_sheet(deleted_ws, deleted_headers, deleted_rows)

            modified_ws = out_wb.create_sheet("Modified_Rows")
            modified_headers = ["Rule No", "Original Excel Row", "Target Column", "Old Value", "New Value", "Reason"]
            write_rows_to_sheet(modified_ws, modified_headers, modified_rows)

            review_ws = out_wb.create_sheet("Review_Items")
            review_headers = ["Rule No", "Reason", "Original Excel Row"] + headers
            write_rows_to_sheet(review_ws, review_headers, review_rows)

            summary_ws = out_wb.create_sheet("Rule_Summary")
            summary_headers = ["Rule No", "Action", "Description", "Affected Rows", "Rows Remaining"]
            write_rows_to_sheet(summary_ws, summary_headers, rule_summary)

            run_ws = out_wb.create_sheet("Run_Info")
            run_ws.append(["Item", "Value"])
            run_ws.append(["Run Time", datetime.now().isoformat(timespec="seconds")])
            run_ws.append(["Input File", self.cleanup_input_file.get().strip()])
            run_ws.append(["Input Sheet", self.cleanup_sheet.get().strip()])
            run_ws.append(["Original Rows", len(rows)])
            run_ws.append(["Final Rows", len(final_rows)])
            run_ws.append(["Deleted Rows", len(deleted_rows)])
            run_ws.append(["Modified Cells", len(modified_rows)])
            run_ws.append(["Review Items", len(review_rows)])
            style_header(run_ws)
            auto_fit_columns(run_ws)

            out_wb.save(output_path)

            # 자동 연계: 2번 결과를 3번 PCWBS 비교와 4번 Display Format 입력으로 전달
            self.pcwbs_bm_file.set(str(output_path))
            self.display_bm_file.set(str(output_path))
            self.pcwbs_bm_sheet.set("Final_Data")
            self.display_bm_sheet.set("Final_Data")

            self.set_status("완료", "현재 작업: 없음", 100)
            self.log(f"수정 완료: {output_path}")
            self.log("자동 연계: 3번 및 4번 탭 입력 파일로 설정")
            self.root.after(0, lambda: messagebox.showinfo(
                "완료",
                f"수정이 완료되었습니다.\n\n"
                f"최종 행: {len(final_rows):,}\n"
                f"삭제 행: {len(deleted_rows):,}\n"
                f"수정 셀: {len(modified_rows):,}\n"
                f"검토 대상: {len(review_rows):,}\n\n"
                f"{output_path}"
            ))
        except Exception as exc:
            self.log(traceback.format_exc())
            self.set_status("오류", "현재 작업: 중단", 0)
            self.root.after(0, lambda exc=exc: messagebox.showerror("오류", str(exc)))
        finally:
            self.set_running(False)


    # ------------------------------------------------------------- PCWBS 비교
    def update_bm_key_mode_ui(self):
        """Enable controls relevant to the selected 3D BM key method."""
        mode = self.pcwbs_bm_key_mode.get()
        use_combo = mode in ["열 조합", "두 방식 교차검증"]
        use_mid = mode in ["MID 텍스트 추출", "두 방식 교차검증"]

        for combo in getattr(self, "bm_component_combos", []):
            combo.configure(state="readonly" if use_combo else "disabled")
        for combo in getattr(self, "bm_component_extract_combos", []):
            combo.configure(state="readonly" if use_combo else "disabled")
        for entry in getattr(self, "bm_component_pad_entries", []):
            entry.configure(state="normal" if use_combo else "disabled")

        if hasattr(self, "pcwbs_bm_iso_combo"):
            self.pcwbs_bm_iso_combo.configure(
                state="readonly" if use_mid else "disabled"
            )
        for name in ["pcwbs_mid_start_entry", "pcwbs_mid_length_entry"]:
            widget = getattr(self, name, None)
            if widget is not None:
                widget.configure(state="normal" if use_mid else "disabled")

    def _set_widgets_state(self, parent, enabled):
        """Enable/disable every input widget inside a frame (recursively)."""
        for child in parent.winfo_children():
            try:
                if isinstance(child, ttk.Combobox):
                    child.configure(state="readonly" if enabled else "disabled")
                elif isinstance(child, (tk.Entry, tk.Checkbutton, tk.Radiobutton, tk.Button, tk.Label)):
                    child.configure(state="normal" if enabled else "disabled")
            except tk.TclError:
                pass
            self._set_widgets_state(child, enabled)

    def _use_mapping_table(self):
        return self.pcwbs_compare_method.get() == COMPARE_METHOD_TABLE

    def update_pcwbs_compare_method_ui(self):
        """Only one of '비교키 생성 설정' and 'Mapping Table' is active at a time."""
        if not hasattr(self, "pcwbs_key_frame") or not hasattr(self, "pcwbs_mapping_frame"):
            return
        use_table = self._use_mapping_table()
        self._set_widgets_state(self.pcwbs_mapping_frame, use_table)
        self._set_widgets_state(self.pcwbs_key_frame, not use_table)
        if not use_table:
            # 생성 방식(열 조합/MID)에 따라 일부 입력은 다시 비활성화
            self.update_bm_key_mode_ui()

    def choose_pcwbs_map_file(self):
        path = filedialog.askopenfilename(
            initialdir=str(self.project_root / "03_reference"),
            filetypes=[("Excel files", "*.xlsx *.xlsm")],
        )
        if path:
            self.pcwbs_map_file.set(path)
            self._load_sheet_names_to_combo(path, self.pcwbs_map_sheet_combo, self.pcwbs_map_sheet)
            self.preview_pcwbs_mapping_table()

    def _normalize_mapping_value(self, value):
        return self._normalize_component(value, "원문", "0")

    def _mapping_separator(self):
        return self.pcwbs_map_separator.get() or "-"

    def _loose_key(self, text):
        """
        Normalize a combined key so that formatting differences do not matter.

        - 대소문자, 앞뒤/연속 공백, 구분자 주변 공백 무시
        - 숫자 앞의 0 무시 (04-4300 == 4-4300)
        """
        text = clean_text(text).upper()
        if not text:
            return ""
        separator = re.escape(self._mapping_separator())
        text = re.sub(rf"\s*{separator}\s*", self._mapping_separator(), text)
        return re.sub(r"(?<![\d.])0+(?=\d)", "", text)

    @staticmethod
    def _resolve_header(name, headers):
        """Find the real header name for a (loosely written) column name."""
        if name in headers:
            return name
        lookup = {canonical_header(header): header for header in headers}
        return lookup.get(canonical_header(name))

    def _parse_combo_header(self, header_text, headers):
        """
        Split a Mapping Table header such as 'SUBTITLE-CIA' into real column names.

        열 이름 자체에 구분자가 들어 있어도(예: SIZE-1) 실제 열 목록과 비교해서 나눕니다.
        Returns a list of column names, or None if it cannot be resolved.
        """
        text = clean_text(header_text)
        if not text:
            return None
        separator = self._mapping_separator()
        tokens = [token.strip() for token in text.split(separator)]

        def resolve_from(start):
            if start == len(tokens):
                return []
            # 긴 열 이름(구분자를 포함한 이름)을 우선 시도
            for end in range(len(tokens), start, -1):
                found = self._resolve_header(separator.join(tokens[start:end]), headers)
                if found:
                    rest = resolve_from(end)
                    if rest is not None:
                        return [found] + rest
            return None

        return resolve_from(0)

    def _combo_display(self, row, columns):
        """Combined value in its original form (for result sheets)."""
        return self._mapping_separator().join(
            self._display_cell_text(row.get(column)) for column in columns
        )

    @staticmethod
    def _display_cell_text(value):
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return clean_text(value)

    def _combo_key(self, row, columns):
        """Join the row values of the given columns with the mapping separator."""
        parts = []
        blanks = []
        for column in columns:
            value = self._normalize_mapping_value(row.get(column))
            if not value:
                blanks.append(column)
            parts.append(value)
        if blanks:
            return "", f"{', '.join(blanks)} 값이 비어 있음"
        return self._loose_key(self._mapping_separator().join(parts)), ""

    def _load_key_mapping_table(self):
        """
        Read the Mapping Table.

        Format (header row = '헤더 행'):
            A열 헤더: 3D BM 열 조합 (예: SUBTITLE-CIA)
            B열 헤더: PCWBS 열 조합 (예: Category5-Category4)
            이후 행: A = 3D BM 조합 값 (04-4300), B = PCWBS 조합 값 (04-Common (4300))
        """
        path = Path(self.pcwbs_map_file.get().strip())
        if not path.is_file():
            raise ValueError("Mapping Table 파일을 선택하세요.")
        try:
            header_row = int(self.pcwbs_map_header_row.get().strip() or "1")
        except ValueError as exc:
            raise ValueError("Mapping Table 헤더 행은 1 이상의 정수여야 합니다.") from exc
        if header_row < 1:
            raise ValueError("Mapping Table 헤더 행은 1 이상의 정수여야 합니다.")

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_name = self.pcwbs_map_sheet.get().strip() or wb.sheetnames[0]
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Mapping Table 시트를 찾을 수 없습니다: {sheet_name}")
            ws = wb[sheet_name]
            rows = ws.iter_rows(min_row=header_row, values_only=True)
            header = list(next(rows, ()))
            bm_header = clean_text(header[0]) if len(header) > 0 else ""
            pcwbs_header = clean_text(header[1]) if len(header) > 1 else ""
            if not bm_header or not pcwbs_header:
                raise ValueError(
                    "Mapping Table 헤더 행의 A열에 3D BM 열 조합(예: SUBTITLE-CIA), "
                    "B열에 PCWBS 열 조합(예: Category5-Category4)을 입력하세요."
                )

            mapping = {}
            issues = []
            pair_count = 0
            for excel_row, values in enumerate(rows, start=header_row + 1):
                bm_value = self._loose_key(
                    self._normalize_mapping_value(values[0] if len(values) > 0 else None)
                )
                pcwbs_value = self._loose_key(
                    self._normalize_mapping_value(values[1] if len(values) > 1 else None)
                )
                if not bm_value and not pcwbs_value:
                    continue
                if not bm_value or not pcwbs_value:
                    issues.append({
                        "Issue Type": "Mapping Table Blank",
                        "Excel Row": excel_row,
                        "Generated Key": self._display_cell_text(
                            values[0] if bm_value else (values[1] if len(values) > 1 else None)
                        ),
                        "Details": "A열 또는 B열 값이 비어 있어 제외함",
                    })
                    continue
                targets = mapping.setdefault(bm_value, [])
                if pcwbs_value not in targets:
                    targets.append(pcwbs_value)
                    pair_count += 1
        finally:
            wb.close()

        return {
            "bm_header": bm_header,
            "pcwbs_header": pcwbs_header,
            "map": mapping,
            "pair_count": pair_count,
            "issues": issues,
        }

    def _prepare_key_mapping(self, bm_headers, ref_headers, bm_label="3D BM"):
        """Load the Mapping Table and resolve its header combinations against both files."""
        table = self._load_key_mapping_table()
        separator = self._mapping_separator()
        bm_columns = self._parse_combo_header(table["bm_header"], bm_headers)
        pcwbs_columns = self._parse_combo_header(table["pcwbs_header"], ref_headers)
        if not bm_columns:
            raise ValueError(
                f"Mapping Table A열 헤더 '{table['bm_header']}'를 {bm_label} 열로 나눌 수 없습니다.\n"
                f"열 이름을 조합 구분자 '{separator}'로 연결했는지, "
                "3D BM 열 불러오기를 먼저 실행했는지 확인하세요."
            )
        if not pcwbs_columns:
            raise ValueError(
                f"Mapping Table B열 헤더 '{table['pcwbs_header']}'를 PCWBS 열로 나눌 수 없습니다.\n"
                f"열 이름을 조합 구분자 '{separator}'로 연결했는지, "
                "PCWBS 열 불러오기를 먼저 실행했는지 확인하세요."
            )
        table["bm_columns"] = bm_columns
        table["pcwbs_columns"] = pcwbs_columns
        table["bm_label"] = " + ".join(bm_columns)
        table["pcwbs_label"] = " + ".join(pcwbs_columns)
        self._active_key_mapping = table
        return table

    def _pcwbs_row_key(self, row, settings):
        """PCWBS key for one row, using the selected comparison method."""
        if self._use_mapping_table():
            return self._combo_key(row, self._active_key_mapping["pcwbs_columns"])
        return self._build_custom_key(row, settings)

    def _bm_row_mapped_keys(self, row):
        """PCWBS keys that one 3D BM row points to through the Mapping Table."""
        table = self._active_key_mapping
        value, error = self._combo_key(row, table["bm_columns"])
        display = self._combo_display(row, table["bm_columns"])
        if not value:
            return [], display, error
        keys = table["map"].get(value, [])
        if not keys:
            return [], display, f"Mapping Table에 없는 3D BM 조합 값: {display}"
        return list(keys), display, ""

    def preview_pcwbs_mapping_table(self):
        try:
            table = self._load_key_mapping_table()
            bm_text = table["bm_header"]
            pcwbs_text = table["pcwbs_header"]
            # 열 목록이 이미 불러와져 있으면 조합이 실제 열로 나뉘는지도 확인
            if self.pcwbs_bm_headers and self.pcwbs_ref_headers:
                resolved = self._prepare_key_mapping(self.pcwbs_bm_headers, self.pcwbs_ref_headers)
                bm_text = resolved["bm_label"]
                pcwbs_text = resolved["pcwbs_label"]
            self.pcwbs_map_info.set(
                f"3D BM [{bm_text}] → PCWBS [{pcwbs_text}] / "
                f"3D BM 조합 값 {len(table['map']):,}개, 매핑 {table['pair_count']:,}건"
                + (f" / 제외 {len(table['issues']):,}건" if table["issues"] else "")
            )
            self.log(f"Mapping Table 불러오기: {table['pair_count']:,}건")
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def show_pcwbs_mapping_example(self):
        headers = ["SUBTITLE-CIA", "Category5-Category4"]
        rows = [
            ["04-4300", "04-Common (4300)"],
            ["79-4300", "79-Common (4300)"],
            ["04-4310", "04-PE(Swing) (4310)"],
            ["12-5100", "12-Utility Area (5100)"],
        ]
        self._show_table_example(
            title="Mapping Table 입력 예시",
            subtitle=(
                "3D BM 열 조합 값과 PCWBS 열 조합 값을 직접 연결하는 표입니다. "
                "열 값이 1:1로 대응되지 않을 때 사용합니다."
            ),
            note=(
                "1행(헤더): A열 = 3D BM 열 조합, B열 = PCWBS 열 조합 (열 이름을 조합 구분자로 연결, 2개 이상 가능)\n"
                "  예) SUBTITLE-CIA  /  Category5-Category4  /  SUBTITLE-CIA-SERIAL\n"
                "2행부터: A열 = 3D BM 조합 값, B열 = 그 값에 해당하는 PCWBS 조합 값\n"
                "대소문자·공백·숫자 앞 0은 무시합니다 (04-4300 = 4-4300). "
                "같은 A값을 여러 행에 쓰면 여러 PCWBS 값에 연결됩니다."
            ),
            headers=headers,
            rows=rows,
            file_name="PCWBS_Mapping_Table_Example.xlsx",
            sheet_title="MAPPING",
        )

    def _show_table_example(self, title, subtitle, note, headers, rows, file_name, sheet_title):
        """Create an example workbook and show it in a popup (same style as Painting Code 예시)."""
        example_path = self.project_root / "03_reference" / file_name
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_header(ws)
        auto_fit_columns(ws, min_width=18)
        example_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(example_path)

        popup = tk.Toplevel(self.root)
        popup.title(title)
        popup.geometry("820x480")
        popup.transient(self.root)
        popup.grab_set()
        popup.configure(bg=self.colors["background"])

        header = tk.Frame(popup, bg=self.colors["navy"], padx=18, pady=12)
        header.pack(fill="x")
        tk.Label(
            header, text=title, bg=self.colors["navy"], fg="#FFFFFF",
            font=("Malgun Gothic", 15, "bold"), anchor="w",
        ).pack(anchor="w")
        tk.Label(
            header, text=subtitle, bg=self.colors["navy"], fg="#D4E1EC",
            font=("Malgun Gothic", 9), anchor="w", wraplength=780, justify="left",
        ).pack(anchor="w", pady=(3, 0))

        note_frame = tk.Frame(
            popup, bg="#FFF8D8", padx=12, pady=9,
            highlightbackground="#E8D792", highlightthickness=1,
        )
        note_frame.pack(fill="x", padx=14, pady=(12, 8))
        tk.Label(
            note_frame, text=note, bg="#FFF8D8", fg="#4B3D00",
            font=("Malgun Gothic", 9, "bold"), anchor="w", justify="left", wraplength=760,
        ).pack(fill="x")

        table = tk.Frame(popup, bg="#FFFFFF", highlightbackground="#D4DAE1", highlightthickness=1)
        table.pack(fill="both", expand=True, padx=14, pady=(0, 8))
        column_labels = [get_column_letter(i) for i in range(1, len(headers) + 1)]
        for column_no, letter in enumerate(column_labels):
            tk.Label(
                table, text=letter, bg="#E3E8ED", relief="solid", bd=1,
                font=("Malgun Gothic", 9, "bold"), width=24,
            ).grid(row=0, column=column_no, sticky="nsew")
        for column_no, header_text in enumerate(headers):
            tk.Label(
                table, text=header_text, bg="#F4E6B1", fg="#111111", relief="solid", bd=1,
                font=("Malgun Gothic", 9, "bold"), width=24, height=2,
            ).grid(row=1, column=column_no, sticky="nsew")
        for row_no, row_values in enumerate(rows, start=2):
            for column_no, value in enumerate(row_values):
                tk.Label(
                    table, text=value, bg="#FFFFFF", fg="#111111", relief="solid", bd=1,
                    font=("Malgun Gothic", 9), width=24,
                ).grid(row=row_no, column=column_no, sticky="nsew")

        button_bar = tk.Frame(popup, bg=self.colors["background"])
        button_bar.pack(fill="x", padx=14, pady=(0, 12))

        def open_example_excel():
            try:
                os.startfile(example_path)
            except (AttributeError, OSError):
                messagebox.showinfo("예시 파일 위치", str(example_path), parent=popup)

        tk.Button(
            button_bar, text="예시 Excel 열기", command=open_example_excel,
            bg=self.colors["navy"], fg="#FFFFFF", activebackground=self.colors["navy_hover"],
            activeforeground="#FFFFFF", relief="flat", bd=0, padx=16, pady=7, font=self.bold_font,
        ).pack(side="left")
        tk.Button(
            button_bar, text="닫기", command=popup.destroy,
            bg=self.colors["gray_button"], fg="#FFFFFF", activebackground=self.colors["gray_hover"],
            activeforeground="#FFFFFF", relief="flat", bd=0, padx=16, pady=7, font=self.bold_font,
        ).pack(side="right")

    def choose_pcwbs_bm_file(self):
        path = filedialog.askopenfilename(
            initialdir=str(Path(self.pcwbs_bm_file.get() or self.project_root).parent),
            filetypes=[("Excel files", "*.xlsx *.xlsm")],
        )
        if path:
            self.pcwbs_bm_file.set(path)
            self.load_pcwbs_bm_metadata()

    def choose_pcwbs_ref_file(self):
        path = filedialog.askopenfilename(
            initialdir=str(Path(self.pcwbs_ref_file.get() or self.project_root).parent),
            filetypes=[("Excel files", "*.xlsx *.xlsm")],
        )
        if path:
            self.pcwbs_ref_file.set(path)
            self.load_pcwbs_ref_metadata()

    def choose_pcwbs_output_file(self):
        path = filedialog.asksaveasfilename(
            initialdir=str(Path(self.pcwbs_output_file.get()).parent),
            initialfile=Path(self.pcwbs_output_file.get()).name,
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if path:
            self.pcwbs_output_file.set(path)

    def _load_sheet_names(self, file_path):
        path = Path(file_path.strip())
        if not path.is_file():
            raise ValueError(f"Excel 파일을 찾을 수 없습니다: {path}")
        wb = load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        wb.close()
        return names

    def load_pcwbs_bm_metadata(self):
        try:
            names = self._load_sheet_names(self.pcwbs_bm_file.get())
            self.pcwbs_bm_sheet_combo["values"] = names
            preferred = "Final_Data" if "Final_Data" in names else ("Merged_Data" if "Merged_Data" in names else names[0])
            self.pcwbs_bm_sheet.set(preferred)
            self.load_pcwbs_bm_headers()
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def load_pcwbs_ref_metadata(self):
        try:
            names = self._load_sheet_names(self.pcwbs_ref_file.get())
            self.pcwbs_ref_sheet_combo["values"] = names
            self.pcwbs_ref_sheet.set(names[0])
            self.load_pcwbs_ref_headers()
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def _read_headers_only(self, file_path, sheet_name, header_row_text):
        try:
            header_row = int(header_row_text)
        except ValueError as exc:
            raise ValueError("헤더 행은 1 이상의 정수여야 합니다.") from exc
        if header_row < 1:
            raise ValueError("헤더 행은 1 이상이어야 합니다.")
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"시트를 찾을 수 없습니다: {sheet_name}")
            ws = wb[sheet_name]
            values = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
            last_col = 0
            for idx, value in enumerate(values, start=1):
                if clean_text(value):
                    last_col = idx
            return make_unique_headers(list(values[:last_col]))
        finally:
            wb.close()

    def load_pcwbs_bm_headers(self):
        try:
            headers = self._read_headers_only(
                self.pcwbs_bm_file.get().strip(), self.pcwbs_bm_sheet.get().strip(), self.pcwbs_bm_header_row.get().strip()
            )
            self.pcwbs_bm_headers = headers
            # 현재 UI에서 실제 생성되는 BM 열 선택 콤보만 갱신합니다.
            # 과거 UI의 subtitle/cia 전용 콤보는 제거되었으므로 참조하지 않습니다.
            for combo in [self.pcwbs_bm_iso_combo] + self.bm_component_combos:
                combo["values"] = headers
            self._set_combo_guess(self.pcwbs_bm_subtitle, headers, ["SUBTITLE", "Subtitle"])
            self._set_combo_guess(self.pcwbs_bm_cia, headers, ["CIA"])
            self._set_combo_guess(self.pcwbs_bm_iso, headers, ["ISO_DWG_ID", "ISO DWG ID", "ISO DWG"])
            self._set_combo_guess(self.bm_component_column[0], headers, ["SUBTITLE", "Subtitle"])
            self._set_combo_guess(self.bm_component_column[1], headers, ["CIA"])
            self.log(f"PCWBS 비교용 BM 열 불러오기: {len(headers)}개")
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def load_pcwbs_ref_headers(self):
        try:
            headers = self._read_headers_only(
                self.pcwbs_ref_file.get().strip(), self.pcwbs_ref_sheet.get().strip(), self.pcwbs_ref_header_row.get().strip()
            )
            self.pcwbs_ref_headers = headers
            for combo in self.pcwbs_component_combos:
                combo["values"] = headers
            # 기본값 자동 추정
            self._set_combo_guess(self.pcwbs_component_column[0], headers, ["Category5", "Category 5"])
            self._set_combo_guess(self.pcwbs_component_column[1], headers, ["Category4", "Category 4"])
            self.log(f"PCWBS 기준 열 불러오기: {len(headers)}개")
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    @staticmethod
    def _set_combo_guess(variable, headers, candidates):
        canonical_map = {canonical_header(header): header for header in headers}
        for candidate in candidates:
            found = canonical_map.get(canonical_header(candidate))
            if found:
                variable.set(found)
                return
        if headers and variable.get() not in headers:
            variable.set(headers[0])

    @staticmethod
    def _extract_component(value, mode):
        text = clean_text(value)
        if not text:
            return ""
        if mode == "원문":
            return text
        if mode == "마지막 괄호 안":
            matches = re.findall(r"\(([^()]*)\)", text)
            return clean_text(matches[-1]) if matches else ""
        if mode == "첫 괄호 안":
            matches = re.findall(r"\(([^()]*)\)", text)
            return clean_text(matches[0]) if matches else ""
        if mode == "숫자만":
            return "".join(re.findall(r"\d+", text))
        if mode == "마지막 숫자":
            matches = re.findall(r"\d+", text)
            return matches[-1] if matches else ""
        return text

    def _normalize_component(self, value, mode, pad_text):
        text = self._extract_component(value, mode)
        if not text:
            return ""
        try:
            pad = max(0, int(clean_text(pad_text) or "0"))
        except ValueError:
            pad = 0
        # 4.0 같은 Excel 숫자 표현 제거
        if re.fullmatch(r"\d+(?:\.0+)?", text):
            text = str(int(float(text)))
        if pad and text.isdigit():
            text = text.zfill(pad)
        return text.upper()

    def _enabled_component_settings(self, is_pcwbs=True):
        enabled = self.pcwbs_component_enabled if is_pcwbs else self.bm_component_enabled
        columns = self.pcwbs_component_column if is_pcwbs else self.bm_component_column
        extracts = self.pcwbs_component_extract if is_pcwbs else self.bm_component_extract
        pads = self.pcwbs_component_pad if is_pcwbs else self.bm_component_pad
        settings = []
        for i in range(5):
            if enabled[i].get():
                settings.append((columns[i].get(), extracts[i].get(), pads[i].get()))
        return settings

    def _build_custom_key(self, row, settings):
        parts = []
        errors = []
        for column, mode, pad in settings:
            value = self._normalize_component(row.get(column), mode, pad)
            if not value:
                errors.append(f"{column} 값 생성 실패")
            parts.append(value)
        if errors:
            return "", "; ".join(errors)
        separator = self.pcwbs_key_separator.get()
        return separator.join(parts), ""

    def _build_bm_keys(self, row):
        """Build a column-combination key, a MID key, or both."""
        mode = self.pcwbs_bm_key_mode.get()
        combination_key = ""
        mid_key = ""
        errors = []

        if mode in ["열 조합", "두 방식 교차검증"]:
            settings = self._enabled_component_settings(False)
            if settings:
                combination_key, error = self._build_custom_key(row, settings)
                if error:
                    errors.append(error)
            else:
                errors.append("3D BM 열 조합 구성요소가 선택되지 않음")

        if mode in ["MID 텍스트 추출", "두 방식 교차검증"]:
            source_text = clean_text(row.get(self.pcwbs_bm_iso.get()))
            try:
                mid_start = int(self.pcwbs_mid_start.get())
                mid_length = int(self.pcwbs_mid_length.get())
            except ValueError:
                mid_start, mid_length = 0, 0

            if mid_start < 1 or mid_length < 1:
                errors.append("MID 시작/길이 설정 오류")
            elif len(source_text) < mid_start:
                errors.append("MID 대상 문자열이 시작 위치보다 짧음")
            else:
                mid_key = source_text[
                    mid_start - 1:mid_start - 1 + mid_length
                ].strip().upper()
                if not mid_key:
                    errors.append("MID 텍스트 추출 결과가 비어 있음")

        # Keep the legacy tuple shape used by the comparison/export code.
        return combination_key, mid_key, combination_key, "; ".join(errors)

    @staticmethod
    def _status(value):
        text = clean_text(value).casefold()
        if text in {"ok", "yes", "y", "true", "1", "있음"}:
            return "OK"
        if text in {"no", "n", "false", "0", "없음"}:
            return "No"
        return ""

    def _collect_pcwbs_inputs(self):
        bm_path = Path(self.pcwbs_bm_file.get().strip())
        ref_path = Path(self.pcwbs_ref_file.get().strip())
        if not bm_path.is_file():
            raise ValueError("정제 완료 3D BM 파일을 선택하세요.")
        if not ref_path.is_file():
            raise ValueError("PCWBS 기준 파일을 선택하세요.")
        try:
            bm_header_row = int(self.pcwbs_bm_header_row.get())
            ref_header_row = int(self.pcwbs_ref_header_row.get())
        except ValueError as exc:
            raise ValueError("헤더 행은 정수여야 합니다.") from exc

        result_column = clean_text(self.pcwbs_result_column.get())
        if not result_column:
            raise ValueError("결과 열명을 입력하세요.")

        if self._use_mapping_table():
            self._prepare_key_mapping(self.pcwbs_bm_headers, self.pcwbs_ref_headers)
            return bm_path, ref_path, bm_header_row, ref_header_row, result_column

        pcwbs_settings = self._enabled_component_settings(True)
        if not pcwbs_settings:
            raise ValueError("PCWBS 조합키 구성요소를 하나 이상 선택하세요.")
        for column, _mode, _pad in pcwbs_settings:
            if column not in self.pcwbs_ref_headers:
                raise ValueError(f"PCWBS 조합 열을 확인하세요: {column}")

        mode = self.pcwbs_bm_key_mode.get()
        if mode in ["열 조합", "두 방식 교차검증"]:
            bm_settings = self._enabled_component_settings(False)
            if not bm_settings:
                raise ValueError("3D BM 열 조합 구성요소를 하나 이상 선택하세요.")
            for column, _mode, _pad in bm_settings:
                if column not in self.pcwbs_bm_headers:
                    raise ValueError(f"3D BM 조합 열을 확인하세요: {column}")

        if mode in ["MID 텍스트 추출", "두 방식 교차검증"]:
            if self.pcwbs_bm_iso.get() not in self.pcwbs_bm_headers:
                raise ValueError("3D BM MID 대상 열을 확인하세요.")
            try:
                mid_start = int(self.pcwbs_mid_start.get())
                mid_length = int(self.pcwbs_mid_length.get())
            except ValueError as exc:
                raise ValueError("MID 시작과 길이는 정수여야 합니다.") from exc
            if mid_start < 1 or mid_length < 1:
                raise ValueError("MID 시작과 길이는 1 이상이어야 합니다.")

        return bm_path, ref_path, bm_header_row, ref_header_row, result_column

    def _stream_bm_keys(self, bm_path, sheet_name, header_row):
        """대용량 BM을 메모리에 전부 올리지 않고 한 행씩 처리한다."""
        wb = load_workbook(bm_path, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"시트를 찾을 수 없습니다: {sheet_name}")
            ws = wb[sheet_name]
            header_values = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
            last_col = 0
            for idx, value in enumerate(header_values, start=1):
                if clean_text(value):
                    last_col = idx
            headers = make_unique_headers(list(header_values[:last_col]))
            needed = set()
            mode = self.pcwbs_bm_key_mode.get()
            use_table = self._use_mapping_table()
            if use_table:
                needed.update(self._active_key_mapping["bm_columns"])
            elif mode in ["열 조합", "두 방식 교차검증"]:
                needed.update(
                    column
                    for column, _mode, _pad
                    in self._enabled_component_settings(False)
                )
            if not use_table and mode in ["MID 텍스트 추출", "두 방식 교차검증"]:
                needed.add(self.pcwbs_bm_iso.get())
            index_map = {header: idx for idx, header in enumerate(headers) if header in needed}

            subtitle_keys, iso_keys, custom_keys = set(), set(), set()
            key_counts = Counter()
            errors, mismatches = [], []
            unmapped_counts = Counter()
            unmapped_first_row = {}
            processed = 0
            estimated = max(ws.max_row - header_row, 1)
            for excel_row, values in enumerate(ws.iter_rows(min_row=header_row + 1, max_col=last_col, values_only=True), start=header_row + 1):
                if is_effectively_empty_row(values):
                    continue
                processed += 1
                row = {header: values[idx] if idx < len(values) else None for header, idx in index_map.items()}
                if use_table:
                    # Mapping Table: 3D BM 값 → PCWBS 값 목록. 미등록 값은 값별로 한 번만 기록.
                    mapped_keys, bm_value, error = self._bm_row_mapped_keys(row)
                    for key in mapped_keys:
                        subtitle_keys.add(key)
                        key_counts[key] += 1
                    if error:
                        unmapped_counts[error] += 1
                        unmapped_first_row.setdefault(error, (excel_row, bm_value))
                    if processed % 5000 == 0:
                        percent = 5 + min(35, (excel_row - header_row) / estimated * 35)
                        self.set_status("PCWBS 비교 중", f"현재 작업: BM 키 계산 {processed:,}행", percent)
                    continue
                subtitle_key, iso_key, custom_key, error = self._build_bm_keys(row)
                for key, target_set in [(subtitle_key, subtitle_keys), (iso_key, iso_keys), (custom_key, custom_keys)]:
                    if key:
                        target_set.add(key)
                        key_counts[key] += 1
                if error:
                    errors.append({"Excel Row": excel_row, "Error": error, "Subtitle Key": subtitle_key, "ISO Key": iso_key, "Custom Key": custom_key})
                if mode == "두 방식 교차검증" and subtitle_key and iso_key and subtitle_key != iso_key:
                    mismatches.append({"Excel Row": excel_row, "Column Combination Key": subtitle_key, "MID Text Key": iso_key})
                if processed % 5000 == 0:
                    percent = 5 + min(35, (excel_row - header_row) / estimated * 35)
                    self.set_status("PCWBS 비교 중", f"현재 작업: BM 키 계산 {processed:,}행", percent)
                    self.log(f"BM 키 계산 진행: {processed:,}행")
            for error, count in unmapped_counts.items():
                first_row, bm_value = unmapped_first_row[error]
                errors.append({
                    "Excel Row": first_row,
                    "Error": f"{error} ({count:,}행)",
                    "Custom Key": bm_value,
                })
            return processed, subtitle_keys, iso_keys, custom_keys, key_counts, errors, mismatches
        finally:
            wb.close()

    def _read_reference_rows_fast(self, ref_path, sheet_name, header_row):
        wb = load_workbook(ref_path, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"시트를 찾을 수 없습니다: {sheet_name}")
            ws = wb[sheet_name]
            header_values = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
            last_col = 0
            for idx, value in enumerate(header_values, start=1):
                if clean_text(value):
                    last_col = idx
            headers = make_unique_headers(list(header_values[:last_col]))
            rows = []
            for excel_row, values in enumerate(ws.iter_rows(min_row=header_row + 1, max_col=last_col, values_only=True), start=header_row + 1):
                if is_effectively_empty_row(values):
                    continue
                row = dict(zip(headers, values))
                row["__Excel_Row__"] = excel_row
                rows.append(row)
            return headers, rows
        finally:
            wb.close()

    def calculate_pcwbs_check(self):
        bm_path, ref_path, bm_header_row, ref_header_row, result_column = self._collect_pcwbs_inputs()
        self.set_status("PCWBS 비교 중", "현재 작업: BM 파일 스트리밍 읽기", 5)
        bm_count, subtitle_keys, iso_keys, custom_keys, bm_key_counts, bm_key_errors, bm_key_mismatches = self._stream_bm_keys(
            bm_path, self.pcwbs_bm_sheet.get(), bm_header_row
        )
        mode = self.pcwbs_bm_key_mode.get()
        use_table = self._use_mapping_table()
        if use_table or mode == "열 조합":
            active_keys = subtitle_keys
        elif mode == "MID 텍스트 추출":
            active_keys = iso_keys
        else:
            active_keys = subtitle_keys | iso_keys

        self.set_status("PCWBS 비교 중", "현재 작업: PCWBS 기준 읽기", 45)
        ref_headers, ref_rows = self._read_reference_rows_fast(ref_path, self.pcwbs_ref_sheet.get(), ref_header_row)

        self.set_status("PCWBS 비교 중", "현재 작업: 행별 일치 여부 계산", 55)
        settings = self._enabled_component_settings(True)
        pcwbs_key_counts = Counter()
        key_error_rows = []
        result_rows = []
        matched_count = 0
        unmatched_count = 0

        shown_keys = {}
        for idx, row in enumerate(ref_rows, start=1):
            key, key_error = self._pcwbs_row_key(row, settings)
            if key:
                pcwbs_key_counts[key] += 1
            if key_error:
                key_error_rows.append({
                    "Issue Type": "PCWBS Key Error",
                    "Excel Row": row["__Excel_Row__"],
                    "Generated Key": key,
                    "Details": key_error,
                })

            is_match = bool(key and key in active_keys)
            match_result = "일치" if is_match else "불일치"
            if is_match:
                matched_count += 1
            else:
                unmatched_count += 1

            if use_table:
                # 비교는 정규화 키로 하고, 결과에는 원래 모양 그대로 표시
                shown_key = self._combo_display(row, self._active_key_mapping["pcwbs_columns"])
            else:
                shown_key = key
            if key:
                shown_keys.setdefault(key, shown_key)
            output_row = {header: row.get(header) for header in ref_headers}
            output_row.update({
                "PCWBS Generated Key": shown_key,
                result_column: match_result,
            })
            result_rows.append(output_row)

            if idx % 1000 == 0:
                self.set_status(
                    "PCWBS 비교 중",
                    f"현재 작업: PCWBS 비교 {idx:,}/{len(ref_rows):,}",
                    55 + min(15, idx / max(len(ref_rows), 1) * 15),
                )

        duplicate_rows = [
            {
                "Issue Type": "PCWBS Duplicate Key",
                "Excel Row": "",
                "Generated Key": shown_keys.get(key, key),
                "Details": f"PCWBS에서 {count}회 중복",
            }
            for key, count in pcwbs_key_counts.items() if count > 1
        ]
        bm_duplicate_rows = [
            {
                "Issue Type": "BM Duplicate Key",
                "Excel Row": "",
                "Generated Key": shown_keys.get(key, key),
                "Details": f"3D BM에서 {count}회 출현",
            }
            for key, count in bm_key_counts.items() if count > 1
        ]
        bm_error_issue_rows = [
            {
                "Issue Type": "BM Key Error",
                "Excel Row": row.get("Excel Row", ""),
                "Generated Key": row.get("Subtitle Key") or row.get("ISO Key") or row.get("Custom Key") or "",
                "Details": row.get("Error", ""),
            }
            for row in bm_key_errors
        ]
        mismatch_issue_rows = [
            {
                "Issue Type": "BM Key Mismatch",
                "Excel Row": row.get("Excel Row", ""),
                "Generated Key": f"{row.get('Column Combination Key', '')} / {row.get('MID Text Key', '')}",
                "Details": "열 조합 키와 MID 텍스트 추출 키가 서로 다름",
            }
            for row in bm_key_mismatches
        ]
        issue_rows = key_error_rows + duplicate_rows + bm_duplicate_rows + bm_error_issue_rows + mismatch_issue_rows
        if use_table:
            issue_rows = self._active_key_mapping["issues"] + issue_rows
            mode = (
                f"Mapping Table ({self._active_key_mapping['bm_label']} → "
                f"{self._active_key_mapping['pcwbs_label']})"
            )
            key_components = f"Mapping Table: {self.pcwbs_map_file.get()}"
        else:
            key_components = " | ".join(f"{c}:{m}:pad{p}" for c, m, p in settings)

        summary = {
            "BM Rows": bm_count,
            "PCWBS Rows": len(ref_rows),
            "BM Active Unique Keys": len(active_keys),
            "Matched": matched_count,
            "Unmatched": unmatched_count,
            "Issue Count": len(issue_rows),
            "Result Column": result_column,
            "BM Key Mode": mode,
            "PCWBS Key Components": key_components,
        }
        return {
            "ref_headers": ref_headers,
            "result_column": result_column,
            "result_rows": result_rows,
            "issue_rows": issue_rows,
            "summary": summary,
        }

    def start_pcwbs_preview_thread(self):
        if self.is_running:
            return
        threading.Thread(target=self.run_pcwbs_preview, daemon=True).start()

    def start_pcwbs_run_thread(self):
        if self.is_running:
            return
        threading.Thread(target=self.run_pcwbs_save, daemon=True).start()

    def run_pcwbs_preview(self):
        self.set_running(True)
        try:
            result = self.calculate_pcwbs_check()
            s = result["summary"]
            message = (
                f"PCWBS 전체 {s['PCWBS Rows']:,}행\n"
                f"일치 {s['Matched']:,}행 / 불일치 {s['Unmatched']:,}행\n"
                f"생성된 3D BM 고유 키 {s['BM Active Unique Keys']:,}개\n"
                f"키 생성·중복·교차검증 참고사항 {s['Issue Count']:,}건"
            )
            self.root.after(0, self.pcwbs_status_text.set, message.replace("\n", " | "))
            self.log("PCWBS 미리보기 완료: " + message.replace("\n", " / "))
            self.set_status("완료", "현재 작업: 없음", 100)
            self.root.after(0, lambda: messagebox.showinfo("PCWBS 비교 미리보기", message))
        except Exception as exc:
            self.log(traceback.format_exc())
            self.set_status("오류", "현재 작업: 중단", 0)
            self.root.after(0, lambda exc=exc: messagebox.showerror("오류", str(exc)))
        finally:
            self.set_running(False)

    def run_pcwbs_save(self):
        self.set_running(True)
        try:
            result = self.calculate_pcwbs_check()
            output_path = Path(self.pcwbs_output_file.get().strip())
            if output_path.suffix.casefold() != ".xlsx":
                raise ValueError("결과 파일은 .xlsx 형식이어야 합니다.")
            output_path.parent.mkdir(parents=True, exist_ok=True)

            self.set_status("PCWBS 비교 중", "현재 작업: 결과 시트 작성", 75)
            out_wb = Workbook()
            result_ws = out_wb.active
            result_ws.title = "PCWBS_Check_Result"
            result_headers = result["ref_headers"] + ["PCWBS Generated Key", result["result_column"]]
            unique_headers = []
            for header in result_headers:
                if header not in unique_headers:
                    unique_headers.append(header)
            write_rows_to_sheet(result_ws, unique_headers, result["result_rows"])

            issue_ws = out_wb.create_sheet("Key_Issues")
            write_rows_to_sheet(
                issue_ws,
                ["Issue Type", "Excel Row", "Generated Key", "Details"],
                result["issue_rows"],
            )

            info_ws = out_wb.create_sheet("Run_Info")
            info_ws.append(["Item", "Value"])
            for key, value in result["summary"].items():
                info_ws.append([key, value])
            info_ws.append(["BM File", self.pcwbs_bm_file.get()])
            info_ws.append(["PCWBS File", self.pcwbs_ref_file.get()])
            info_ws.append(["Run Time", datetime.now().isoformat(timespec="seconds")])
            style_header(info_ws)
            auto_fit_columns(info_ws)

            self.set_status("PCWBS 비교 중", "현재 작업: Excel 파일 저장", 90)
            self.log(f"Excel 저장 시작: {output_path}")
            out_wb.save(output_path)
            self.set_status("완료", "현재 작업: 없음", 100)
            self.log(f"PCWBS 비교 결과 저장 완료: {output_path}")
            self.root.after(
                0,
                lambda: messagebox.showinfo(
                    "완료",
                    "PCWBS 비교가 완료되었습니다.\n\n"
                    f"일치: {result['summary']['Matched']:,}행\n"
                    f"불일치: {result['summary']['Unmatched']:,}행\n\n"
                    f"저장 위치:\n{output_path}",
                ),
            )
        except Exception as exc:
            self.log(traceback.format_exc())
            self.set_status("오류", "현재 작업: 중단", 0)
            self.root.after(0, lambda exc=exc: messagebox.showerror("오류", str(exc)))
        finally:
            self.set_running(False)


    # ------------------------------------------------------------------ Display Format
    def _build_display_tab(self):
        output_dir = self.project_root / "04_output"
        self.display_bm_file = tk.StringVar(value=str(output_dir / "BM_Modified_Result.xlsx"))
        self.display_bm_sheet = tk.StringVar(value="Final_Data")
        self.display_bm_header_row = tk.StringVar(value="1")
        self.display_output_file = tk.StringVar(value=str(output_dir / "Display_Format_Base.xlsx"))
        self.display_part_default = tk.StringVar(value="PPA")
        self.display_unit_default = tk.StringVar(value="mm")
        self.display_unit_item_column = tk.StringVar(value="Item")
        self.display_unit_operator = tk.StringVar(value="목록 중 하나 포함")
        self.display_unit_values = tk.StringVar(value="BT")
        self.display_cd_site = tk.StringVar(value="")
        self.display_times = tk.StringVar(value="")
        self.display_mtrl_group = tk.StringVar(value="RUSSIA")
        self.display_filename = tk.StringVar(value="Display_Format_01.xlsx")
        self.display_main_sub_blank = tk.BooleanVar(value=True)
        self.display_strip_main_sub_units = tk.BooleanVar(value=True)
        self.display_use_pcwbs_mapping = tk.BooleanVar(value=True)

        file_frame = tk.LabelFrame(self.display_tab, text="입력 및 결과 파일", padx=8, pady=8)
        file_frame.pack(fill="x", padx=10, pady=(10, 5))
        file_frame.columnconfigure(1, weight=1)
        self._file_row(file_frame, 0, "수정 완료 3D BM", self.display_bm_file, self.choose_display_bm_file)
        tk.Label(file_frame, text="3D BM 시트").grid(row=1, column=0, sticky="w", pady=4)
        self.display_bm_sheet_combo = ttk.Combobox(file_frame, textvariable=self.display_bm_sheet, state="readonly", width=28)
        self.display_bm_sheet_combo.grid(row=1, column=1, sticky="w", padx=6)
        tk.Label(file_frame, text="헤더 행").grid(row=1, column=2, sticky="e")
        tk.Entry(file_frame, textvariable=self.display_bm_header_row, width=7).grid(row=1, column=3, padx=5)
        tk.Button(file_frame, text="3D BM 열 불러오기", command=self.load_display_bm_columns, width=17).grid(row=1, column=4, padx=5)
        self._file_row(file_frame, 2, "결과 파일", self.display_output_file, self.choose_display_output_file, save=True)

        pcwbs_map_frame = tk.LabelFrame(
            self.display_tab,
            text="Category1~10 PCWBS 매핑",
            padx=8,
            pady=5,
        )
        pcwbs_map_frame.pack(fill="x", padx=10, pady=3)
        pcwbs_map_frame.grid_columnconfigure(0, weight=1)

        tk.Label(
            pcwbs_map_frame,
            text=(
                "PCWBS 기준 파일·시트와 비교 방식(비교키 생성 설정 또는 Mapping Table)은 "
                "3번 'PCWBS 비교' 탭의 현재 설정을 그대로 사용합니다."
            ),
            anchor="w",
        ).grid(row=0, column=0, sticky="w")

        tk.Label(
            pcwbs_map_frame,
            text=(
                "A~AS 매핑에서 입력 방식을 'PCWBS 파일 매핑'으로 지정한 "
                "Category 열만 PCWBS에서 가져옵니다."
            ),
            fg="#555555",
            anchor="w",
        ).grid(row=1, column=0, sticky="w", pady=(3, 0))

        settings = tk.LabelFrame(self.display_tab, text="Display Format 기본 설정", padx=8, pady=8)
        settings.pack(fill="x", padx=10, pady=5)
        labels = [
            ("Part 기본값", self.display_part_default, 0, 0),
            ("Unit Size 기본값", self.display_unit_default, 0, 2),
            ("CD_SITE", self.display_cd_site, 0, 4),
            ("Times", self.display_times, 1, 0),
            ("Mtrl Group", self.display_mtrl_group, 1, 2),
            ("FileName 기본값", self.display_filename, 1, 4),
        ]
        for label, var, r, c in labels:
            tk.Label(settings, text=label).grid(row=r, column=c, sticky="w", padx=(0, 4), pady=3)
            tk.Entry(settings, textvariable=var, width=22).grid(row=r, column=c + 1, sticky="w", padx=(0, 12), pady=3)
        tk.Checkbutton(
            settings,
            text="Main과 Sub가 같으면 Sub를 빈칸 처리",
            variable=self.display_main_sub_blank,
        ).grid(row=2, column=0, columnspan=3, sticky="w", pady=4)
        tk.Checkbutton(
            settings,
            text="Main/Sub 값에서 단위 제거 (예: 25 mm → 25)",
            variable=self.display_strip_main_sub_units,
        ).grid(row=3, column=0, columnspan=3, sticky="w", pady=4)
        tk.Label(settings, text="Unit Size 예외 Item 열").grid(row=2, column=3, sticky="e")
        self.display_unit_item_combo = ttk.Combobox(settings, textvariable=self.display_unit_item_column, state="readonly", width=19)
        self.display_unit_item_combo.grid(row=2, column=4, padx=5)
        self.display_unit_operator_combo = ttk.Combobox(
            settings, textvariable=self.display_unit_operator, values=OPERATORS, state="readonly", width=18
        )
        self.display_unit_operator_combo.grid(row=2, column=5, padx=5)
        tk.Entry(settings, textvariable=self.display_unit_values, width=30).grid(row=2, column=6, padx=5)
        tk.Label(settings, text="일치 시: U x mm").grid(row=2, column=7, sticky="w")

        buttons = tk.Frame(self.display_tab)
        buttons.pack(side="bottom", fill="x", padx=10, pady=(2, 7))
        tk.Button(buttons, text="매핑 저장", command=self.save_display_mapping, width=14).pack(side="left")
        tk.Button(buttons, text="매핑 불러오기", command=self.load_display_mapping, width=14).pack(side="left", padx=5)
        self.display_preview_button = tk.Button(
            buttons,
            text="변환 미리보기",
            command=lambda: self.start_thread(self.preview_display_format),
            width=16,
        )
        self.display_preview_button.pack(side="left", padx=12)
        self.display_generate_button = tk.Button(
            buttons,
            text="Display Format 생성",
            command=lambda: self.start_thread(self.generate_display_format),
            width=20,
        )
        self.display_generate_button.pack(side="left", padx=5)

        edit_frame = tk.LabelFrame(
            self.display_tab,
            text="선택 행 매핑 설정",
            padx=8,
            pady=7,
        )
        edit_frame.pack(
            side="bottom",
            fill="x",
            padx=10,
            pady=(3, 4),
        )
        edit_frame.grid_columnconfigure(3, weight=1)

        self.display_selected_header = tk.StringVar(value="")
        self.display_mapping_mode = tk.StringVar(value="3D BM 열 복사")
        self.display_mapping_source = tk.StringVar(value="")

        tk.Label(edit_frame, text="Display Format 열").grid(
            row=0,
            column=0,
            sticky="w",
            padx=(0, 4),
        )
        tk.Label(
            edit_frame,
            textvariable=self.display_selected_header,
            width=20,
            anchor="w",
            relief="sunken",
            padx=4,
        ).grid(row=0, column=1, sticky="w", padx=(0, 10))

        tk.Label(edit_frame, text="입력 방식").grid(
            row=0,
            column=2,
            sticky="e",
            padx=(0, 4),
        )
        self.display_mapping_mode_combo = ttk.Combobox(
            edit_frame,
            textvariable=self.display_mapping_mode,
            values=MAPPING_MODES,
            state="readonly",
            width=18,
        )
        self.display_mapping_mode_combo.grid(
            row=0,
            column=3,
            sticky="w",
            padx=(0, 10),
        )
        self.display_mapping_mode_combo.bind(
            "<<ComboboxSelected>>",
            self.on_display_mapping_mode_change,
        )

        tk.Label(edit_frame, text="3D BM 열 또는 값").grid(
            row=1,
            column=0,
            sticky="w",
            pady=(6, 0),
            padx=(0, 4),
        )
        self.display_source_combo = ttk.Combobox(
            edit_frame,
            textvariable=self.display_mapping_source,
            width=36,
        )
        self.display_source_combo.grid(
            row=1,
            column=1,
            columnspan=3,
            sticky="ew",
            pady=(6, 0),
            padx=(0, 10),
        )

        self.display_apply_mapping_button = tk.Button(
            edit_frame,
            text="선택 매핑 적용",
            command=self.apply_display_mapping,
            width=15,
        )
        self.display_apply_mapping_button.grid(
            row=0,
            column=4,
            rowspan=2,
            sticky="ns",
            padx=(4, 4),
        )

        self.display_auto_map_button = tk.Button(
            edit_frame,
            text="헤더명 자동 매핑",
            command=self.auto_map_display_headers,
            width=15,
        )
        self.display_auto_map_button.grid(
            row=0,
            column=5,
            rowspan=2,
            sticky="ns",
            padx=(4, 0),
        )

        self.display_mode_help = tk.StringVar(value="")
        tk.Label(
            edit_frame,
            textvariable=self.display_mode_help,
            fg="#1f4e79",
            anchor="w",
            justify="left",
            wraplength=1100,
        ).grid(row=2, column=0, columnspan=6, sticky="w", pady=(6, 0))
        self._update_display_mode_help()


        body = tk.Frame(self.display_tab)
        body.pack(fill="both", expand=True, padx=10, pady=(4, 3))

        body.grid_rowconfigure(0, weight=1)
        body.grid_columnconfigure(0, weight=5)
        body.grid_columnconfigure(1, weight=1)

        map_frame = tk.LabelFrame(body, text="A~AS 열 매핑", padx=6, pady=6)
        map_frame.grid(row=0, column=0, sticky="nsew")
        tree_container = tk.Frame(map_frame)
        tree_container.pack(fill="both", expand=True)
        tree_container.grid_rowconfigure(0, weight=1)
        tree_container.grid_columnconfigure(0, weight=1)

        self.display_map_tree = ttk.Treeview(
            tree_container,
            columns=("display", "mode", "source"),
            show="headings",
            height=9,
        )
        for col, title, width in [
            ("display", "Display Format 열", 230),
            ("mode", "입력 방식", 150),
            ("source", "3D BM 열 또는 값", 300),
        ]:
            self.display_map_tree.heading(col, text=title)
            self.display_map_tree.column(col, width=width, anchor="w")

        display_y_scroll = ttk.Scrollbar(
            tree_container,
            orient="vertical",
            command=self.display_map_tree.yview,
        )
        display_x_scroll = ttk.Scrollbar(
            tree_container,
            orient="horizontal",
            command=self.display_map_tree.xview,
        )
        self.display_map_tree.configure(
            yscrollcommand=display_y_scroll.set,
            xscrollcommand=display_x_scroll.set,
        )
        self.display_map_tree.grid(row=0, column=0, sticky="nsew")
        display_y_scroll.grid(row=0, column=1, sticky="ns")
        display_x_scroll.grid(row=1, column=0, sticky="ew")
        self.display_map_tree.bind(
            "<<TreeviewSelect>>",
            self.on_display_mapping_select,
        )

        rule_frame = tk.LabelFrame(body, text="Part 조건 규칙", padx=6, pady=6)
        rule_frame.grid(
            row=0,
            column=1,
            sticky="nsew",
            padx=(8, 0),
        )
        rule_frame.grid_rowconfigure(5, weight=1)
        self.part_rule_column = tk.StringVar(value="Class")
        self.part_rule_operator = tk.StringVar(value="포함")
        self.part_rule_value = tk.StringVar(value="")
        self.part_rule_result = tk.StringVar(value="PPU")
        tk.Label(rule_frame, text="조건 열").grid(row=0, column=0, sticky="w")
        self.part_rule_column_combo = ttk.Combobox(rule_frame, textvariable=self.part_rule_column, state="readonly", width=18)
        self.part_rule_column_combo.grid(row=0, column=1, padx=4, pady=3)
        tk.Label(rule_frame, text="조건").grid(row=1, column=0, sticky="w")
        ttk.Combobox(rule_frame, textvariable=self.part_rule_operator, values=OPERATORS, state="readonly", width=18).grid(row=1, column=1, padx=4, pady=3)
        tk.Label(rule_frame, text="값").grid(row=2, column=0, sticky="w")
        tk.Entry(rule_frame, textvariable=self.part_rule_value, width=21).grid(row=2, column=1, padx=4, pady=3)
        tk.Label(rule_frame, text="Part").grid(row=3, column=0, sticky="w")
        ttk.Combobox(rule_frame, textvariable=self.part_rule_result, values=PART_VALUES, state="readonly", width=18).grid(row=3, column=1, padx=4, pady=3)
        tk.Button(rule_frame, text="규칙 추가", command=self.add_part_rule).grid(row=4, column=0, columnspan=2, pady=5)
        self.part_rule_list = tk.Listbox(rule_frame, width=38, height=13)
        self.part_rule_list.grid(
            row=5,
            column=0,
            columnspan=2,
            sticky="nsew",
            pady=4,
        )
        tk.Button(rule_frame, text="선택 규칙 삭제", command=self.delete_part_rule).grid(row=6, column=0, columnspan=2, pady=3)

        self.reset_display_mappings()

    def reset_display_mappings(self):
        self.display_mappings = {}
        for header in DISPLAY_FORMAT_HEADERS:
            mode, source = DISPLAY_DEFAULT_MODES.get(header, ("3D BM 열 복사", header))
            self.display_mappings[header] = {"mode": mode, "source": source}
        if hasattr(self, "display_map_tree"):
            self.refresh_display_mapping_tree()

    def refresh_display_mapping_tree(self):
        self.display_map_tree.delete(*self.display_map_tree.get_children())
        for header in DISPLAY_FORMAT_HEADERS:
            mapping = self.display_mappings.get(
                header,
                {"mode": "빈칸", "source": ""},
            )
            visible_source = mapping.get("source", "")
            if mapping.get("mode") in ["빈칸", "자동번호", "특수규칙"]:
                visible_source = ""
            self.display_map_tree.insert(
                "",
                "end",
                iid=header,
                values=(header, mapping["mode"], visible_source),
            )

    def choose_display_bm_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm")])
        if path:
            self.display_bm_file.set(path)
            self._load_sheet_names_to_combo(path, self.display_bm_sheet_combo, self.display_bm_sheet)

    def load_display_pcwbs_metadata(self):
        """Use the PCWBS file, sheet and key settings already configured in tab 3."""
        try:
            self.load_pcwbs_ref_metadata()
            self.log(
                "Display Format용 PCWBS 설정 확인: "
                f"{self.pcwbs_ref_sheet.get()} / 열 {len(self.pcwbs_ref_headers)}개"
            )
        except Exception as exc:
            messagebox.showerror("오류", str(exc))

    def choose_display_output_file(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if path:
            self.display_output_file.set(path)

    def _load_sheet_names_to_combo(self, path, combo, variable):
        wb = load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        wb.close()
        combo["values"] = names
        if names and variable.get() not in names:
            variable.set(names[0])

    def load_display_bm_columns(self):
        path = Path(self.display_bm_file.get().strip())
        if not path.exists():
            messagebox.showerror("오류", "3D BM 파일을 찾을 수 없습니다.")
            return
        self._load_sheet_names_to_combo(path, self.display_bm_sheet_combo, self.display_bm_sheet)
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[self.display_bm_sheet.get()]
        header_row = int(self.display_bm_header_row.get() or "1")
        headers = make_unique_headers([cell.value for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row))])
        wb.close()
        self.display_bm_headers = headers
        for combo in [self.display_source_combo, self.part_rule_column_combo, self.display_unit_item_combo]:
            combo["values"] = headers
        self.auto_map_display_headers()
        self.log(f"Display Format용 BM 열 불러오기: {len(headers)}개")

    def auto_map_display_headers(self):
        if not self.display_bm_headers:
            self.refresh_display_mapping_tree()
            return
        lookup = {canonical_header(h): h for h in self.display_bm_headers}
        for target in DISPLAY_FORMAT_HEADERS:
            if target in DISPLAY_DEFAULT_MODES:
                continue
            key = canonical_header(target)
            if key in lookup:
                self.display_mappings[target] = {"mode": "3D BM 열 복사", "source": lookup[key]}
        # common aliases
        aliases = {
            "Drawing Number": ["DRAWING NUMBER", "Drawing Number", "DWG Number"],
            "CWP No.": ["CWP No.", "ISO_DWG_ID", "ISO DWG ID"],
            "Class": ["CLASS", "Class", "Spec"],
            "Item": ["ITEM", "Item", "Symbol"],
            "Qty": ["QTY", "Qty", "Quantity"],
            "Main": ["MAIN", "Main"],
            "Sub": ["SUB", "Sub"],
            "Insulation Symbol": ["INSULATION SYMBOL", "Insulation Symbol"],
            "Remark": ["REMARK", "Remark"],
        }
        for target, candidates in aliases.items():
            for candidate in candidates:
                found = lookup.get(canonical_header(candidate))
                if found:
                    self.display_mappings[target] = {"mode": "3D BM 열 복사", "source": found}
                    break
        self.refresh_display_mapping_tree()

    def on_display_mapping_select(self, _event=None):
        selected = self.display_map_tree.selection()
        if not selected:
            return
        header = selected[0]
        mapping = self.display_mappings[header]
        self.display_selected_header.set(header)
        self.display_mapping_mode.set(mapping["mode"])
        source = mapping.get("source", "")
        if mapping["mode"] in ["빈칸", "자동번호", "특수규칙"]:
            source = ""
        self.display_mapping_source.set(source)
        self.on_display_mapping_mode_change()

    def on_display_mapping_mode_change(self, _event=None):
        mode = self.display_mapping_mode.get()
        if mode in ["빈칸", "자동번호", "특수규칙"]:
            self.display_mapping_source.set("")
            self.display_source_combo.configure(state="disabled")
        else:
            self.display_source_combo.configure(state="normal")
        self._update_display_mode_help()

    def _update_display_mode_help(self):
        """Explain the selected input mode below the mapping editor."""
        if not hasattr(self, "display_mode_help"):
            return
        mode = self.display_mapping_mode.get()
        header = self.display_selected_header.get()
        text = f"[{mode}] " + MAPPING_MODE_HELP.get(mode, "")
        if mode == "특수규칙" and header and header not in SPECIAL_RULE_HEADERS:
            text += f"\n※ '{header}' 열에는 특수규칙이 없습니다. Part 또는 Unit Size 열에서만 선택하세요."
        if mode == "PCWBS 파일 매핑" and header and not header.startswith("Category"):
            text += f"\n※ '{header}' 열은 PCWBS 파일 매핑 대상이 아닙니다."
        self.display_mode_help.set(text)

    def apply_display_mapping(self):
        header = self.display_selected_header.get()
        if not header:
            return
        mode = self.display_mapping_mode.get()
        source = self.display_mapping_source.get()
        if mode == "특수규칙" and header not in SPECIAL_RULE_HEADERS:
            messagebox.showwarning(
                "안내",
                f"특수규칙은 {', '.join(SPECIAL_RULE_HEADERS)} 열에서만 사용할 수 있습니다.\n"
                f"'{header}' 열은 다른 입력 방식을 선택하세요.",
            )
            return
        if mode == "PCWBS 파일 매핑" and not header.startswith("Category"):
            messagebox.showwarning(
                "안내",
                "PCWBS 파일 매핑은 Category1~10 열에서만 사용할 수 있습니다.",
            )
            return
        if mode in ["빈칸", "자동번호", "특수규칙"]:
            source = ""
        self.display_mappings[header] = {
            "mode": mode,
            "source": source,
        }
        self.refresh_display_mapping_tree()

    def add_part_rule(self):
        column = self.part_rule_column.get().strip()
        value = self.part_rule_value.get()
        if not column:
            messagebox.showwarning("경고", "조건 열을 선택하세요.")
            return
        rule = {
            "column": column,
            "operator": self.part_rule_operator.get(),
            "value": value,
            "result": self.part_rule_result.get(),
        }
        self.part_rules.append(rule)
        self.part_rule_list.insert("end", f"{column} | {rule['operator']} | {value} → {rule['result']}")

    def delete_part_rule(self):
        selection = self.part_rule_list.curselection()
        if not selection:
            return
        index = selection[0]
        self.part_rules.pop(index)
        self.part_rule_list.delete(index)

    def _display_config(self):
        return {
            "mappings": self.display_mappings,
            "part_rules": self.part_rules,
            "part_default": self.display_part_default.get(),
            "unit_default": self.display_unit_default.get(),
            "unit_item_column": self.display_unit_item_column.get(),
            "unit_operator": self.display_unit_operator.get(),
            "unit_values": self.display_unit_values.get(),
            "cd_site": self.display_cd_site.get(),
            "times": self.display_times.get(),
            "mtrl_group": self.display_mtrl_group.get(),
            "filename": self.display_filename.get(),
            "main_sub_blank": self.display_main_sub_blank.get(),
            "strip_main_sub_units": self.display_strip_main_sub_units.get(),
            "use_pcwbs_mapping": self._display_requires_pcwbs_mapping(),
        }

    def save_display_mapping(self):
        path = filedialog.asksaveasfilename(
            initialdir=self.project_root / "03_reference",
            defaultextension=".json",
            filetypes=[("JSON", "*.json")],
        )
        if not path:
            return
        Path(path).write_text(json.dumps(self._display_config(), ensure_ascii=False, indent=2), encoding="utf-8")
        messagebox.showinfo("완료", "Display Format 매핑을 저장했습니다.")

    def load_display_mapping(self):
        path = filedialog.askopenfilename(
            initialdir=self.project_root / "03_reference",
            filetypes=[("JSON", "*.json")],
        )
        if not path:
            return
        config = json.loads(Path(path).read_text(encoding="utf-8"))
        self.display_mappings = config.get("mappings", self.display_mappings)
        for mapping in self.display_mappings.values():
            mapping["mode"] = LEGACY_MAPPING_MODES.get(mapping.get("mode"), mapping.get("mode"))
        if config.get("use_pcwbs_mapping", False):
            for number in range(1, 6):
                header = f"Category{number}"
                current = self.display_mappings.get(header, {})
                if current.get("mode") == "3D BM 열 복사":
                    self.display_mappings[header] = {
                        "mode": "PCWBS 파일 매핑",
                        "source": header,
                    }
        self.part_rules = config.get("part_rules", [])
        self.display_part_default.set(config.get("part_default", "PPA"))
        self.display_unit_default.set(config.get("unit_default", "mm"))
        self.display_unit_item_column.set(config.get("unit_item_column", "Item"))
        self.display_unit_operator.set(config.get("unit_operator", "목록 중 하나 포함"))
        self.display_unit_values.set(config.get("unit_values", ""))
        self.display_cd_site.set(config.get("cd_site", ""))
        self.display_times.set(config.get("times", ""))
        self.display_mtrl_group.set(config.get("mtrl_group", "RUSSIA"))
        self.display_filename.set(config.get("filename", "Display_Format_01.xlsx"))
        self.display_main_sub_blank.set(config.get("main_sub_blank", True))
        self.display_strip_main_sub_units.set(
            config.get("strip_main_sub_units", True)
        )
        self.display_use_pcwbs_mapping.set(self._display_requires_pcwbs_mapping())
        self.part_rule_list.delete(0, "end")
        for rule in self.part_rules:
            self.part_rule_list.insert("end", f"{rule['column']} | {rule['operator']} | {rule['value']} → {rule['result']}")
        self.refresh_display_mapping_tree()
        messagebox.showinfo("완료", "Display Format 매핑을 불러왔습니다.")

    def _iter_excel_dict_rows(self, path, sheet_name, header_row):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name]
        iterator = ws.iter_rows(min_row=header_row, values_only=True)
        headers = make_unique_headers(next(iterator))
        for excel_row, values in enumerate(iterator, start=header_row + 1):
            row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
            if not is_effectively_empty_row(values):
                yield excel_row, row
        wb.close()

    def _part_value_for_row(self, row):
        for rule in self.part_rules:
            if compare_value(row.get(rule["column"]), rule["operator"], rule["value"]):
                return rule["result"]
        return self.display_part_default.get()

    def _unit_value_for_row(self, row):
        column = self.display_unit_item_column.get()
        if column and compare_value(
            row.get(column),
            self.display_unit_operator.get(),
            self.display_unit_values.get(),
        ):
            return "U x mm"
        return self.display_unit_default.get()

    def _pcwbs_mapped_category_targets(self):
        """Return Display Category headers configured for PCWBS file mapping."""
        targets = []
        for number in range(1, 11):
            header = f"Category{number}"
            mapping = self.display_mappings.get(header, {})
            if mapping.get("mode") == "PCWBS 파일 매핑":
                source = clean_text(mapping.get("source")) or header
                targets.append((header, source))
        return targets

    def _display_requires_pcwbs_mapping(self):
        return bool(self._pcwbs_mapped_category_targets())

    def _validate_display_key_settings(self):
        """Validate tab 3 key settings for Display Category mapping."""
        if self._use_mapping_table():
            self._prepare_key_mapping(
                self.display_bm_headers,
                self.pcwbs_ref_headers,
                bm_label="4번 탭 입력 3D BM",
            )
            return

        pcwbs_settings = self._enabled_component_settings(True)
        if not pcwbs_settings:
            raise ValueError("3번 탭에서 PCWBS 조합키 구성요소를 하나 이상 선택하세요.")
        for column, _mode, _pad in pcwbs_settings:
            if column not in self.pcwbs_ref_headers:
                raise ValueError(f"PCWBS 조합 열을 확인하세요: {column}")

        mode = self.pcwbs_bm_key_mode.get()
        if mode in ["열 조합", "두 방식 교차검증"]:
            bm_settings = self._enabled_component_settings(False)
            if not bm_settings:
                raise ValueError("3번 탭에서 3D BM 열 조합 구성요소를 하나 이상 선택하세요.")
            for column, _extract, _pad in bm_settings:
                if column not in self.display_bm_headers:
                    raise ValueError(
                        f"Display 입력 BM에서 조합 열을 찾을 수 없습니다: {column}"
                    )

        if mode in ["MID 텍스트 추출", "두 방식 교차검증"]:
            column = self.pcwbs_bm_iso.get()
            if column not in self.display_bm_headers:
                raise ValueError(
                    f"Display 입력 BM에서 MID 대상 열을 찾을 수 없습니다: {column}"
                )

    def _load_pcwbs_category_map(self):
        """Create key -> selected PCWBS Category values using tab 3 key settings."""
        targets = self._pcwbs_mapped_category_targets()
        if not targets:
            return {}, []

        ref_path = Path(self.pcwbs_ref_file.get().strip())
        if not ref_path.is_file():
            raise ValueError("Category 매핑에 사용할 PCWBS 기준 파일을 선택하세요.")

        self.pcwbs_ref_headers = self._read_headers_only(
            str(ref_path),
            self.pcwbs_ref_sheet.get().strip(),
            self.pcwbs_ref_header_row.get().strip(),
        )
        self._validate_display_key_settings()

        canonical_map = {canonical_header(h): h for h in self.pcwbs_ref_headers}
        resolved_targets = []
        missing_columns = []
        for display_header, requested_source in targets:
            found = canonical_map.get(canonical_header(requested_source))
            if not found:
                found = canonical_map.get(canonical_header(display_header))
            if not found:
                missing_columns.append(f"{display_header} ← {requested_source}")
            else:
                resolved_targets.append((display_header, found))

        if missing_columns:
            raise ValueError(
                "PCWBS 파일 매핑으로 설정했지만 기준 파일에 열이 없습니다:\n"
                + "\n".join(missing_columns)
            )

        _headers, ref_rows = self._read_reference_rows_fast(
            ref_path,
            self.pcwbs_ref_sheet.get(),
            int(self.pcwbs_ref_header_row.get() or "1"),
        )
        settings = self._enabled_component_settings(True)
        category_map = {}
        issues = []

        for row in ref_rows:
            key, error = self._pcwbs_row_key(row, settings)
            if error or not key:
                issues.append({
                    "Issue Type": "PCWBS Key Error",
                    "Excel Row": row.get("__Excel_Row__"),
                    "Generated Key": key,
                    "Details": error or "키가 비어 있음",
                })
                continue

            mapped_values = {
                display_header: row.get(source_column)
                for display_header, source_column in resolved_targets
            }
            if key not in category_map:
                category_map[key] = mapped_values
            elif category_map[key] != mapped_values:
                issues.append({
                    "Issue Type": "PCWBS Conflicting Duplicate Key",
                    "Excel Row": row.get("__Excel_Row__"),
                    "Generated Key": key,
                    "Details": (
                        "동일 키에 서로 다른 PCWBS Category 값이 존재함: "
                        f"{category_map[key]} / {mapped_values}"
                    ),
                })

        if self._use_mapping_table():
            issues = self._active_key_mapping["issues"] + issues
        return category_map, issues

    def _candidate_bm_keys_for_mapping(self, row):
        if self._use_mapping_table():
            keys, _bm_value, error = self._bm_row_mapped_keys(row)
            return keys, error

        combination_key, mid_key, _legacy_key, error = self._build_bm_keys(row)
        mode = self.pcwbs_bm_key_mode.get()
        if mode == "열 조합":
            candidates = [combination_key]
        elif mode == "MID 텍스트 추출":
            candidates = [mid_key]
        else:
            candidates = [combination_key, mid_key]

        unique_candidates = []
        for key in candidates:
            if key and key not in unique_candidates:
                unique_candidates.append(key)
        return unique_candidates, error

    def _resolve_pcwbs_categories(self, row, category_map):
        """Return selected mapped Category values, selected key, and an issue message."""
        candidates, key_error = self._candidate_bm_keys_for_mapping(row)
        matched = [(key, category_map[key]) for key in candidates if key in category_map]

        if not matched:
            issue = key_error or (
                "PCWBS에 일치하는 조합키가 없음"
                + (f" ({', '.join(candidates)})" if candidates else "")
            )
            return None, "", issue

        first_key, first_values = matched[0]
        conflicting = [
            (key, values)
            for key, values in matched[1:]
            if values != first_values
        ]
        if conflicting:
            return first_values, first_key, (
                "복수 BM 키가 서로 다른 PCWBS Category 값에 매핑됨: "
                + ", ".join(key for key, _values in matched)
            )
        return first_values, first_key, ""

    @staticmethod
    def _strip_size_unit(value):
        """Remove common unit suffixes from Main/Sub values."""
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return value
        value_text = clean_text(value)
        if not value_text:
            return None
        return re.sub(
            r'(?i)\s*(?:mm|millimeter(?:s)?|cm|meter(?:s)?|metre(?:s)?|inch(?:es)?|in\.?|")\s*$',
            "",
            value_text,
        ).strip()

    def _build_display_row(
        self,
        row,
        number,
        pcwbs_category_map=None,
        mapping_issues=None,
        excel_row=None,
    ):
        result = {}
        for header in DISPLAY_FORMAT_HEADERS:
            mapping = self.display_mappings.get(header, {"mode": "빈칸", "source": ""})
            mode = mapping["mode"]
            source = mapping["source"]
            if mode == "자동번호":
                value = number
            elif mode == "특수규칙":
                if header == "Part":
                    value = self._part_value_for_row(row)
                elif header == "Unit Size":
                    value = self._unit_value_for_row(row)
                else:
                    value = None
            elif mode == "3D BM 열 복사":
                value = row.get(source)
            elif mode == "PCWBS 파일 매핑":
                value = None
            elif mode == "고정값":
                value = source
            else:
                value = None
            result[header] = value

        if self._display_requires_pcwbs_mapping() and pcwbs_category_map is not None:
            mapped_categories, used_key, issue = self._resolve_pcwbs_categories(row, pcwbs_category_map)
            if mapped_categories is not None:
                for category_header, value in mapped_categories.items():
                    result[category_header] = value
            if issue and mapping_issues is not None:
                mapping_issues.append([
                    excel_row,
                    used_key,
                    issue,
                    clean_text(row.get(self.pcwbs_bm_subtitle.get())),
                    clean_text(row.get(self.pcwbs_bm_cia.get())),
                    clean_text(row.get(self.pcwbs_bm_iso.get())),
                ])

        result["CD_SITE"] = self.display_cd_site.get()
        result["PrpsProj"] = "2"
        result["Times"] = self.display_times.get()
        result["Purpose"] = "PO"
        result["Mtrl Group"] = self.display_mtrl_group.get()
        result["FileName"] = self.display_filename.get()

        if self.display_strip_main_sub_units.get():
            result["Main"] = self._strip_size_unit(result.get("Main"))
            result["Sub"] = self._strip_size_unit(result.get("Sub"))

        if self.display_main_sub_blank.get():
            if (
                clean_text(result.get("Main"))
                and clean_text(result.get("Main"))
                == clean_text(result.get("Sub"))
            ):
                result["Sub"] = None
        return result

    def _calculate_display_preview(self, limit=None):
        path = Path(self.display_bm_file.get().strip())
        if not path.exists():
            raise ValueError("3D BM 파일을 찾을 수 없습니다.")
        header_row = int(self.display_bm_header_row.get() or "1")
        count = 0
        missing_sources = set()
        for header, mapping in self.display_mappings.items():
            if mapping["mode"] == "3D BM 열 복사" and mapping["source"] not in self.display_bm_headers:
                missing_sources.add(mapping["source"])

        category_map = {}
        pcwbs_reference_issues = []
        if self._display_requires_pcwbs_mapping():
            category_map, pcwbs_reference_issues = self._load_pcwbs_category_map()

        matched = 0
        unmatched = 0
        for _excel_row, row in self._iter_excel_dict_rows(
            path,
            self.display_bm_sheet.get(),
            header_row,
        ):
            count += 1
            if self._display_requires_pcwbs_mapping():
                categories, _key, _issue = self._resolve_pcwbs_categories(
                    row,
                    category_map,
                )
                if categories is None:
                    unmatched += 1
                else:
                    matched += 1
            if count % 5000 == 0:
                self.set_status(
                    "Display Format 미리보기",
                    f"현재 작업: {count:,}행 점검",
                    min(95, 10 + (count % 85000) / 1000),
                )
                self.log(f"Display Format 미리보기 진행: {count:,}행")
            if limit and count >= limit:
                break
        return (
            count,
            sorted(x for x in missing_sources if x),
            matched,
            unmatched,
            len(category_map),
            len(pcwbs_reference_issues),
        )

    def preview_display_format(self):
        self.set_running(True)
        self.set_status(
            "Display Format 미리보기",
            "현재 작업: 입력 파일 확인",
            0,
        )
        try:
            count, missing, matched, unmatched, pcwbs_keys, pcwbs_ref_issues = self._calculate_display_preview()
            message = (
                f"변환 대상 데이터: {count:,}행\n"
                f"Display Format 열: {len(DISPLAY_FORMAT_HEADERS)}개\n"
                f"존재하지 않는 BM 원본 열: {len(missing)}개"
            )
            if self._display_requires_pcwbs_mapping():
                message += (
                    f"\nPCWBS 고유 조합키: {pcwbs_keys:,}개"
                    f"\nPCWBS Category 매핑 성공: {matched:,}행"
                    f"\nPCWBS Category 매핑 불일치: {unmatched:,}행"
                    f"\nPCWBS 기준 키 이슈: {pcwbs_ref_issues:,}건"
                )
            if missing:
                message += "\n\n누락 원본 열:\n" + "\n".join(missing[:20])
            self.root.after(0, lambda: messagebox.showinfo("Display Format 미리보기", message))
            self.set_status("완료", "현재 작업: 없음", 100)
        except Exception as exc:
            self.log(traceback.format_exc())
            self.root.after(0, lambda exc=exc: messagebox.showerror("오류", str(exc)))
        finally:
            self.set_running(False)

    def generate_display_format(self):
        self.set_running(True)
        try:
            input_path = Path(self.display_bm_file.get().strip())
            output_path = Path(self.display_output_file.get().strip())
            output_path.parent.mkdir(parents=True, exist_ok=True)
            header_row = int(self.display_bm_header_row.get() or "1")

            pcwbs_category_map = {}
            pcwbs_reference_issues = []
            if self._display_requires_pcwbs_mapping():
                pcwbs_category_map, pcwbs_reference_issues = self._load_pcwbs_category_map()

            out_wb = Workbook(write_only=True)
            out_ws = out_wb.create_sheet("Display_Format")
            out_ws.append(DISPLAY_FORMAT_HEADERS)
            mapping_issues = []

            count = 0
            count_wb = load_workbook(input_path, read_only=True, data_only=True)
            count_ws = count_wb[self.display_bm_sheet.get()]
            estimated_total = max(1, count_ws.max_row - header_row)
            count_wb.close()

            self.set_status(
                "Display Format 생성",
                "현재 작업: 데이터 변환 시작",
                0,
            )
            for excel_row, row in self._iter_excel_dict_rows(
                input_path,
                self.display_bm_sheet.get(),
                header_row,
            ):
                count += 1
                display_row = self._build_display_row(
                    row,
                    count,
                    pcwbs_category_map=pcwbs_category_map,
                    mapping_issues=mapping_issues,
                    excel_row=excel_row,
                )
                out_ws.append([display_row.get(header) for header in DISPLAY_FORMAT_HEADERS])
                if count % 2000 == 0:
                    percent = min(92, count / estimated_total * 92)
                    self.set_status(
                        "Display Format 생성 중",
                        f"현재 작업: {count:,} / {estimated_total:,}행",
                        percent,
                    )
                    self.log(
                        f"Display Format 생성 진행: {count:,} / "
                        f"{estimated_total:,}행"
                    )

            if self._display_requires_pcwbs_mapping():
                issue_ws = out_wb.create_sheet("PCWBS_Mapping_Issues")
                issue_ws.append([
                    "Source Excel Row",
                    "Matched Key",
                    "Issue",
                    "BM Subtitle",
                    "BM CIA",
                    "BM ISO DWG",
                ])
                for issue in mapping_issues:
                    issue_ws.append(issue)

                ref_issue_ws = out_wb.create_sheet("PCWBS_Reference_Issues")
                ref_issue_ws.append(["Issue Type", "Excel Row", "Generated Key", "Details"])
                for issue in pcwbs_reference_issues:
                    ref_issue_ws.append([
                        issue.get("Issue Type"),
                        issue.get("Excel Row"),
                        issue.get("Generated Key"),
                        issue.get("Details"),
                    ])

            out_wb.save(output_path)

            # 자동 연계: 4번 결과를 5번 Temperature & Painting 입력으로 전달
            self.paint_display_file.set(str(output_path))
            self.paint_display_sheet.set("Display_Format")

            self.set_status("완료", "현재 작업: 없음", 100)
            self.log(
                f"Display Format 생성 완료: {output_path} / {count:,}행 / "
                f"PCWBS 매핑 이슈 {len(mapping_issues):,}건"
            )
            self.log("자동 연계: 5번 탭 Display Format 입력 파일로 설정")
            self.root.after(
                0,
                lambda: messagebox.showinfo(
                    "완료",
                    f"{count:,}행 생성 완료\n"
                    f"PCWBS 매핑 이슈: {len(mapping_issues):,}건\n\n{output_path}",
                ),
            )
        except Exception as exc:
            self.log(traceback.format_exc())
            self.root.after(0, lambda exc=exc: messagebox.showerror("오류", str(exc)))
        finally:
            self.set_running(False)

    # ------------------------------------------------------------------ Temperature & Painting
    def _build_paint_tab(self):
        output_dir = self.project_root / "04_output"

        self.paint_display_file = tk.StringVar(
            value=str(output_dir / "Display_Format_Base.xlsx")
        )
        self.paint_display_sheet = tk.StringVar(value="Display_Format")

        self.paint_line_file = tk.StringVar(value="")
        self.paint_line_sheet = tk.StringVar(value="")
        self.paint_line_header_start = tk.StringVar(value="2")
        self.paint_line_header_end = tk.StringVar(value="4")
        self.paint_line_data_start = tk.StringVar(value="6")

        self.paint_drawing_col = tk.StringVar(value="Drawing Number")
        self.paint_class_col = tk.StringVar(value="Class")
        self.paint_insulation_col = tk.StringVar(value="Insulation Symbol")

        self.paint_line_no_col = tk.StringVar(value="")
        self.paint_operating_temp_col = tk.StringVar(value="")
        self.paint_max_operating_temp_col = tk.StringVar(value="")
        self.paint_line_match_mode = tk.StringVar(value=LINE_MATCH_STRIP_SHEET)
        self.paint_sheet_separator = tk.StringVar(value="-")
        self.paint_ins_temp_source = tk.StringVar(value=INS_TEMP_SOURCE_LINE)
        self.paint_ins_rule_file = tk.StringVar(value="")
        self.paint_ins_rule_sheet = tk.StringVar(value="")
        self.paint_symbol_source = tk.StringVar(value=PAINT_SOURCE_TABLE)
        self.paint_rule_file = tk.StringVar(value="")
        self.paint_rule_sheet = tk.StringVar(value="")
        self.paint_selected_header_text = tk.StringVar(
            value="선택한 열의 전체명이 여기에 표시됩니다."
        )
        self.paint_ambient_temp = tk.StringVar(value="35")

        self.paint_table_file = tk.StringVar(value="")
        self.paint_table_sheet = tk.StringVar(value="")
        self.paint_table_header = tk.StringVar(value="1")

        self.paint_output_file = tk.StringVar(
            value=str(output_dir / "Display_Format_TempPaint.xlsx")
        )
        self.paint_insulation_values = []

        file_frame = tk.LabelFrame(
            self.paint_tab,
            text="입력 및 결과 파일",
            padx=8,
            pady=7,
        )
        file_frame.pack(fill="x", padx=10, pady=(10, 4))
        file_frame.columnconfigure(1, weight=1)

        self._file_row(
            file_frame,
            0,
            "Display Format 파일",
            self.paint_display_file,
            self.choose_paint_display_file,
        )
        tk.Label(file_frame, text="Display 시트").grid(
            row=1, column=0, sticky="w", pady=3
        )
        self.paint_display_sheet_combo = ttk.Combobox(
            file_frame,
            textvariable=self.paint_display_sheet,
            state="readonly",
            width=28,
        )
        self.paint_display_sheet_combo.grid(
            row=1, column=1, sticky="w", padx=6, pady=3
        )
        tk.Button(
            file_frame,
            text="Display 열 불러오기",
            command=self.load_paint_display_columns,
            width=17,
        ).grid(row=1, column=4, padx=5)

        self._file_row(
            file_frame,
            2,
            "Line List",
            self.paint_line_file,
            self.choose_paint_line_file,
        )
        tk.Label(file_frame, text="Line List 시트").grid(
            row=3, column=0, sticky="w", pady=3
        )
        self.paint_line_sheet_combo = ttk.Combobox(
            file_frame,
            textvariable=self.paint_line_sheet,
            state="readonly",
            width=28,
        )
        self.paint_line_sheet_combo.grid(
            row=3, column=1, sticky="w", padx=6, pady=3
        )

        line_header_frame = tk.Frame(file_frame)
        line_header_frame.grid(
            row=3,
            column=2,
            columnspan=2,
            sticky="w",
            padx=5,
        )
        tk.Label(line_header_frame, text="헤더 시작").pack(side="left")
        tk.Entry(
            line_header_frame,
            textvariable=self.paint_line_header_start,
            width=5,
        ).pack(side="left", padx=(3, 8))
        tk.Label(line_header_frame, text="헤더 종료").pack(side="left")
        tk.Entry(
            line_header_frame,
            textvariable=self.paint_line_header_end,
            width=5,
        ).pack(side="left", padx=(3, 8))
        tk.Label(line_header_frame, text="데이터 시작").pack(side="left")
        tk.Entry(
            line_header_frame,
            textvariable=self.paint_line_data_start,
            width=5,
        ).pack(side="left", padx=3)

        tk.Button(
            file_frame,
            text="Line List 열 불러오기",
            command=self.load_paint_line_columns,
            width=17,
        ).grid(row=3, column=4, padx=5)

        self._file_row(
            file_frame,
            4,
            "Painting Code Table",
            self.paint_table_file,
            self.choose_paint_table_file,
        )
        tk.Label(file_frame, text="Painting Table 시트").grid(
            row=5, column=0, sticky="w", pady=3
        )
        self.paint_table_sheet_combo = ttk.Combobox(
            file_frame,
            textvariable=self.paint_table_sheet,
            state="readonly",
            width=28,
        )
        self.paint_table_sheet_combo.grid(
            row=5, column=1, sticky="w", padx=6, pady=3
        )
        tk.Label(file_frame, text="헤더 행").grid(
            row=5, column=2, sticky="e"
        )
        tk.Entry(
            file_frame,
            textvariable=self.paint_table_header,
            width=6,
        ).grid(row=5, column=3, sticky="w", padx=5)
        tk.Button(
            file_frame,
            text="Painting Code 예시",
            command=self.show_painting_code_example,
            width=17,
        ).grid(row=5, column=4, padx=5)

        self._file_row(
            file_frame,
            6,
            "결과 파일",
            self.paint_output_file,
            self.choose_paint_output_file,
            save=True,
        )

        map_frame = tk.LabelFrame(
            self.paint_tab,
            text="열 매핑 및 온도 정규화",
            padx=8,
            pady=5,
        )
        map_frame.pack(fill="x", padx=10, pady=4)
        for column_no in range(3):
            map_frame.grid_columnconfigure(column_no, weight=1, uniform="paint_map")

        self.paint_display_combos = []
        self.paint_line_combos = []

        def add_combo_row(parent, row_no, label, variable, combo_list, width=30):
            tk.Label(parent, text=label).grid(
                row=row_no, column=0, sticky="e", padx=(0, 4), pady=2
            )
            combo = ttk.Combobox(
                parent,
                textvariable=variable,
                state="readonly",
                width=width,
            )
            combo.grid(row=row_no, column=1, sticky="ew", padx=(0, 4), pady=2)
            combo.bind(
                "<<ComboboxSelected>>",
                lambda _event, var=variable: self.show_paint_header_full_name(
                    var.get()
                ),
            )
            combo_list.append(combo)
            return combo

        # ---- 공통: Display Drawing Number ↔ Line List Line No. 매칭
        common_frame = tk.LabelFrame(
            map_frame,
            text="공통 · Line List 매칭",
            padx=6,
            pady=4,
        )
        common_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 4))
        common_frame.grid_columnconfigure(1, weight=1)
        add_combo_row(
            common_frame, 0, "Display Drawing Number",
            self.paint_drawing_col, self.paint_display_combos,
        )
        add_combo_row(
            common_frame, 1, "Line List Line No.",
            self.paint_line_no_col, self.paint_line_combos,
        )
        tk.Label(common_frame, text="Line No. 매칭 방식").grid(
            row=2, column=0, sticky="e", padx=(0, 4), pady=2
        )
        self.paint_line_match_combo = ttk.Combobox(
            common_frame,
            textvariable=self.paint_line_match_mode,
            values=LINE_MATCH_MODES,
            state="readonly",
            width=30,
        )
        self.paint_line_match_combo.grid(row=2, column=1, sticky="ew", pady=2)
        self.paint_line_match_combo.bind(
            "<<ComboboxSelected>>",
            lambda _event: self.show_paint_header_full_name(
                self.paint_line_match_mode.get()
            ),
        )
        sheet_frame = tk.Frame(common_frame)
        sheet_frame.grid(row=3, column=0, columnspan=2, sticky="w", pady=2)
        tk.Label(sheet_frame, text="Sh't No. 구분자").pack(side="left")
        tk.Entry(
            sheet_frame,
            textvariable=self.paint_sheet_separator,
            width=4,
        ).pack(side="left", padx=4)
        tk.Label(
            sheet_frame,
            text="예: 304-PW-0051-01 → 304-PW-0051",
            fg="#555555",
        ).pack(side="left", padx=4)

        # ---- ① Line List Insulation Temp → Display Format 'Insulation Temp' 열
        insulation_temp_frame = tk.LabelFrame(
            map_frame,
            text="① Line List Insulation Temp → Display Format Insulation Temp 열",
            padx=6,
            pady=4,
        )
        insulation_temp_frame.grid(row=0, column=1, sticky="nsew", padx=4)
        insulation_temp_frame.grid_columnconfigure(1, weight=1)
        self.paint_ins_frame = insulation_temp_frame
        tk.Label(insulation_temp_frame, text="온도 입력 방식").grid(
            row=0, column=0, sticky="e", padx=(0, 4), pady=2
        )
        self.paint_ins_source_combo = ttk.Combobox(
            insulation_temp_frame,
            textvariable=self.paint_ins_temp_source,
            values=INS_TEMP_SOURCES,
            state="readonly",
            width=30,
        )
        self.paint_ins_source_combo.grid(row=0, column=1, sticky="ew", pady=2)
        self.paint_ins_source_combo.bind(
            "<<ComboboxSelected>>",
            lambda _event: self.update_insulation_temp_source_ui(),
        )
        # row 1: Line List 방식
        add_combo_row(
            insulation_temp_frame, 1, "Insulation 적용 Temp.",
            self.paint_operating_temp_col, self.paint_line_combos,
        )
        # row 2~3: 규칙 파일 방식
        self.paint_ins_rule_widgets = self._build_rule_file_rows(
            insulation_temp_frame,
            2,
            "온도 규칙 파일",
            self.paint_ins_rule_file,
            self.paint_ins_rule_sheet,
            self.show_insulation_temp_rule_example,
        )
        tk.Label(insulation_temp_frame, text="AMB 대체 온도").grid(
            row=4, column=0, sticky="e", padx=(0, 4), pady=2
        )
        tk.Entry(
            insulation_temp_frame,
            textvariable=self.paint_ambient_temp,
            width=8,
        ).grid(row=4, column=1, sticky="w", pady=2)
        self.paint_ins_help = tk.StringVar(value="")
        help_label = tk.Label(
            insulation_temp_frame,
            textvariable=self.paint_ins_help,
            fg="#555555",
            anchor="w",
            justify="left",
            wraplength=380,
        )
        help_label.grid(row=5, column=0, columnspan=2, sticky="w", pady=(2, 0))
        insulation_temp_frame.bind(
            "<Configure>",
            lambda event, label=help_label: label.configure(
                wraplength=max(200, event.width - 20)
            ),
            add="+",
        )

        # ---- ② Paint Code → Display Format 'Paint Symbol' 열
        paint_code_frame = tk.LabelFrame(
            map_frame,
            text="② Paint Code → Display Format Paint Symbol 열",
            padx=6,
            pady=4,
        )
        paint_code_frame.grid(row=0, column=2, sticky="nsew", padx=(4, 0))
        paint_code_frame.grid_columnconfigure(1, weight=1)
        self.paint_code_frame = paint_code_frame
        tk.Label(paint_code_frame, text="Paint 입력 방식").grid(
            row=0, column=0, sticky="e", padx=(0, 4), pady=2
        )
        self.paint_source_combo = ttk.Combobox(
            paint_code_frame,
            textvariable=self.paint_symbol_source,
            values=PAINT_SOURCES,
            state="readonly",
            width=30,
        )
        self.paint_source_combo.grid(row=0, column=1, sticky="ew", pady=2)
        self.paint_source_combo.bind(
            "<<ComboboxSelected>>",
            lambda _event: self.update_paint_source_ui(),
        )
        # row 1~3: Painting Code Table 계산 방식
        add_combo_row(
            paint_code_frame, 1, "Display Class",
            self.paint_class_col, self.paint_display_combos,
        )
        add_combo_row(
            paint_code_frame, 2, "Display Insulation Symbol",
            self.paint_insulation_col, self.paint_display_combos,
        )
        add_combo_row(
            paint_code_frame, 3, "Painting 적용 Temp.",
            self.paint_max_operating_temp_col, self.paint_line_combos,
        )
        # row 4~5: 규칙 파일 방식
        self.paint_rule_widgets = self._build_rule_file_rows(
            paint_code_frame,
            4,
            "Paint 규칙 파일",
            self.paint_rule_file,
            self.paint_rule_sheet,
            self.show_paint_rule_example,
        )
        self.paint_code_help = tk.StringVar(value="")
        help_label = tk.Label(
            paint_code_frame,
            textvariable=self.paint_code_help,
            fg="#555555",
            anchor="w",
            justify="left",
            wraplength=380,
        )
        help_label.grid(row=6, column=0, columnspan=2, sticky="w", pady=(2, 0))
        paint_code_frame.bind(
            "<Configure>",
            lambda event, label=help_label: label.configure(
                wraplength=max(200, event.width - 20)
            ),
            add="+",
        )

        full_name_frame = tk.LabelFrame(
            map_frame,
            text="선택 열 전체명",
            padx=6,
            pady=2,
        )
        full_name_frame.grid(
            row=1,
            column=0,
            columnspan=3,
            sticky="ew",
            pady=(5, 2),
        )
        tk.Label(
            full_name_frame,
            textvariable=self.paint_selected_header_text,
            anchor="w",
            justify="left",
            wraplength=1220,
            fg="#1f4e79",
        ).pack(fill="x")

        tk.Label(
            map_frame,
            text=(
                "Insulation 적용 Temp. / Painting 적용 Temp.에는 프로젝트 기준에 따라 Line List의 "
                "Operating / Max. Operating / Design Temp. 중 원하는 열을 선택하세요. "
                "복수 숫자는 가장 큰 값을 적용합니다."
            ),
            fg="#555555",
            anchor="w",
            justify="left",
            wraplength=1220,
        ).grid(
            row=2,
            column=0,
            columnspan=3,
            sticky="w",
            pady=(2, 0),
        )
        self.update_insulation_temp_source_ui()

        rule_body = tk.Frame(self.paint_tab)
        rule_body.pack(fill="both", expand=True, padx=10, pady=4)
        rule_body.grid_rowconfigure(0, weight=1)
        rule_body.grid_columnconfigure(0, weight=1)
        rule_body.grid_columnconfigure(1, weight=1)

        material_frame = tk.LabelFrame(
            rule_body,
            text="② Paint Code · Class → Material Group 규칙",
            padx=6,
            pady=6,
        )
        material_frame.grid(row=0, column=0, sticky="nsew")
        material_frame.grid_rowconfigure(1, weight=1)

        self.material_rule_operator = tk.StringVar(value="포함")
        self.material_rule_value = tk.StringVar(value="")
        self.material_rule_result = tk.StringVar(value="CS")

        ttk.Combobox(
            material_frame,
            textvariable=self.material_rule_operator,
            values=OPERATORS,
            state="readonly",
            width=18,
        ).grid(row=0, column=0, padx=3)
        tk.Entry(
            material_frame,
            textvariable=self.material_rule_value,
            width=25,
        ).grid(row=0, column=1, padx=3)
        ttk.Combobox(
            material_frame,
            textvariable=self.material_rule_result,
            values=["CS", "SS", "HDPE", "GALVA"],
            width=12,
        ).grid(row=0, column=2, padx=3)
        tk.Button(
            material_frame,
            text="추가",
            command=self.add_material_rule,
        ).grid(row=0, column=3, padx=3)

        self.material_rule_list = tk.Listbox(
            material_frame,
            height=8,
        )
        self.material_rule_list.grid(
            row=1,
            column=0,
            columnspan=4,
            sticky="nsew",
            pady=4,
        )
        tk.Button(
            material_frame,
            text="선택 삭제",
            command=self.delete_material_rule,
        ).grid(row=2, column=0, columnspan=4)

        insulation_frame = tk.LabelFrame(
            rule_body,
            text="② Paint Code · Insulation Symbol → Paint Suffix 규칙",
            padx=6,
            pady=6,
        )
        insulation_frame.grid(
            row=0,
            column=1,
            sticky="nsew",
            padx=(8, 0),
        )
        insulation_frame.grid_rowconfigure(1, weight=1)

        self.ins_rule_operator = tk.StringVar(value="포함")
        self.ins_rule_value = tk.StringVar(value="")
        self.ins_rule_result = tk.StringVar(value="H")

        ttk.Combobox(
            insulation_frame,
            textvariable=self.ins_rule_operator,
            values=OPERATORS,
            state="readonly",
            width=18,
        ).grid(row=0, column=0, padx=3)
        self.ins_rule_value_combo = ttk.Combobox(
            insulation_frame,
            textvariable=self.ins_rule_value,
            values=[],
            width=28,
        )
        self.ins_rule_value_combo.grid(row=0, column=1, padx=3)
        ttk.Combobox(
            insulation_frame,
            textvariable=self.ins_rule_result,
            values=["N", "H", "P", "M", "R", "C"],
            width=12,
        ).grid(row=0, column=2, padx=3)
        tk.Button(
            insulation_frame,
            text="추가",
            command=self.add_insulation_rule,
        ).grid(row=0, column=3, padx=3)

        self.ins_rule_list = tk.Listbox(
            insulation_frame,
            height=8,
        )
        self.ins_rule_list.grid(
            row=1,
            column=0,
            columnspan=4,
            sticky="nsew",
            pady=4,
        )
        tk.Button(
            insulation_frame,
            text="선택 삭제",
            command=self.delete_insulation_rule,
        ).grid(row=2, column=0, columnspan=4)

        button_frame = tk.Frame(self.paint_tab)
        button_frame.pack(fill="x", padx=10, pady=(5, 10))

        self.paint_save_rules_button = tk.Button(
            button_frame,
            text="규칙 저장",
            command=self.save_temperature_painting_rules,
            width=14,
            height=2,
        )
        self.paint_save_rules_button.pack(side="left")

        self.paint_load_rules_button = tk.Button(
            button_frame,
            text="규칙 불러오기",
            command=self.load_temperature_painting_rules,
            width=14,
            height=2,
        )
        self.paint_load_rules_button.pack(side="left", padx=(6, 16))

        self.paint_run_button = tk.Button(
            button_frame,
            text="Temperature & Painting 실행",
            command=lambda: self.start_thread(
                self.run_temperature_painting
            ),
            width=34,
            height=2,
            font=("Arial", 10, "bold"),
        )
        self.paint_run_button.pack(side="left")

        self.paint_material_frame = material_frame
        self.paint_suffix_frame = insulation_frame
        self.update_paint_source_ui()

    def show_paint_header_full_name(self, header_name):
        """Show a complete header name below the comboboxes."""
        self.paint_selected_header_text.set(
            clean_text(header_name)
            or "선택한 열의 전체명이 여기에 표시됩니다."
        )

    @staticmethod
    def _remove_display_sheet_suffix(value, separator="-"):
        """
        Remove the Sh't No. (text after the last separator) from a Display Drawing Number.

        Example:
            304-PW-0051-01 -> 304-PW-0051
            304-PW-0051-A1 -> 304-PW-0051
        """
        normalized = clean_text(value)
        if not normalized:
            return ""
        separator = separator or "-"
        if separator not in normalized:
            return normalized
        return normalized.rsplit(separator, 1)[0].strip()

    def _match_line_record(
        self,
        drawing_number,
        line_map,
        sorted_line_keys,
        match_cache,
    ):
        """
        Match a Display Drawing Number to a Line List Line No.

        Returns:
            (record_or_none, matched_key, issue_reason)
        """
        cache_key = (
            self.paint_line_match_mode.get(),
            self.paint_sheet_separator.get(),
            clean_text(drawing_number),
        )
        if cache_key in match_cache:
            return match_cache[cache_key]

        mode = self.paint_line_match_mode.get()
        display_key = self._normalize_line_key(drawing_number)

        if mode == LINE_MATCH_STRIP_SHEET:
            display_key = self._normalize_line_key(
                self._remove_display_sheet_suffix(
                    drawing_number,
                    self.paint_sheet_separator.get(),
                )
            )

        if not display_key:
            result = (None, "", "Display Drawing Number가 비어 있음")
            match_cache[cache_key] = result
            return result

        if mode in [LINE_MATCH_EXACT, LINE_MATCH_STRIP_SHEET]:
            record = line_map.get(display_key)
            result = (
                record,
                display_key if record else "",
                "" if record else "Line No. 매칭 실패",
            )
            match_cache[cache_key] = result
            return result

        if mode == LINE_MATCH_DISPLAY_CONTAINS:
            matches = [
                key
                for key in sorted_line_keys
                if key and key in display_key
            ]
        else:
            matches = [
                key
                for key in sorted_line_keys
                if display_key and display_key in key
            ]

        if not matches:
            result = (None, "", "Line No. 포함 매칭 실패")
            match_cache[cache_key] = result
            return result

        # The longest key is the most specific. If more than one equally
        # specific key exists, do not select one arbitrarily.
        longest_length = len(matches[0])
        most_specific = [
            key for key in matches if len(key) == longest_length
        ]
        if len(most_specific) > 1:
            result = (
                None,
                "",
                "Line No. 포함 매칭 결과가 여러 개임: "
                + ", ".join(most_specific[:10]),
            )
            match_cache[cache_key] = result
            return result

        matched_key = most_specific[0]
        result = (
            line_map[matched_key],
            matched_key,
            (
                "포함 방식으로 매칭"
                if matched_key != display_key
                else ""
            ),
        )
        match_cache[cache_key] = result
        return result

    def _ins_temp_uses_rule_file(self):
        return self.paint_ins_temp_source.get() == INS_TEMP_SOURCE_RULE

    def _ins_temp_uses_line_list(self):
        return not self._ins_temp_uses_rule_file()

    def _paint_uses_rule_file(self):
        return self.paint_symbol_source.get() == PAINT_SOURCE_RULE

    def _line_list_required(self):
        """Line List is needed only when one of the two sides reads it."""
        return self._ins_temp_uses_line_list() or not self._paint_uses_rule_file()

    def _show_grid_rows(self, frame, rows, visible):
        """Show or hide every widget placed on the given grid rows."""
        # grid_slaves()는 숨긴 위젯을 돌려주지 않으므로 처음 본 위젯을 기억해 둔다.
        cache = self.__dict__.setdefault("_grid_row_widgets", {})
        for row_no in rows:
            key = (str(frame), row_no)
            if key not in cache:
                cache[key] = frame.grid_slaves(row=row_no)
            for widget in cache[key]:
                if visible:
                    widget.grid()
                else:
                    widget.grid_remove()

    def _build_rule_file_rows(self, parent, row, label, file_var, sheet_var, example_command):
        """File + sheet + 예시 rows for a 'Display 열 기준 규칙 파일'."""
        tk.Label(parent, text=label).grid(row=row, column=0, sticky="e", padx=(0, 4), pady=2)
        file_frame = tk.Frame(parent)
        file_frame.grid(row=row, column=1, sticky="ew", pady=2)
        file_frame.grid_columnconfigure(0, weight=1)
        tk.Entry(file_frame, textvariable=file_var).grid(row=0, column=0, sticky="ew")

        tk.Label(parent, text="규칙 시트").grid(row=row + 1, column=0, sticky="e", padx=(0, 4), pady=2)
        sheet_frame = tk.Frame(parent)
        sheet_frame.grid(row=row + 1, column=1, sticky="w", pady=2)
        sheet_combo = ttk.Combobox(sheet_frame, textvariable=sheet_var, state="readonly", width=18)
        sheet_combo.pack(side="left")
        tk.Button(sheet_frame, text="규칙 파일 예시", command=example_command).pack(side="left", padx=6)

        def choose_file():
            path = filedialog.askopenfilename(
                initialdir=str(self.project_root / "03_reference"),
                filetypes=[("Excel", "*.xlsx *.xlsm")],
            )
            if path:
                file_var.set(path)
                self._load_sheet_names_to_combo(path, sheet_combo, sheet_var)

        tk.Button(file_frame, text="파일 선택", command=choose_file).grid(row=0, column=1, padx=(4, 0))
        return {"rows": [row, row + 1], "sheet_combo": sheet_combo}

    def update_insulation_temp_source_ui(self):
        """Show only the inputs used by the selected Insulation Temp. source."""
        if not hasattr(self, "paint_ins_rule_widgets"):
            return
        use_rule = self._ins_temp_uses_rule_file()
        self._show_grid_rows(self.paint_ins_frame, [1], not use_rule)
        self._show_grid_rows(self.paint_ins_frame, self.paint_ins_rule_widgets["rows"], use_rule)
        if use_rule:
            self.paint_ins_help.set(
                "규칙 파일의 조건을 위에서부터 검사해 처음 일치한 규칙 1개만 적용합니다 "
                "(VLOOKUP처럼 먼저 나온 조건 우선). 조건이 겹치면 원하는 규칙을 위에 두세요."
            )
        else:
            self.paint_ins_help.set(
                "Line List에서 선택한 온도 열(Operating / Max. Operating / Design 등) 값을 입력합니다."
            )

    def update_paint_source_ui(self):
        """Show only the inputs used by the selected Paint Symbol source."""
        if not hasattr(self, "paint_rule_widgets"):
            return
        use_rule = self._paint_uses_rule_file()
        self._show_grid_rows(self.paint_code_frame, [1, 2, 3], not use_rule)
        self._show_grid_rows(self.paint_code_frame, self.paint_rule_widgets["rows"], use_rule)
        for frame in [getattr(self, "paint_material_frame", None), getattr(self, "paint_suffix_frame", None)]:
            if frame is not None:
                self._set_widgets_state(frame, not use_rule)
        if use_rule:
            self.paint_code_help.set(
                "규칙 파일의 조건을 위에서부터 검사해 처음 일치한 규칙의 Paint Symbol을 입력합니다 "
                "(VLOOKUP처럼 먼저 나온 조건 우선). 아래 Material Group / Paint Suffix 규칙은 사용하지 않습니다."
            )
        else:
            self.paint_code_help.set(
                "Paint Symbol은 Insulation과 Temp를 바탕으로 계산합니다: "
                "Class → Material Group, Insulation Symbol → Paint Suffix로 Paint Type(예: CS-H)을 정하고, "
                "Painting 적용 Temp. 구간에 맞는 값을 Painting Code Table에서 찾습니다."
            )

    def show_insulation_temp_rule_example(self):
        self._show_table_example(
            title="Insulation Temperature 규칙 파일 예시",
            subtitle=(
                "Display Format의 열 값(Fluid, Class 등)에 따라 Insulation Temp를 일괄 입력하는 규칙 파일입니다."
            ),
            note=(
                "1행은 제목(자유롭게 작성), 2행부터 규칙을 입력합니다.\n"
                "A열 = Display Format 열 이름(예: Fluid, Class)  /  B열 = 그 열의 값(정확히 일치, 대소문자 무시, "
                "'*' = 모든 값)  /  C열 = 입력할 온도 (숫자, 범위, AMB 가능)\n"
                "조건이 겹치면 위에 있는 규칙이 우선입니다 (VLOOKUP 방식). "
                "예) Fluid=FW, Class=A1C인 행은 아래 표에서 2행(Fluid FW → 45)이 먼저 일치하므로 45."
            ),
            headers=["Display Format 열", "값", "Insulation Temp"],
            rows=[
                ["Fluid", "SC", "180"],
                ["Fluid", "FW", "45"],
                ["Class", "A1C", "35"],
                ["Class", "B2S", "AMB"],
                ["Fluid", "*", "60"],
            ],
            file_name="Insulation_Temp_Rule_Example.xlsx",
            sheet_title="INS TEMP RULE",
        )

    def show_paint_rule_example(self):
        self._show_table_example(
            title="Paint Symbol 규칙 파일 예시",
            subtitle=(
                "Display Format의 열 값(Insulation Symbol, Class 등)에 따라 Paint Symbol을 직접 입력하는 규칙 파일입니다."
            ),
            note=(
                "1행은 제목(자유롭게 작성), 2행부터 규칙을 입력합니다.\n"
                "A열 = Display Format 열 이름(예: Insulation Symbol, Class)  /  B열 = 그 열의 값(정확히 일치, "
                "대소문자 무시, '*' = 모든 값)  /  C열 = 입력할 Paint Symbol\n"
                "조건이 겹치면 위에 있는 규칙이 우선입니다 (VLOOKUP 방식)."
            ),
            headers=["Display Format 열", "값", "Paint Symbol"],
            rows=[
                ["Insulation Symbol", "H", "B1→B2"],
                ["Insulation Symbol", "P", "B1→B2"],
                ["Class", "A1C", "A1→A4"],
                ["Insulation Symbol", "*", "No"],
            ],
            file_name="Paint_Symbol_Rule_Example.xlsx",
            sheet_title="PAINT RULE",
        )

    def _load_display_rule_file(self, path_text, sheet_text, display_headers, label, value_parser):
        """
        Read a 'Display 열 기준 규칙 파일'.

        A: Display Format column, B: value ('*' = any), C: result.
        value_parser(raw) -> (parsed_value or None, original_text).
        Returns a list of (display_column, value_casefold, result, excel_row), in file order.
        """
        path = Path(path_text.strip())
        if not path.is_file():
            raise ValueError(f"{label} 규칙 파일을 선택하세요.")

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_name = sheet_text.strip() or wb.sheetnames[0]
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"{label} 규칙 시트를 찾을 수 없습니다: {sheet_name}")
            ws = wb[sheet_name]
            rules = []
            errors = []
            for excel_row, values in enumerate(
                ws.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                values = list(values) + [None] * (3 - len(values))
                column_name = clean_text(values[0])
                if not column_name and is_effectively_empty_row(values[:3]):
                    continue
                display_column = self._resolve_header(column_name, display_headers)
                if not display_column:
                    errors.append(f"{excel_row}행: Display Format에 '{column_name}' 열이 없음")
                    continue
                result, original = value_parser(values[2])
                if result is None:
                    errors.append(f"{excel_row}행: C열 값 오류 ('{original}')")
                    continue
                rules.append((
                    display_column,
                    clean_text(values[1]).casefold(),
                    result,
                    excel_row,
                ))
        finally:
            wb.close()

        if errors:
            raise ValueError(
                f"{label} 규칙 파일을 확인하세요:\n" + "\n".join(errors[:20])
            )
        if not rules:
            raise ValueError(f"{label} 규칙 파일에 규칙이 없습니다. (2행부터 입력)")
        return rules

    def _load_insulation_temp_rules(self, display_headers, ambient_temperature):
        def parse_temperature(raw):
            temperature, _status, original = self._normalize_temperature(raw, ambient_temperature)
            return temperature, original

        return self._load_display_rule_file(
            self.paint_ins_rule_file.get(),
            self.paint_ins_rule_sheet.get(),
            display_headers,
            "Insulation Temp",
            parse_temperature,
        )

    def _load_paint_symbol_rules(self, display_headers):
        def parse_symbol(raw):
            text = clean_text(raw)
            return (text or None), text

        return self._load_display_rule_file(
            self.paint_rule_file.get(),
            self.paint_rule_sheet.get(),
            display_headers,
            "Paint Symbol",
            parse_symbol,
        )

    @staticmethod
    def _value_from_display_rules(row_values, header_index, rules):
        """VLOOKUP style: return (result, rule_excel_row) of the FIRST matching rule."""
        for display_column, rule_value, result, excel_row in rules:
            cell = clean_text(row_values[header_index[display_column]]).casefold()
            if rule_value == "*" or cell == rule_value:
                return result, excel_row
        return None, None

    def _temperature_painting_rule_config(self):
        return {
            "display_sheet": self.paint_display_sheet.get(),
            "line_sheet": self.paint_line_sheet.get(),
            "line_header_start": self.paint_line_header_start.get(),
            "line_header_end": self.paint_line_header_end.get(),
            "line_data_start": self.paint_line_data_start.get(),
            "drawing_column": self.paint_drawing_col.get(),
            "class_column": self.paint_class_col.get(),
            "insulation_column": self.paint_insulation_col.get(),
            "line_no_column": self.paint_line_no_col.get(),
            "operating_temp_column": self.paint_operating_temp_col.get(),
            "maximum_operating_temp_column": self.paint_max_operating_temp_col.get(),
            "line_match_mode": self.paint_line_match_mode.get(),
            "sheet_separator": self.paint_sheet_separator.get(),
            "insulation_temp_source": self.paint_ins_temp_source.get(),
            "insulation_temp_rule_file": self.paint_ins_rule_file.get(),
            "insulation_temp_rule_sheet": self.paint_ins_rule_sheet.get(),
            "paint_symbol_source": self.paint_symbol_source.get(),
            "paint_symbol_rule_file": self.paint_rule_file.get(),
            "paint_symbol_rule_sheet": self.paint_rule_sheet.get(),
            "ambient_temperature": self.paint_ambient_temp.get(),
            "painting_table_sheet": self.paint_table_sheet.get(),
            "painting_table_header": self.paint_table_header.get(),
            "material_rules": self.material_rules,
            "insulation_rules": self.insulation_rules,
        }

    def save_temperature_painting_rules(self):
        path = filedialog.asksaveasfilename(
            initialdir=self.project_root / "03_reference",
            defaultextension=".json",
            filetypes=[("JSON", "*.json")],
        )
        if not path:
            return
        Path(path).write_text(
            json.dumps(
                self._temperature_painting_rule_config(),
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        messagebox.showinfo(
            "완료",
            "Temperature & Painting 규칙을 저장했습니다.",
        )

    def load_temperature_painting_rules(self):
        path = filedialog.askopenfilename(
            initialdir=self.project_root / "03_reference",
            filetypes=[("JSON", "*.json")],
        )
        if not path:
            return

        config = json.loads(
            Path(path).read_text(encoding="utf-8")
        )

        self.paint_display_sheet.set(
            config.get("display_sheet", self.paint_display_sheet.get())
        )
        self.paint_line_sheet.set(
            config.get("line_sheet", self.paint_line_sheet.get())
        )
        self.paint_line_header_start.set(
            config.get("line_header_start", "2")
        )
        self.paint_line_header_end.set(
            config.get("line_header_end", "4")
        )
        self.paint_line_data_start.set(
            config.get("line_data_start", "6")
        )
        self.paint_drawing_col.set(
            config.get("drawing_column", "Drawing Number")
        )
        self.paint_class_col.set(
            config.get("class_column", "Class")
        )
        self.paint_insulation_col.set(
            config.get("insulation_column", "Insulation Symbol")
        )
        self.paint_line_no_col.set(
            config.get("line_no_column", "")
        )
        self.paint_operating_temp_col.set(
            config.get("operating_temp_column", "")
        )
        self.paint_max_operating_temp_col.set(
            config.get("maximum_operating_temp_column", "")
        )
        line_match_mode = config.get("line_match_mode", LINE_MATCH_STRIP_SHEET)
        self.paint_line_match_mode.set(
            LEGACY_LINE_MATCH_MODES.get(line_match_mode, line_match_mode)
        )
        self.paint_sheet_separator.set(config.get("sheet_separator", "-"))
        ins_source = config.get("insulation_temp_source", INS_TEMP_SOURCE_LINE)
        self.paint_ins_temp_source.set(LEGACY_INS_TEMP_SOURCES.get(ins_source, ins_source))
        self.paint_symbol_source.set(config.get("paint_symbol_source", PAINT_SOURCE_TABLE))
        self.paint_rule_file.set(config.get("paint_symbol_rule_file", ""))
        self.paint_rule_sheet.set(config.get("paint_symbol_rule_sheet", ""))
        self.paint_ins_rule_file.set(config.get("insulation_temp_rule_file", ""))
        self.paint_ins_rule_sheet.set(config.get("insulation_temp_rule_sheet", ""))
        self.update_insulation_temp_source_ui()
        self.update_paint_source_ui()
        self.paint_ambient_temp.set(
            config.get("ambient_temperature", "35")
        )
        self.paint_table_sheet.set(
            config.get("painting_table_sheet", "")
        )
        self.paint_table_header.set(
            config.get("painting_table_header", "1")
        )

        self.material_rules = config.get("material_rules", [])
        self.insulation_rules = config.get("insulation_rules", [])

        self.material_rule_list.delete(0, "end")
        for rule in self.material_rules:
            self.material_rule_list.insert(
                "end",
                f"{rule['operator']} {rule['value']} → {rule['result']}",
            )

        self.ins_rule_list.delete(0, "end")
        for rule in self.insulation_rules:
            self.ins_rule_list.insert(
                "end",
                f"{rule['operator']} {rule['value']} → {rule['result']}",
            )

        messagebox.showinfo(
            "완료",
            "Temperature & Painting 규칙을 불러왔습니다.",
        )

    def _create_painting_code_example_file(self):
        """Create the example workbook and return its path."""
        example_path = (
            self.project_root
            / "03_reference"
            / "Painting_Code_Table_Example.xlsx"
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "PAINT CODE"

        ws.merge_cells("B1:E1")
        ws["B1"] = "MAX OPERATING TEMPERATURE (°C)"
        ws["A2"] = "PAINT TYPE"
        ws["B2"] = "t ≤ 120 °C"
        ws["C2"] = "120 °C < t ≤ 200 °C"
        ws["D2"] = "200 °C < t ≤ 400 °C"
        ws["E2"] = "400 °C < t ≤ 500 °C"

        sample_rows = [
            ["CS-N", "A1→A4", "C1→C3", "C1→C3", "F1→F3"],
            ["CS-H", "B1→B2", "B1→B2", "C1→C2", "F1→F3"],
            ["CS-P", "B1→B2", "B1→B2", "C1→C2", "F1→F3"],
            ["CS-M", "B1→B2", "B1→B2", "C1→C2", "F1→F3"],
            ["CS-R", "B1→B2", "B1→B2", "C1→C2", "F1→F3"],
            ["CS-C", "B1→B2", "B1→B2", "C1→C2", "F1→F3"],
            ["SS-N", "No", "No", "No", "No"],
            ["SS-H", "B1→B2", "B1→B2", "F1→F3", "F1→F3"],
            ["SS-P", "B1→B2", "B1→B2", "F1→F3", "F1→F3"],
            ["SS-M", "B1→B2", "B1→B2", "F1→F3", "F1→F3"],
            ["SS-R", "B1→B2", "B1→B2", "F1→F3", "F1→F3"],
            ["HDPE-N", "No", "No", "No", "No"],
            ["HDPE-H", "No", "No", "No", "No"],
            ["HDPE-P", "No", "No", "No", "No"],
            ["HDPE-M", "No", "No", "No", "No"],
            ["HDPE-R", "No", "No", "No", "No"],
        ]
        for row in sample_rows:
            ws.append(row)

        title_fill = PatternFill("solid", fgColor="F4E6B1")
        white_fill = PatternFill("solid", fgColor="FFFFFF")

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
        for cell in ws[2]:
            cell.fill = title_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row_no in range(10, 14):
            for col_no in range(4, 6):
                ws.cell(row_no, col_no).fill = white_fill

        widths = {
            "A": 16,
            "B": 19,
            "C": 23,
            "D": 23,
            "E": 23,
        }
        for column, width in widths.items():
            ws.column_dimensions[column].width = width

        ws["A20"] = (
            "NOTE: Temperature range columns may continue beyond column E. "
            "Add F, G, H... columns as needed; the program reads all populated range columns."
        )
        ws.merge_cells("A20:E20")
        ws["A20"].font = Font(italic=True, color="666666")
        ws["A20"].alignment = Alignment(wrap_text=True)

        ws.freeze_panes = "A3"
        example_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(example_path)
        return example_path, sample_rows

    def show_painting_code_example(self):
        """Show the Painting Code Table example immediately in a popup."""
        example_path, sample_rows = self._create_painting_code_example_file()

        popup = tk.Toplevel(self.root)
        popup.title("Painting Code Table 예시")
        popup.geometry("980x650")
        popup.minsize(850, 560)
        popup.transient(self.root)
        popup.grab_set()
        popup.configure(bg=self.colors["background"])

        header = tk.Frame(
            popup,
            bg=self.colors["navy"],
            padx=18,
            pady=12,
        )
        header.pack(fill="x")
        tk.Label(
            header,
            text="Painting Code Table 입력 예시",
            bg=self.colors["navy"],
            fg="#FFFFFF",
            font=("Malgun Gothic", 15, "bold"),
            anchor="w",
        ).pack(anchor="w")
        tk.Label(
            header,
            text=(
                "A열에는 Paint Type을 입력하고, B열 이후에는 필요한 만큼 "
                "Painting 적용 Temp. 구간별 Paint Symbol 열을 추가할 수 있습니다."
            ),
            bg=self.colors["navy"],
            fg="#D4E1EC",
            font=("Malgun Gothic", 9),
            anchor="w",
        ).pack(anchor="w", pady=(3, 0))

        note = tk.Frame(
            popup,
            bg="#FFF8D8",
            padx=12,
            pady=9,
            highlightbackground="#E8D792",
            highlightthickness=1,
        )
        note.pack(fill="x", padx=14, pady=(12, 8))
        tk.Label(
            note,
            text=(
                "중요: 실제 온도 범위 헤더는 2행입니다. "
                "온도 구간 열은 E열을 넘어 F, G, H열 등으로 계속 추가할 수 있습니다. "
                "5번 탭의 Painting Table '헤더 행'에는 2를 입력하세요."
            ),
            bg="#FFF8D8",
            fg="#4B3D00",
            font=("Malgun Gothic", 9, "bold"),
            anchor="w",
        ).pack(fill="x")

        table_card = tk.Frame(
            popup,
            bg="#FFFFFF",
            highlightbackground="#D4DAE1",
            highlightthickness=1,
        )
        table_card.pack(
            fill="both",
            expand=True,
            padx=14,
            pady=(0, 8),
        )

        canvas = tk.Canvas(
            table_card,
            bg="#FFFFFF",
            highlightthickness=0,
        )
        vertical_scroll = ttk.Scrollbar(
            table_card,
            orient="vertical",
            command=canvas.yview,
        )
        horizontal_scroll = ttk.Scrollbar(
            table_card,
            orient="horizontal",
            command=canvas.xview,
        )
        canvas.configure(
            yscrollcommand=vertical_scroll.set,
            xscrollcommand=horizontal_scroll.set,
        )

        canvas.grid(row=0, column=0, sticky="nsew")
        vertical_scroll.grid(row=0, column=1, sticky="ns")
        horizontal_scroll.grid(row=1, column=0, sticky="ew")
        table_card.grid_rowconfigure(0, weight=1)
        table_card.grid_columnconfigure(0, weight=1)

        table = tk.Frame(canvas, bg="#FFFFFF")
        table_window = canvas.create_window(
            (0, 0),
            window=table,
            anchor="nw",
        )

        headers = [
            "PAINT TYPE",
            "t ≤ 120 °C",
            "120 °C < t ≤ 200 °C",
            "200 °C < t ≤ 400 °C",
            "400 °C < t ≤ 500 °C",
        ]

        tk.Label(
            table,
            text="",
            bg="#F4E6B1",
            relief="solid",
            bd=1,
            width=18,
            height=2,
        ).grid(row=0, column=0, sticky="nsew")
        tk.Label(
            table,
            text="MAX OPERATING TEMPERATURE (°C)",
            bg="#F4E6B1",
            fg="#111111",
            relief="solid",
            bd=1,
            font=("Malgun Gothic", 10, "bold"),
            height=2,
        ).grid(
            row=0,
            column=1,
            columnspan=4,
            sticky="nsew",
        )

        for column_no, header_text in enumerate(headers):
            tk.Label(
                table,
                text=header_text,
                bg="#F4E6B1",
                fg="#111111",
                relief="solid",
                bd=1,
                font=("Malgun Gothic", 9, "bold"),
                width=24 if column_no else 18,
                height=2,
            ).grid(
                row=1,
                column=column_no,
                sticky="nsew",
            )

        for row_no, row_values in enumerate(sample_rows, start=2):
            for column_no, value in enumerate(row_values):
                tk.Label(
                    table,
                    text=value,
                    bg="#FFFFFF",
                    fg="#111111",
                    relief="solid",
                    bd=1,
                    font=("Malgun Gothic", 9),
                    width=24 if column_no else 18,
                    height=1,
                ).grid(
                    row=row_no,
                    column=column_no,
                    sticky="nsew",
                )

        def update_scroll_region(_event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        table.bind("<Configure>", update_scroll_region)

        button_bar = tk.Frame(
            popup,
            bg=self.colors["background"],
        )
        button_bar.pack(fill="x", padx=14, pady=(0, 12))

        def open_example_excel():
            try:
                os.startfile(example_path)
            except (AttributeError, OSError):
                messagebox.showinfo(
                    "예시 파일 위치",
                    str(example_path),
                    parent=popup,
                )

        tk.Button(
            button_bar,
            text="예시 Excel 열기",
            command=open_example_excel,
            bg=self.colors["navy"],
            fg="#FFFFFF",
            activebackground=self.colors["navy_hover"],
            activeforeground="#FFFFFF",
            relief="flat",
            bd=0,
            padx=16,
            pady=7,
            font=self.bold_font,
        ).pack(side="left")

        tk.Button(
            button_bar,
            text="닫기",
            command=popup.destroy,
            bg=self.colors["gray_button"],
            fg="#FFFFFF",
            activebackground=self.colors["gray_hover"],
            activeforeground="#FFFFFF",
            relief="flat",
            bd=0,
            padx=16,
            pady=7,
            font=self.bold_font,
        ).pack(side="right")

    def choose_paint_display_file(self):
        path = filedialog.askopenfilename(
            filetypes=[("Excel", "*.xlsx *.xlsm")]
        )
        if path:
            self.paint_display_file.set(path)
            self._load_sheet_names_to_combo(
                path,
                self.paint_display_sheet_combo,
                self.paint_display_sheet,
            )

    def choose_paint_line_file(self):
        path = filedialog.askopenfilename(
            filetypes=[("Excel", "*.xlsx *.xlsm")]
        )
        if path:
            self.paint_line_file.set(path)
            self._load_sheet_names_to_combo(
                path,
                self.paint_line_sheet_combo,
                self.paint_line_sheet,
            )

    def choose_paint_table_file(self):
        path = filedialog.askopenfilename(
            filetypes=[("Excel", "*.xlsx *.xlsm")]
        )
        if path:
            self.paint_table_file.set(path)
            self._load_sheet_names_to_combo(
                path,
                self.paint_table_sheet_combo,
                self.paint_table_sheet,
            )

    def choose_paint_output_file(self):
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if path:
            self.paint_output_file.set(path)

    def _read_headers(self, path, sheet, row):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet]
        values = [
            cell.value
            for cell in next(
                ws.iter_rows(min_row=row, max_row=row)
            )
        ]
        wb.close()
        return make_unique_headers(values)

    def _read_multiline_headers(
        self,
        path,
        sheet_name,
        header_start,
        header_end,
    ):
        """Build unique column names from a multi-row, merged-cell header."""
        wb = load_workbook(path, read_only=False, data_only=True)
        ws = wb[sheet_name]

        if header_start < 1 or header_end < header_start:
            wb.close()
            raise ValueError("Line List 헤더 시작/종료 행을 확인하세요.")

        max_column = ws.max_column
        merged_values = {}

        for merged_range in ws.merged_cells.ranges:
            if (
                merged_range.max_row < header_start
                or merged_range.min_row > header_end
            ):
                continue
            anchor_value = ws.cell(
                merged_range.min_row,
                merged_range.min_col,
            ).value
            for row_no in range(
                max(header_start, merged_range.min_row),
                min(header_end, merged_range.max_row) + 1,
            ):
                for col_no in range(
                    merged_range.min_col,
                    merged_range.max_col + 1,
                ):
                    merged_values[(row_no, col_no)] = anchor_value

        raw_headers = []
        for col_no in range(1, max_column + 1):
            parts = []
            for row_no in range(header_start, header_end + 1):
                value = merged_values.get(
                    (row_no, col_no),
                    ws.cell(row_no, col_no).value,
                )
                value_text = clean_text(value)
                if value_text and value_text not in parts:
                    parts.append(value_text)
            raw_headers.append(" / ".join(parts) if parts else f"Column{col_no}")

        wb.close()
        return make_unique_headers(raw_headers)

    def load_paint_display_columns(self):
        path = self.paint_display_file.get()
        self._load_sheet_names_to_combo(
            path,
            self.paint_display_sheet_combo,
            self.paint_display_sheet,
        )
        headers = self._read_headers(
            path,
            self.paint_display_sheet.get(),
            1,
        )
        for combo in self.paint_display_combos:
            combo["values"] = headers

        insulation_column = self.paint_insulation_col.get()
        if insulation_column not in headers:
            canonical_lookup = {
                canonical_header(header): header
                for header in headers
            }
            insulation_column = canonical_lookup.get(
                canonical_header("Insulation Symbol"),
                "",
            )
            if insulation_column:
                self.paint_insulation_col.set(insulation_column)

        unique_values = set()
        if insulation_column in headers:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[self.paint_display_sheet.get()]
            iterator = ws.iter_rows(values_only=True)
            loaded_headers = make_unique_headers(next(iterator))
            insulation_index = loaded_headers.index(insulation_column)

            for values in iterator:
                if insulation_index >= len(values):
                    continue
                value_text = clean_text(values[insulation_index])
                if value_text:
                    unique_values.add(value_text)
                if len(unique_values) >= 500:
                    break
            wb.close()

        self.paint_insulation_values = sorted(
            unique_values,
            key=lambda value: value.casefold(),
        )
        self.ins_rule_value_combo["values"] = self.paint_insulation_values

        self.log(
            f"Display 열 불러오기: {len(headers)}개 / "
            f"Insulation Symbol 고유값 {len(self.paint_insulation_values)}개"
        )

    def load_paint_line_columns(self):
        path = Path(self.paint_line_file.get().strip())
        if not path.exists():
            messagebox.showerror("오류", "Line List 파일을 찾을 수 없습니다.")
            return

        self._load_sheet_names_to_combo(
            path,
            self.paint_line_sheet_combo,
            self.paint_line_sheet,
        )

        try:
            header_start = int(
                self.paint_line_header_start.get().strip()
            )
            header_end = int(
                self.paint_line_header_end.get().strip()
            )
            data_start = int(
                self.paint_line_data_start.get().strip()
            )
        except ValueError:
            messagebox.showerror(
                "오류",
                "헤더 시작·종료 및 데이터 시작 행은 정수여야 합니다.",
            )
            return

        if data_start <= header_end:
            messagebox.showerror(
                "오류",
                "데이터 시작 행은 헤더 종료 행보다 뒤에 있어야 합니다.",
            )
            return

        headers = self._read_multiline_headers(
            path,
            self.paint_line_sheet.get(),
            header_start,
            header_end,
        )
        for combo in self.paint_line_combos:
            combo["values"] = headers

        if headers:
            self.show_paint_header_full_name(headers[0])

        self.log(
            "Line List 다중 헤더 불러오기: "
            f"{len(headers)}개 열 / 헤더 {header_start}~{header_end}행 / "
            f"데이터 {data_start}행부터"
        )

    def add_material_rule(self):
        rule = {
            "operator": self.material_rule_operator.get(),
            "value": self.material_rule_value.get(),
            "result": self.material_rule_result.get(),
        }
        self.material_rules.append(rule)
        self.material_rule_list.insert(
            "end",
            f"{rule['operator']} {rule['value']} → {rule['result']}",
        )

    def delete_material_rule(self):
        selected = self.material_rule_list.curselection()
        if selected:
            self.material_rules.pop(selected[0])
            self.material_rule_list.delete(selected[0])

    def add_insulation_rule(self):
        rule = {
            "operator": self.ins_rule_operator.get(),
            "value": self.ins_rule_value.get(),
            "result": self.ins_rule_result.get(),
        }
        self.insulation_rules.append(rule)
        self.ins_rule_list.insert(
            "end",
            f"{rule['operator']} {rule['value']} → {rule['result']}",
        )

    def delete_insulation_rule(self):
        selected = self.ins_rule_list.curselection()
        if selected:
            self.insulation_rules.pop(selected[0])
            self.ins_rule_list.delete(selected[0])

    def _classify_by_rules(self, value, rules):
        for rule in rules:
            if compare_value(
                value,
                rule["operator"],
                rule["value"],
            ):
                return rule["result"]
        return ""

    @staticmethod
    def _normalize_line_key(value):
        return clean_text(value).casefold()

    def _normalize_temperature(self, value, ambient_temperature):
        """
        Convert a Line List temperature to one numeric severe value.

        Returns:
            (normalized_value, status, detail)
        """
        original = clean_text(value)
        if not original:
            return None, "INVALID", "빈칸"

        compact = re.sub(r"\s+", "", original).upper()
        if compact in {"NA", "N/A", "N.A", "N.A.", "HOLD", "-"}:
            return None, "INVALID", original

        if re.fullmatch(r"AMB\.?", compact):
            return ambient_temperature, "AMBIENT", original

        # A hyphen between two digits is treated as a range separator,
        # while a leading minus sign remains a negative number.
        prepared = original.replace("–", "-").replace("—", "-")
        prepared = re.sub(r"(?<=\d)\s*-\s*(?=\d)", " | ", prepared)

        numbers = [
            float(number)
            for number in re.findall(
                r"(?<![\d.])-?\d+(?:\.\d+)?",
                prepared,
            )
        ]
        if not numbers:
            return None, "INVALID", original

        severe_value = max(numbers)
        if severe_value.is_integer():
            severe_value = int(severe_value)
        return severe_value, "OK", original

    def _read_line_list_rows(
        self,
        path,
        sheet_name,
        headers,
        data_start,
    ):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name]

        for excel_row, values in enumerate(
            ws.iter_rows(min_row=data_start, values_only=True),
            start=data_start,
        ):
            if is_effectively_empty_row(values):
                continue
            row = {
                headers[index]: (
                    values[index]
                    if index < len(values)
                    else None
                )
                for index in range(len(headers))
            }
            yield excel_row, row

        wb.close()

    def _parse_temp_band(self, text):
        raw = clean_text(text)
        normalized = (
            raw.replace("℃", "")
            .replace("°C", "")
            .replace("°С", "")
            .replace("C", "")
            .replace("С", "")
        )
        numbers = [
            float(value)
            for value in re.findall(r"-?\d+(?:\.\d+)?", normalized)
        ]
        if not numbers:
            return None

        lower_text = normalized.lower().replace(" ", "")

        if len(numbers) == 1:
            number = numbers[0]
            if (
                lower_text.startswith("t≤")
                or lower_text.startswith("t<=")
                or lower_text.endswith("≥t")
            ):
                return (None, number, True, True)
            if (
                lower_text.startswith("t<")
                or lower_text.endswith(">t")
            ):
                return (None, number, True, False)
            if (
                lower_text.startswith("t≥")
                or lower_text.startswith("t>=")
            ):
                return (number, None, True, True)
            if lower_text.startswith("t>"):
                return (number, None, False, True)

        if len(numbers) >= 2:
            minimum = numbers[0]
            maximum = numbers[1]
            min_inclusive = (
                f"{minimum:g}≤t" in lower_text
                or f"{minimum:g}<=t" in lower_text
            )
            max_inclusive = (
                "t≤" in lower_text
                or "t<=" in lower_text
            )
            return (
                minimum,
                maximum,
                min_inclusive,
                max_inclusive,
            )

        return None

    def _load_paint_matrix(self):
        path = Path(self.paint_table_file.get().strip())
        if not path.exists():
            raise ValueError("Painting Code Table 파일을 선택하세요.")

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[self.paint_table_sheet.get()]
        header_row = int(self.paint_table_header.get() or "1")
        row_iterator = ws.iter_rows(
            min_row=header_row,
            values_only=True,
        )
        headers = [clean_text(value) for value in next(row_iterator)]

        matrix = {}
        for values in row_iterator:
            paint_type = clean_text(values[0] if values else None)
            if not paint_type:
                continue
            for index in range(
                1,
                min(len(headers), len(values)),
            ):
                band = self._parse_temp_band(headers[index])
                symbol = clean_text(values[index])
                if band and symbol:
                    matrix.setdefault(
                        paint_type.upper(),
                        [],
                    ).append((*band, symbol))

        wb.close()
        return matrix

    def _paint_symbol(self, paint_type, temperature, matrix):
        numeric_temperature = normalize_numeric(temperature)
        if numeric_temperature is None:
            return ""

        for (
            minimum,
            maximum,
            min_inclusive,
            max_inclusive,
            symbol,
        ) in matrix.get(paint_type.upper(), []):
            left_ok = (
                True
                if minimum is None
                else (
                    numeric_temperature >= minimum
                    if min_inclusive
                    else numeric_temperature > minimum
                )
            )
            right_ok = (
                True
                if maximum is None
                else (
                    numeric_temperature <= maximum
                    if max_inclusive
                    else numeric_temperature < maximum
                )
            )
            if left_ok and right_ok:
                return symbol
        return ""

    def run_temperature_painting(self):
        self.set_running(True)
        self.set_status(
            "Temperature & Painting",
            "현재 작업: 입력값 검증",
            0,
        )

        try:
            display_path = Path(self.paint_display_file.get().strip())
            line_path = Path(self.paint_line_file.get().strip())
            output_path = Path(self.paint_output_file.get().strip())

            use_line_list = self._line_list_required()
            ins_from_line = self._ins_temp_uses_line_list()
            paint_from_rule = self._paint_uses_rule_file()

            if not display_path.exists():
                raise ValueError("Display Format 파일을 선택하세요.")
            if use_line_list and not line_path.exists():
                raise ValueError("Line List 파일을 선택하세요.")

            try:
                ambient_temperature = float(
                    self.paint_ambient_temp.get().strip()
                )
                if use_line_list:
                    header_start = int(
                        self.paint_line_header_start.get().strip()
                    )
                    header_end = int(
                        self.paint_line_header_end.get().strip()
                    )
                    data_start = int(
                        self.paint_line_data_start.get().strip()
                    )
                else:
                    header_start = header_end = data_start = ""
            except ValueError as exc:
                raise ValueError(
                    "헤더 행, 데이터 시작 행 및 AMB 대체 온도를 확인하세요."
                ) from exc

            line_map = {}
            duplicate_keys = set()
            conflicting_duplicate_keys = set()

            if use_line_list:
                if data_start <= header_end:
                    raise ValueError(
                        "Line List 데이터 시작 행은 헤더 종료 행보다 뒤여야 합니다."
                    )

                line_headers = self._read_multiline_headers(
                    line_path,
                    self.paint_line_sheet.get(),
                    header_start,
                    header_end,
                )

                required_line_columns = [self.paint_line_no_col.get()]
                if ins_from_line:
                    required_line_columns.append(self.paint_operating_temp_col.get())
                if not paint_from_rule:
                    required_line_columns.append(self.paint_max_operating_temp_col.get())
                missing_line_columns = [
                    column
                    for column in required_line_columns
                    if column not in line_headers
                ]
                if missing_line_columns:
                    raise ValueError(
                        "Line List 선택 열을 찾을 수 없습니다:\n"
                        + "\n".join(missing_line_columns)
                    )

                self.set_status(
                    "Temperature & Painting",
                    "현재 작업: Line List 매핑표 생성",
                    5,
                )

                for excel_row, row in self._read_line_list_rows(
                    line_path,
                    self.paint_line_sheet.get(),
                    line_headers,
                    data_start,
                ):
                    line_number = clean_text(
                        row.get(self.paint_line_no_col.get())
                    )
                    key = self._normalize_line_key(line_number)
                    if not key:
                        continue

                    line_record = {
                        "excel_row": excel_row,
                        "line_number": line_number,
                        "operating_original": (
                            row.get(self.paint_operating_temp_col.get())
                            if ins_from_line
                            else None
                        ),
                        "maximum_original": (
                            row.get(self.paint_max_operating_temp_col.get())
                            if not paint_from_rule
                            else None
                        ),
                    }

                    if key in line_map:
                        duplicate_keys.add(key)
                        previous = line_map[key]
                        if (
                            clean_text(previous["operating_original"])
                            != clean_text(line_record["operating_original"])
                            or clean_text(previous["maximum_original"])
                            != clean_text(line_record["maximum_original"])
                        ):
                            conflicting_duplicate_keys.add(key)
                    else:
                        line_map[key] = line_record

            sorted_line_keys = sorted(
                line_map.keys(),
                key=len,
                reverse=True,
            )
            line_match_cache = {}

            paint_matrix = {} if paint_from_rule else self._load_paint_matrix()

            wb_in = load_workbook(
                display_path,
                read_only=True,
                data_only=True,
            )
            ws_in = wb_in[self.paint_display_sheet.get()]
            iterator = ws_in.iter_rows(values_only=True)
            headers = make_unique_headers(next(iterator))
            header_index = {
                header: index
                for index, header in enumerate(headers)
            }

            required_display_columns = [
                "Insulation Temp (Operating Temp)",
                "Paint Symbol",
            ]
            if use_line_list:
                required_display_columns.append(self.paint_drawing_col.get())
            if not paint_from_rule:
                required_display_columns += [
                    self.paint_class_col.get(),
                    self.paint_insulation_col.get(),
                ]
            missing_display_columns = [
                column
                for column in required_display_columns
                if column not in header_index
            ]
            if missing_display_columns:
                wb_in.close()
                raise ValueError(
                    "Display Format 선택 열을 찾을 수 없습니다:\n"
                    + "\n".join(missing_display_columns)
                )

            insulation_temp_rules = []
            paint_symbol_rules = []
            try:
                if not ins_from_line:
                    insulation_temp_rules = self._load_insulation_temp_rules(
                        headers,
                        ambient_temperature,
                    )
                if paint_from_rule:
                    paint_symbol_rules = self._load_paint_symbol_rules(headers)
            except Exception:
                wb_in.close()
                raise

            drawing_index = header_index.get(self.paint_drawing_col.get())
            class_index = header_index.get(self.paint_class_col.get())
            insulation_index = header_index.get(self.paint_insulation_col.get())

            wb_out = Workbook(write_only=True)
            display_ws = wb_out.create_sheet("Display_Format")
            display_ws.append(headers)

            temperature_issue_ws = wb_out.create_sheet(
                "Temperature_Issues"
            )
            temperature_issue_ws.append([
                "Display Excel Row",
                "Drawing Number",
                "Line List Row",
                "Line No.",
                "Temperature Type",
                "Original Value",
                "Issue Reason",
            ])

            painting_issue_ws = wb_out.create_sheet(
                "Painting_Issues"
            )
            painting_issue_ws.append([
                "Display Excel Row",
                "Drawing Number",
                "Class",
                "Insulation Symbol",
                "Material Group",
                "Paint Suffix",
                "Paint Type",
                "Painting 적용 Temp.",
                "Issue Reason",
            ])

            summary = {
                "Total Display Rows": 0,
                "Line No. Matched": 0,
                "Line No. Unmatched": 0,
                "Line No. Matched by Contains": 0,
                "Line No. Ambiguous": 0,
                "Duplicate Line No.": 0,
                "Conflicting Duplicate Line No.": 0,
                "Insulation Temp. by Rule File": 0,
                "Insulation Temp. by Line List": 0,
                "Insulation Temp. Issue": 0,
                "Painting Temp. Converted": 0,
                "Painting Temp. Issue": 0,
                "Paint Symbol by Rule File": 0,
                "Paint Symbol Determined": 0,
                "Paint Symbol Issue": 0,
            }

            raw_max_row = ws_in.max_row
            if isinstance(raw_max_row, int) and raw_max_row > 1:
                estimated_total = raw_max_row - 1
            else:
                estimated_total = None

            self.set_status(
                "Temperature & Painting",
                "현재 작업: Display Format 처리 시작",
                10,
            )

            for display_row_no, values in enumerate(
                iterator,
                start=2,
            ):
                if is_effectively_empty_row(values):
                    continue

                summary["Total Display Rows"] += 1
                row_values = list(values) + [
                    None
                ] * (len(headers) - len(values))

                drawing_number = (
                    clean_text(row_values[drawing_index])
                    if drawing_index is not None
                    else ""
                )

                insulation_value_temp = None
                maximum_value = None
                line_operating_value = None
                line_operating_original = ""
                line_record = None

                if use_line_list:
                    (
                        line_record,
                        line_key,
                        line_match_issue,
                    ) = self._match_line_record(
                        drawing_number,
                        line_map,
                        sorted_line_keys,
                        line_match_cache,
                    )

                    if line_record is None:
                        summary["Line No. Unmatched"] += 1
                        if "여러 개" in line_match_issue:
                            summary["Line No. Ambiguous"] += 1
                        if ins_from_line and not paint_from_rule:
                            temperature_type = "Insulation / Painting 적용 Temp."
                        elif ins_from_line:
                            temperature_type = "Insulation 적용 Temp."
                        else:
                            temperature_type = "Painting 적용 Temp."
                        temperature_issue_ws.append([
                            display_row_no,
                            drawing_number,
                            "",
                            "",
                            temperature_type,
                            "",
                            line_match_issue or "Line No. 매칭 실패",
                        ])
                        if not paint_from_rule:
                            summary["Painting Temp. Issue"] += 1
                    else:
                        summary["Line No. Matched"] += 1
                        if line_match_issue == "포함 방식으로 매칭":
                            summary["Line No. Matched by Contains"] += 1

                        if line_match_issue:
                            temperature_issue_ws.append([
                                display_row_no,
                                drawing_number,
                                line_record["excel_row"],
                                line_record["line_number"],
                                "Line No.",
                                "",
                                line_match_issue,
                            ])

                        if line_key in duplicate_keys:
                            summary["Duplicate Line No."] += 1
                            issue_reason = "Line No. 중복"
                            if line_key in conflicting_duplicate_keys:
                                summary["Conflicting Duplicate Line No."] += 1
                                issue_reason = "Line No. 중복 및 온도값 불일치"
                            temperature_issue_ws.append([
                                display_row_no,
                                drawing_number,
                                line_record["excel_row"],
                                line_record["line_number"],
                                "Line No.",
                                "",
                                issue_reason,
                            ])

                        if ins_from_line:
                            (
                                line_operating_value,
                                _operating_status,
                                line_operating_original,
                            ) = self._normalize_temperature(
                                line_record["operating_original"],
                                ambient_temperature,
                            )

                        if not paint_from_rule:
                            (
                                maximum_value,
                                _maximum_status,
                                maximum_original,
                            ) = self._normalize_temperature(
                                line_record["maximum_original"],
                                ambient_temperature,
                            )
                            if maximum_value is None:
                                summary["Painting Temp. Issue"] += 1
                                temperature_issue_ws.append([
                                    display_row_no,
                                    drawing_number,
                                    line_record["excel_row"],
                                    line_record["line_number"],
                                    "Painting 적용 Temp.",
                                    maximum_original,
                                    "유효한 온도값 없음",
                                ])
                            else:
                                summary["Painting Temp. Converted"] += 1

                # ① Insulation Temperature → Display Format Insulation Temp 열
                if ins_from_line:
                    insulation_value_temp = line_operating_value
                    if insulation_value_temp is not None:
                        summary["Insulation Temp. by Line List"] += 1
                    elif line_record is not None:
                        temperature_issue_ws.append([
                            display_row_no,
                            drawing_number,
                            line_record["excel_row"],
                            line_record["line_number"],
                            "Insulation 적용 Temp.",
                            line_operating_original,
                            "유효한 온도값 없음",
                        ])
                else:
                    insulation_value_temp, _rule_row = self._value_from_display_rules(
                        row_values,
                        header_index,
                        insulation_temp_rules,
                    )
                    if insulation_value_temp is not None:
                        summary["Insulation Temp. by Rule File"] += 1
                    else:
                        temperature_issue_ws.append([
                            display_row_no,
                            drawing_number,
                            "",
                            "",
                            "Insulation 적용 Temp.",
                            "",
                            "Insulation Temp 규칙 미일치",
                        ])
                if insulation_value_temp is None:
                    summary["Insulation Temp. Issue"] += 1

                row_values[
                    header_index[
                        "Insulation Temp (Operating Temp)"
                    ]
                ] = insulation_value_temp

                # ② Paint Code → Display Format Paint Symbol 열
                class_value = (
                    row_values[class_index] if class_index is not None else None
                )
                insulation_value = (
                    row_values[insulation_index] if insulation_index is not None else None
                )
                material_group = ""
                paint_suffix = ""
                paint_type = ""
                paint_symbol = ""
                painting_issue_reasons = []

                if paint_from_rule:
                    paint_symbol, _rule_row = self._value_from_display_rules(
                        row_values,
                        header_index,
                        paint_symbol_rules,
                    )
                    paint_symbol = paint_symbol or ""
                    if paint_symbol:
                        summary["Paint Symbol by Rule File"] += 1
                    else:
                        painting_issue_reasons.append("Paint 규칙 미일치")
                else:
                    material_group = self._classify_by_rules(
                        class_value,
                        self.material_rules,
                    )
                    paint_suffix = self._classify_by_rules(
                        insulation_value,
                        self.insulation_rules,
                    )
                    paint_type = (
                        f"{material_group}-{paint_suffix}"
                        if material_group and paint_suffix
                        else ""
                    )

                    if not material_group:
                        painting_issue_reasons.append(
                            "Material Group 분류 실패"
                        )
                    if not paint_suffix:
                        painting_issue_reasons.append(
                            "Paint Suffix 분류 실패"
                        )
                    if maximum_value is None:
                        painting_issue_reasons.append(
                            "Painting 적용 Temp. 미결정"
                        )

                    if (
                        material_group
                        and paint_suffix
                        and maximum_value is not None
                    ):
                        paint_symbol = self._paint_symbol(
                            paint_type,
                            maximum_value,
                            paint_matrix,
                        )
                        if not paint_symbol:
                            painting_issue_reasons.append(
                                "Painting Code Table 매핑 실패"
                            )

                row_values[
                    header_index["Paint Symbol"]
                ] = paint_symbol

                if paint_symbol:
                    summary["Paint Symbol Determined"] += 1
                else:
                    summary["Paint Symbol Issue"] += 1

                if painting_issue_reasons:
                    painting_issue_ws.append([
                        display_row_no,
                        drawing_number,
                        clean_text(class_value),
                        clean_text(insulation_value),
                        material_group,
                        paint_suffix,
                        paint_type,
                        maximum_value,
                        "; ".join(painting_issue_reasons),
                    ])

                display_ws.append(row_values[:len(headers)])

                processed = summary["Total Display Rows"]
                if processed % 2000 == 0:
                    if estimated_total:
                        percent = min(
                            94,
                            10 + processed / estimated_total * 84,
                        )
                        item_text = (
                            f"현재 작업: {processed:,} / "
                            f"{estimated_total:,}행"
                        )
                        log_text = (
                            f"Temperature & Painting 진행: "
                            f"{processed:,} / {estimated_total:,}행"
                        )
                    else:
                        percent = min(
                            94,
                            10 + (processed % 84000) / 1000,
                        )
                        item_text = f"현재 작업: {processed:,}행 처리"
                        log_text = (
                            f"Temperature & Painting 진행: "
                            f"{processed:,}행"
                        )

                    self.set_status(
                        "Temperature & Painting",
                        item_text,
                        percent,
                    )
                    self.log(log_text)

            summary_ws = wb_out.create_sheet(
                "LineList_Mapping_Summary"
            )
            summary_ws.append(["Item", "Count"])
            for item, count in summary.items():
                summary_ws.append([item, count])
            summary_ws.append([
                "Line No. Matching Mode",
                self.paint_line_match_mode.get(),
            ])
            summary_ws.append(["AMB Replacement Temperature", ambient_temperature])
            summary_ws.append(["Line List Header Start Row", header_start])
            summary_ws.append(["Line List Header End Row", header_end])
            summary_ws.append(["Line List Data Start Row", data_start])
            summary_ws.append([
                "Sh't No. Separator",
                self.paint_sheet_separator.get(),
            ])
            summary_ws.append([
                "Insulation Temp. Source",
                self.paint_ins_temp_source.get(),
            ])
            summary_ws.append([
                "Insulation Temp. Line List Column",
                (
                    self.paint_operating_temp_col.get()
                    if self._ins_temp_uses_line_list()
                    else ""
                ),
            ])
            summary_ws.append([
                "Insulation Temp. Rule File",
                (
                    self.paint_ins_rule_file.get()
                    if self._ins_temp_uses_rule_file()
                    else ""
                ),
            ])
            summary_ws.append([
                "Painting Temp. Line List Column",
                (
                    self.paint_max_operating_temp_col.get()
                    if not self._paint_uses_rule_file()
                    else ""
                ),
            ])
            summary_ws.append([
                "Paint Symbol Source",
                self.paint_symbol_source.get(),
            ])
            summary_ws.append([
                "Paint Symbol Rule File",
                self.paint_rule_file.get() if self._paint_uses_rule_file() else "",
            ])

            wb_in.close()
            output_path.parent.mkdir(
                parents=True,
                exist_ok=True,
            )

            self.set_status(
                "Temperature & Painting",
                "현재 작업: Excel 파일 저장",
                96,
            )
            wb_out.save(output_path)

            # 자동 연계: 5번 결과를 6번 File Split & Export 입력으로 전달
            self.split_input_file.set(str(output_path))
            self.split_input_sheet.set("Display_Format")

            self.set_status(
                "완료",
                "현재 작업: 없음",
                100,
            )
            self.log(
                "Temperature & Painting 완료: "
                f"{output_path} / {summary['Total Display Rows']:,}행"
            )
            self.log("자동 연계: 6번 탭 Display Format 입력 파일로 설정")
            self.root.after(
                0,
                lambda: messagebox.showinfo(
                    "완료",
                    (
                        f"{summary['Total Display Rows']:,}행 처리 완료\n"
                        f"Line No. 미매칭: "
                        f"{summary['Line No. Unmatched']:,}건\n"
                        f"온도 이슈: "
                        f"Insulation {summary['Insulation Temp. Issue']:,}건 / "
                        f"Painting {summary['Painting Temp. Issue']:,}건\n"
                        f"Paint Symbol 미결정: "
                        f"{summary['Paint Symbol Issue']:,}건\n\n"
                        f"{output_path}"
                    ),
                ),
            )

        except Exception as exc:
            self.log(traceback.format_exc())
            self.root.after(
                0,
                lambda exc=exc: messagebox.showerror(
                    "오류",
                    str(exc),
                ),
            )
        finally:
            self.set_running(False)

    # ------------------------------------------------------------------ File Split & Export
    def _build_split_tab(self):
        output_dir = self.project_root / "04_output"
        self.split_input_file = tk.StringVar(value=str(output_dir / "Display_Format_TempPaint.xlsx"))
        self.split_input_sheet = tk.StringVar(value="Display_Format")
        self.split_output_folder = tk.StringVar(value=str(output_dir / "Final_Export"))
        self.split_base_name = tk.StringVar(value="Display_Format")
        self.split_max_rows = tk.StringVar(value="60000")
        self.split_drawing_col = tk.StringVar(value="Drawing Number")
        self.split_filename_col = tk.StringVar(value="FileName")
        self.split_keep_boundary = tk.BooleanVar(value=True)

        frame = tk.LabelFrame(self.split_tab, text="File Split & Export 전용 설정", padx=8, pady=8)
        frame.pack(fill="x", padx=10, pady=10)
        frame.columnconfigure(1, weight=1)
        self._file_row(frame, 0, "Display Format 파일", self.split_input_file, self.choose_split_input)
        tk.Label(frame, text="시트").grid(row=1, column=0, sticky="w")
        self.split_input_sheet_combo = ttk.Combobox(frame, textvariable=self.split_input_sheet, state="readonly", width=28)
        self.split_input_sheet_combo.grid(row=1, column=1, sticky="w", padx=6)
        tk.Button(frame, text="열 불러오기", command=self.load_split_columns).grid(row=1, column=4, padx=5)
        tk.Label(frame, text="저장 폴더").grid(row=2, column=0, sticky="w")
        tk.Entry(frame, textvariable=self.split_output_folder).grid(row=2, column=1, sticky="ew", padx=6)
        tk.Button(frame, text="폴더 선택", command=self.choose_split_output).grid(row=2, column=4, padx=5)

        settings = tk.LabelFrame(self.split_tab, text="분할 설정", padx=8, pady=8)
        settings.pack(fill="x", padx=10, pady=5)
        for i, (label, var) in enumerate([
            ("기본 파일명", self.split_base_name),
            ("기준 행 수", self.split_max_rows),
        ]):
            tk.Label(settings, text=label).grid(row=0, column=i * 2, padx=4, sticky="e")
            tk.Entry(settings, textvariable=var, width=24).grid(row=0, column=i * 2 + 1, padx=4)
        tk.Label(settings, text="Drawing Number 열").grid(row=1, column=0, padx=4, sticky="e")
        self.split_drawing_combo = ttk.Combobox(settings, textvariable=self.split_drawing_col, state="readonly", width=24)
        self.split_drawing_combo.grid(row=1, column=1, padx=4)
        tk.Label(settings, text="FileName 열").grid(row=1, column=2, padx=4, sticky="e")
        self.split_filename_combo = ttk.Combobox(settings, textvariable=self.split_filename_col, state="readonly", width=24)
        self.split_filename_combo.grid(row=1, column=3, padx=4)
        tk.Checkbutton(settings, text="Drawing Number가 바뀌는 첫 행에서만 파일 분할", variable=self.split_keep_boundary).grid(
            row=2, column=0, columnspan=4, sticky="w", pady=5
        )

        buttons = tk.Frame(self.split_tab)
        buttons.pack(fill="x", padx=10, pady=8)
        self.split_preview_button = tk.Button(
            buttons,
            text="분할 미리보기",
            command=lambda: self.start_thread(self.preview_split),
            width=18,
        )
        self.split_preview_button.pack(side="left")
        self.split_run_button = tk.Button(
            buttons,
            text="최종 파일 분할 저장",
            command=lambda: self.start_thread(self.run_split_export),
            width=22,
        )
        self.split_run_button.pack(side="left", padx=8)

    def choose_split_input(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm")])
        if path:
            self.split_input_file.set(path)
            self._load_sheet_names_to_combo(path, self.split_input_sheet_combo, self.split_input_sheet)

    def choose_split_output(self):
        path = filedialog.askdirectory()
        if path:
            self.split_output_folder.set(path)

    def load_split_columns(self):
        path = self.split_input_file.get()
        self._load_sheet_names_to_combo(path, self.split_input_sheet_combo, self.split_input_sheet)
        headers = self._read_headers(path, self.split_input_sheet.get(), 1)
        self.split_drawing_combo["values"] = headers
        self.split_filename_combo["values"] = headers

    def _calculate_split_groups(self):
        path = Path(self.split_input_file.get().strip())
        if not path.exists():
            raise ValueError("Display Format 파일을 선택하세요.")
        max_rows = int(self.split_max_rows.get())
        if max_rows <= 0:
            raise ValueError("기준 행 수는 1 이상이어야 합니다.")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[self.split_input_sheet.get()]
        iterator = ws.iter_rows(values_only=True)
        headers = make_unique_headers(next(iterator))
        if self.split_drawing_col.get() not in headers:
            raise ValueError("Drawing Number 열을 찾을 수 없습니다.")
        drawing_idx = headers.index(self.split_drawing_col.get())
        groups = []
        current = []
        current_drawing = None
        processed = 0
        raw_max_row = ws.max_row
        if isinstance(raw_max_row, int) and raw_max_row > 1:
            estimated_total = raw_max_row - 1
        else:
            estimated_total = None

        for values in iterator:
            if is_effectively_empty_row(values):
                continue
            processed += 1
            if processed % 2000 == 0:
                if estimated_total:
                    progress = min(
                        90,
                        processed / estimated_total * 90,
                    )
                    item_text = (
                        f"현재 작업: {processed:,} / "
                        f"{estimated_total:,}행"
                    )
                else:
                    progress = min(
                        90,
                        5 + (processed % 85000) / 1000,
                    )
                    item_text = f"현재 작업: {processed:,}행 처리"

                self.set_status(
                    "File Split 계산",
                    item_text,
                    progress,
                )
                self.log(
                    f"File Split 경계 계산 진행: {processed:,}행"
                )
            values = list(values) + [None] * (len(headers) - len(values))
            drawing = clean_text(values[drawing_idx])
            should_split = False
            if current and len(current) >= max_rows:
                if self.split_keep_boundary.get():
                    should_split = drawing != current_drawing
                else:
                    should_split = True
            if should_split:
                groups.append(current)
                current = []
            current.append(values[:len(headers)])
            current_drawing = drawing
        if current:
            groups.append(current)
        wb.close()
        return headers, groups

    def preview_split(self):
        self.set_running(True)
        self.set_status(
            "File Split 미리보기",
            "현재 작업: 분할 경계 계산",
            0,
        )
        try:
            _headers, groups = self._calculate_split_groups()
            lines = [f"예상 파일 수: {len(groups)}개"]
            for i, group in enumerate(groups[:20], start=1):
                lines.append(f"{i:02d}: {len(group):,}행")
            if len(groups) > 20:
                lines.append("...")
            self.set_status("완료", "현재 작업: 없음", 100)
            self.root.after(
                0,
                lambda: messagebox.showinfo(
                    "분할 미리보기",
                    "\n".join(lines),
                ),
            )
        except Exception as exc:
            self.log(traceback.format_exc())
            self.root.after(
                0,
                lambda exc=exc: messagebox.showerror("오류", str(exc)),
            )
        finally:
            self.set_running(False)

    def run_split_export(self):
        self.set_running(True)
        self.set_status(
            "File Split & Export",
            "현재 작업: 분할 경계 계산",
            0,
        )
        try:
            headers, groups = self._calculate_split_groups()
            output_folder = Path(self.split_output_folder.get().strip())
            output_folder.mkdir(parents=True, exist_ok=True)
            base = self.split_base_name.get().strip() or "Display_Format"
            filename_col = self.split_filename_col.get()
            filename_idx = (
                headers.index(filename_col)
                if filename_col in headers
                else None
            )
            no_idx = headers.index("NO") if "NO" in headers else None

            created = []
            for i, group in enumerate(groups, start=1):
                filename = f"{base}_{i:02d}.xlsx"

                # Each split file starts its NO column again from 1.
                for local_no, row in enumerate(group, start=1):
                    if filename_idx is not None:
                        row[filename_idx] = filename
                    if no_idx is not None:
                        row[no_idx] = local_no

                wb = Workbook(write_only=True)
                ws = wb.create_sheet("Display_Format")
                ws.append(headers)
                for row in group:
                    ws.append(row)
                output_path = output_folder / filename
                wb.save(output_path)
                created.append((filename, len(group)))
                self.set_status("File Split & Export", f"현재 작업: {filename}", i / len(groups) * 100)

            summary_path = output_folder / f"{base}_Split_Summary.xlsx"
            swb = Workbook()
            sws = swb.active
            sws.title = "Split_Summary"
            sws.append(["File Name", "Data Rows", "NO Numbering"])
            for filename, count in created:
                sws.append([filename, count, f"1 ~ {count:,}"])
            style_header(sws)
            auto_fit_columns(sws)
            swb.save(summary_path)

            self.set_status("완료", "현재 작업: 없음", 100)
            self.root.after(0, lambda: messagebox.showinfo(
                "완료",
                f"{len(created)}개 파일 생성 완료\n\n저장 폴더:\n{output_folder}",
            ))
        except Exception as exc:
            self.log(traceback.format_exc())
            self.root.after(0, lambda exc=exc: messagebox.showerror("오류", str(exc)))
        finally:
            self.set_running(False)


if __name__ == "__main__":
    root = tk.Tk()
    app = BmDmcsTool(root)
    root.mainloop()
